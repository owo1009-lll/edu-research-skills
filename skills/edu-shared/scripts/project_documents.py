"""Read only explicitly supplied files into addressable units; report coverage gaps."""
import csv
import io
import json
from pathlib import Path
import re
import zipfile

from private_store import digest


def unit(locator, text, kind="text", **meta):
    return {"locator": locator, "text": str(text), "kind": kind, **meta}


def json_units(value, prefix=""):
    if isinstance(value, dict):
        for key, child in value.items():
            escaped = str(key).replace("~", "~0").replace("/", "~1")
            yield from json_units(child, prefix + "/" + escaped)
    elif isinstance(value, list):
        for i, child in enumerate(value):
            yield from json_units(child, prefix + "/" + str(i))
    else:
        yield unit("json:" + prefix, "" if value is None else value, "scalar")


def read_document(path):
    """The caller supplies scope. No recursive directory search and no external links."""
    path = Path(path)
    result = {"path": str(path.resolve()), "sha256": None, "units": [], "unchecked": [],
              "status": "read"}
    if not path.is_file():
        result["status"] = "missing"
        result["unchecked"].append({"location": str(path), "reason": "File is missing; no directory search performed."})
        return result
    data = path.read_bytes()
    result["sha256"] = digest(data)
    out, gaps = result["units"], result["unchecked"]
    ext = path.suffix.lower()
    try:
        if ext in {".md", ".txt"}:
            for n, line in enumerate(data.decode("utf-8-sig").splitlines(), 1):
                if line.strip():
                    out.append(unit(f"line:{n}", line))
        elif ext == ".json":
            out.extend(json_units(json.loads(data.decode("utf-8-sig"))))
        elif ext == ".csv":
            rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
            for r, row in enumerate(rows, 1):
                for c, value in enumerate(row, 1):
                    out.append(unit(f"csv:{r}:{c}", value, "cell",
                                    row_context=row, column=rows[0][c - 1] if len(rows[0]) >= c else str(c)))
        elif ext == ".xlsx":
            from openpyxl import load_workbook
            book = load_workbook(path, data_only=False, read_only=False, keep_links=False)
            cached = load_workbook(path, data_only=True, read_only=False, keep_links=False)
            for sheet in book:
                for row in sheet:
                    for cell in row:
                        if cell.value is None:
                            continue
                        if cell.data_type == "f":
                            value = cached[sheet.title][cell.coordinate].value
                            gaps.append({"location": f"{sheet.title}!{cell.coordinate}",
                                         "reason": "Formula not recalculated; cached value is not authoritative."})
                            out.append(unit(f"xlsx:{sheet.title}!{cell.coordinate}", value,
                                            "formula", formula=cell.value))
                        else:
                            out.append(unit(f"xlsx:{sheet.title}!{cell.coordinate}", cell.value, "cell"))
                if sheet._charts or sheet._images:
                    gaps.append({"location": sheet.title, "reason": "Embedded Excel charts/images not parsed."})
            book.close()
            cached.close()
        elif ext == ".docx":
            from lxml import etree as ET
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            with zipfile.ZipFile(path) as z:
                root = ET.fromstring(z.read("word/document.xml"),
                                     ET.XMLParser(resolve_entities=False, no_network=True))
                for n, p in enumerate(root.xpath(".//w:p", namespaces=ns)):
                    text = "".join(p.xpath(".//w:t/text()", namespaces=ns))
                    risky = bool(p.xpath(".//w:fldChar|.//w:instrText|.//w:ins|.//w:del|.//w:drawing|.//w:txbxContent",
                                         namespaces=ns))
                    in_table = bool(p.xpath("ancestor::w:tc", namespaces=ns))
                    out.append(unit(f"docx:p:{n}", text, "table_paragraph" if in_table else "paragraph",
                                    editable=not risky))
                    if risky:
                        gaps.append({"location": f"docx:p:{n}",
                                     "reason": "Field, revision, drawing or textbox: text is a clue; embedded content requires review."})
                for name in z.namelist():
                    if re.match(r"word/(header\d+|footer\d+|footnotes|endnotes)\.xml$", name):
                        gaps.append({"location": name, "reason": "Ancillary Word part not checked."})
        elif ext == ".pptx":
            from pptx import Presentation
            prs = Presentation(path)

            def shapes(items, slide, prefix=""):
                for shape in items:
                    nodes = shape._element.xpath(".//p:cNvPr")
                    try:
                        ident = json.loads(nodes[0].get("descr", "{}")).get("edu_ppt", {}).get("id")
                    except (ValueError, IndexError):
                        ident = None
                    loc = f"ppt:{slide}:{prefix}{ident or shape.shape_id}"
                    if shape.shape_type == 6:
                        shapes(shape.shapes, slide, prefix + str(shape.shape_id) + "/")
                    elif shape.has_chart:
                        chart = shape.chart
                        for i, cat in enumerate(chart.plots[0].categories):
                            out.append(unit(f"{loc}:category:{i}", cat.label, "chart_category"))
                        for s, series in enumerate(chart.series):
                            out.append(unit(f"{loc}:series:{s}:name", series.name, "chart_series"))
                            for i, val in enumerate(series.values):
                                out.append(unit(f"{loc}:series:{s}:value:{i}", val, "chart_value"))
                        try:
                            axis = chart.value_axis
                            for key in ("minimum_scale", "maximum_scale", "major_unit"):
                                out.append(unit(f"{loc}:axis:{key}", getattr(axis, key), "axis"))
                            if axis.has_title:
                                out.append(unit(f"{loc}:axis:title", axis.axis_title.text_frame.text, "unit"))
                        except (ValueError, AttributeError):
                            gaps.append({"location": loc, "reason": "Chart axes unavailable; check chart with its plan."})
                        gaps.append({"location": loc, "reason": "Chart label formatting/geometry require edu-ppt check_deck and rendering."})
                    elif shape.has_table:
                        for r, row in enumerate(shape.table.rows):
                            for c, cell in enumerate(row.cells):
                                out.append(unit(f"{loc}:cell:{r}:{c}", cell.text, "cell"))
                    elif shape.has_text_frame:
                        out.append(unit(loc + ":text", shape.text, "slide_text"))
                    elif shape.shape_type == 13:
                        gaps.append({"location": loc, "reason": "Image/raster chart: values not read; use source asset/plan."})
                    else:
                        gaps.append({"location": loc, "reason": "Nontext drawing: geometry/meaning not checked."})
            for n, slide in enumerate(prs.slides, 1):
                shapes(slide.shapes, n)
                if slide.has_notes_slide:
                    out.append(unit(f"ppt:{n}:notes", slide.notes_slide.notes_text_frame.text, "notes"))
        elif ext == ".pdf":
            import fitz
            with fitz.open(path) as doc:
                for n, page in enumerate(doc, 1):
                    blocks = page.get_text("blocks")
                    if not any(str(b[4]).strip() for b in blocks if b[6] == 0):
                        gaps.append({"location": f"pdf:page:{n}", "reason": "No readable text; scanned/image-only page not checked."})
                    for i, block in enumerate(blocks):
                        if block[6] == 0:
                            out.append(unit(f"pdf:{n}:block:{i}", block[4].strip(), "pdf_text",
                                            bbox=list(block[:4])))
                    if page.get_images():
                        gaps.append({"location": f"pdf:page:{n}", "reason": "Image content not checked; text extraction does not establish table reading order."})
        else:
            raise ValueError(f"Unsupported extension: {ext}")
    except Exception as error:
        # Keep successfully parsed units while explicitly exposing partial coverage.
        result["status"] = "partial" if out else "failed"
        gaps.append({"location": str(path), "reason": f"{type(error).__name__}: {error}"})
    locs = [x["locator"] for x in out]
    if len(locs) != len(set(locs)):
        result["status"] = "partial"
        gaps.append({"location": str(path), "reason": "Duplicate native object IDs; matching is ambiguous."})
    return result


def locate(document, binding, previous=None):
    matches = [u for u in document["units"] if u["locator"] == binding["locator"]]
    if len(matches) != 1:
        matches = []
    current = matches[0] if matches else None
    if previous and previous.get("text") is not None and (not current or current["text"] != previous["text"]):
        # Exact full-unit relocation only for prose. Repeated quotations are ambiguous.
        if previous.get("kind") in {"text", "paragraph", "table_paragraph", "pdf_text"}:
            candidates = [u for u in document["units"] if u["text"] == previous["text"]]
            if len(candidates) == 1:
                return candidates[0], "relocated"
            if len(candidates) > 1:
                return None, "ambiguous relocation"
    if current is None:
        return None, "location missing or duplicated"
    return current, "located"


def replace_docx(source, edits, output):
    """Surgical text edits on plain runs. Preserve ZIP parts, comments and formatting."""
    from lxml import etree as ET
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    with zipfile.ZipFile(source) as z:
        entries = [(i, z.read(i.filename)) for i in z.infolist()]
    if any(i.filename.startswith("_xmlsignatures/") for i, _ in entries):
        raise ValueError("Signed documents cannot be edited.")
    body = dict((i.filename, data) for i, data in entries)["word/document.xml"]
    root = ET.fromstring(body, ET.XMLParser(resolve_entities=False, no_network=True))
    paragraphs = root.xpath(".//w:p", namespaces=ns)
    for edit in edits:
        p = paragraphs[int(edit["locator"].split(":")[-1])]
        if p.xpath(".//w:fldChar|.//w:ins|.//w:del|.//w:drawing|.//w:txbxContent", namespaces=ns):
            raise ValueError("Complex Word object requires manual editing on a copy.")
        nodes = p.xpath(".//w:t", namespaces=ns)
        text = "".join(n.text or "" for n in nodes)
        if text != edit["before"] or text.count(edit["old"]) != 1:
            raise ValueError("Stale or ambiguous Word edit.")
        start = text.index(edit["old"])
        end, offset, inserted = start + len(edit["old"]), 0, False
        for node in nodes:
            original = node.text or ""
            a, b = offset, offset + len(original)
            if b > start and a < end:
                left = original[:max(0, start - a)]
                right = original[max(0, end - a):] if b > end else ""
                node.text = left + (edit["new"] if not inserted else "") + right
                node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                inserted = True
            offset = b
    payload = ET.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)
    with zipfile.ZipFile(output, "x", zipfile.ZIP_DEFLATED) as z:
        for info, data in entries:
            z.writestr(info, payload if info.filename == "word/document.xml" else data)
