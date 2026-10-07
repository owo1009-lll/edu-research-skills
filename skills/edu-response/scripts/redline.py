"""Mark the changes between an original and a revised manuscript in red (修改稿标红).

    python redline.py --original 原稿.docx|.md --revised 修改稿.docx|.md --output 修改稿_标红.docx
                      [--show-deletions] [--lang zh|en]

Paragraphs and table rows are aligned in document order; tables come out as Word tables in their original place.
Inside a changed paragraph or cell, Chinese text is compared character by character (numbers and Latin words as
whole tokens, short unchanged gaps between two changes merged into the change), English text word by word. Added text
is red; with --show-deletions, removed text is shown in red strikethrough (many Chinese journals want only the
additions marked). Also writes <output>.changes.csv listing each changed paragraph or row, to help fill in the
modification locations. Character formatting and layout of the original are not kept; when they matter, use Word's
审阅 > 比较 on the two files instead.
"""
from pathlib import Path
import argparse
import csv
import difflib
import re

RED = (0xC0, 0x00, 0x00)


def blocks(path):
    """Return [('p', level, text)] and [('row', table_no, [cells])] in document order; level 0 is the title."""
    out = []
    if path.suffix.lower() == '.docx':
        import docx
        from docx.oxml.ns import qn
        d = docx.Document(path)
        styles = {s.style_id: s.name for s in d.styles}
        tables = 0
        for child in d.element.body.iterchildren():
            if child.tag == qn('w:p'):
                text = ''.join(t.text or '' for t in child.iter(qn('w:t'))).strip()
                if not text:
                    continue
                style = child.find(qn('w:pPr') + '/' + qn('w:pStyle'))
                name = styles.get(style.get(qn('w:val')), '') if style is not None else ''
                m = re.match(r'Heading (\d)', name)
                out.append(('p', 0 if name == 'Title' else min(int(m.group(1)), 3) if m else None, text))
            elif child.tag == qn('w:tbl'):
                tables += 1
                for tr in child.iter(qn('w:tr')):
                    out.append(('row', tables, [''.join(t.text or '' for t in tc.iter(qn('w:t'))).strip() for tc in tr.iter(qn('w:tc'))]))
        return out
    tables, in_table = 0, False
    for line in path.read_text(encoding='utf-8').splitlines():
        stripped = line.strip()
        if stripped.startswith('|'):
            if not in_table:
                tables, in_table = tables + 1, True
            if not re.fullmatch(r'\|?[\s:|-]+\|?', stripped):
                out.append(('row', tables, [c.strip() for c in stripped.strip('|').split('|')]))
            continue
        in_table = False
        if not stripped:
            continue
        m = re.match(r'^(#{1,6})\s+(.*)$', line)
        level = (0 if len(m.group(1)) == 1 and not out else min(len(m.group(1)), 3)) if m else None  # leading # is the title
        out.append(('p', level, m.group(2).strip() if m else stripped))
    return out


def key(b):
    return b[2] if b[0] == 'p' else ' | '.join(b[2])


def tokens(text, lang):
    if lang == 'zh':
        return re.findall(r'\d+(?:[.,]\d+)*%?|[A-Za-z]+(?:[\'-][A-Za-z]+)*|\s+|.', text)
    return re.findall(r'\s+|\w+|[^\w\s]', text)


def diff_segments(old, new, lang):
    """[(kind, text)] with kind equal / insert / delete; short equal gaps between changes are merged into them."""
    a, b = tokens(old, lang), tokens(new, lang)
    ops = [list(op) for op in difflib.SequenceMatcher(None, a, b, autojunk=False).get_opcodes()]
    for i in range(1, len(ops) - 1):
        if ops[i][0] == 'equal' and ops[i - 1][0] != 'equal' and ops[i + 1][0] != 'equal' \
                and len(''.join(a[ops[i][1]:ops[i][2]]).strip()) <= 2:
            ops[i][0] = 'replace'  # a one- or two-character gap between two changes reads better as one change
    merged = []
    for tag, a1, a2, b1, b2 in ops:
        if merged and tag != 'equal' and merged[-1][0] != 'equal':
            merged[-1] = ['replace', merged[-1][1], a2, merged[-1][3], b2]
        else:
            merged.append([tag, a1, a2, b1, b2])
    out = []
    for tag, a1, a2, b1, b2 in merged:
        if tag == 'equal':
            out.append(('equal', ''.join(b[b1:b2])))
        else:
            out.append(('delete', ''.join(a[a1:a2])))
            out.append(('insert', ''.join(b[b1:b2])))
    return [(k, t) for k, t in out if t]


def add(paragraph, text, kind, zh):
    from docx.shared import RGBColor
    from docx.oxml.ns import qn
    if not text:
        return
    run = paragraph.add_run(text)
    if zh:
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    if kind in ('insert', 'delete'):
        run.font.color.rgb = RGBColor(*RED)
    if kind == 'delete':
        run.font.strike = True


def fill(paragraph, old, new, show_deletions, zh, lang):
    """Write new text into a paragraph, marking the differences from old (None: everything is new)."""
    if old is None:
        add(paragraph, new, 'insert', zh)
        return
    for kind, text in diff_segments(old, new, lang):
        if kind != 'delete' or show_deletions:
            add(paragraph, text, kind, zh)


def redline(original, revised, output, show_deletions=False, lang=None):
    from docx import Document
    from docx.oxml.ns import qn
    old, new = blocks(Path(original)), blocks(Path(revised))
    sample = ''.join(key(b) for b in new[:50])
    lang = lang or ('zh' if len(re.findall(r'[一-鿿]', sample)) > len(re.findall(r'[A-Za-z]+', sample)) else 'en')
    zh = lang == 'zh'
    doc = Document()
    if zh:
        doc.styles['Normal'].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    changes, state = [], {'table': None, 'table_no': None}

    def where():
        return f'table row {len(state["table"].rows)}' if state['table'] is not None else f'paragraph {len(doc.paragraphs)}'

    def emit(o, n, kind):
        """o, n: blocks or None. kind: equal / modified / added / deleted."""
        b = n or o
        if b[0] == 'row':
            if state['table'] is None or state['table_no'] != b[1]:
                state['table'] = doc.add_table(rows=0, cols=max(1, len(b[2])))
                state['table'].style = 'Table Grid'
                state['table_no'] = b[1]
            cells = n[2] if n else o[2]
            row = state['table'].add_row().cells
            for i, text in enumerate(cells[:len(row)]):
                p = row[i].paragraphs[0]
                if kind == 'deleted':
                    add(p, text, 'delete', zh)
                elif kind == 'modified' and o[0] == 'row' and len(o[2]) == len(cells):
                    fill(p, o[2][i], text, show_deletions, zh, lang)
                else:
                    fill(p, None if kind in ('added', 'modified') else text, text, show_deletions, zh, lang)
        else:
            state['table'] = state['table_no'] = None
            p = doc.add_heading('', b[1]) if b[1] is not None else doc.add_paragraph()
            if kind == 'deleted':
                add(p, o[2], 'delete', zh)
            elif kind == 'modified' and o[0] == 'p':
                fill(p, o[2], n[2], show_deletions, zh, lang)
            else:
                fill(p, None if kind == 'added' else n[2], n[2], show_deletions, zh, lang)
        if kind != 'equal':
            changes.append({'location': where(), 'type': kind, 'original': key(o)[:200] if o else '', 'revised': key(n)[:200] if n else ''})

    sm = difflib.SequenceMatcher(None, [key(b) for b in old], [key(b) for b in new], autojunk=False)
    for op, i1, i2, j1, j2 in sm.get_opcodes():
        if op == 'equal':
            for b in new[j1:j2]:
                emit(None, b, 'equal')
            continue
        olds, news = old[i1:i2], new[j1:j2]
        for k in range(max(len(olds), len(news))):
            o = olds[k] if k < len(olds) else None
            n = news[k] if k < len(news) else None
            if o and n and (o[0] != n[0] or difflib.SequenceMatcher(None, key(o), key(n), autojunk=False).ratio() < 0.4):
                if show_deletions:
                    emit(o, None, 'deleted')
                else:
                    changes.append({'location': where(), 'type': 'deleted', 'original': key(o)[:200], 'revised': ''})
                emit(None, n, 'added')
            elif o and n:
                emit(o, n, 'modified')
            elif n:
                emit(None, n, 'added')
            elif show_deletions:
                emit(o, None, 'deleted')
            else:
                changes.append({'location': where(), 'type': 'deleted', 'original': key(o)[:200], 'revised': ''})
    Path(output).parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    with open(str(output) + '.changes.csv', 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=['location', 'type', 'original', 'revised'])
        w.writeheader()
        w.writerows(changes)
    return changes


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--original', required=True, type=Path)
    ap.add_argument('--revised', required=True, type=Path)
    ap.add_argument('--output', required=True, type=Path)
    ap.add_argument('--show-deletions', action='store_true')
    ap.add_argument('--lang', choices=['zh', 'en'])
    a = ap.parse_args()
    changes = redline(a.original, a.revised, a.output, a.show_deletions, a.lang)
    kinds = {k: sum(c['type'] == k for c in changes) for k in ('modified', 'added', 'deleted')}
    print(f'{a.output}: ' + ', '.join(f'{v} {k}' for k, v in kinds.items()) + f'; list in {a.output}.changes.csv')


if __name__ == '__main__':
    main()
