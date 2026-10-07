"""Check a response letter or 修改说明 against the comment list and the revised manuscript.

    python check_response.py --comments comments.csv --response 修改说明.md|.docx --revised 修改稿.docx|.md [--json]

comments.csv needs columns id and comment (verbatim); ids look like R1-1, R2-3a or E-1. In the response an id may
appear as R1-1, 1-1, 1.1 or "Comment 1.1" at the start of a line, a table cell or after 意见/Comment.
For every comment the script reports: no response found; no location in the revised manuscript; quoted revised
text that does not occur in the revised manuscript (quotes of deleted original wording are skipped); numbers that do
not occur in the revised manuscript; unfinished placeholders such as 【待补…】; defensive wording; deferral to future
research. It also counts placeholders left in the response and the revised manuscript: both must be zero before
submission. These are prompts to check, not verdicts.
"""
from pathlib import Path
import argparse
import csv
import json
import re
import sys

LOCATION = (r'(?:第\s*\d+\s*[页行]|第[一二三四五六七八九十\d]+\s*(?:部分|节|段|章|自然段)|[（(][一二三四五六七八九十]+[）)]|[表图]\s*\d+|'
            r'(?<![\d.])\d+(?:\.\d+)+\s*节|(?<![\d.])\d+\s*节|'  # numbered sections: 3.7节, 4节
            r'摘要|引言|问题提出|文献综述|研究设计|研究方法|研究结果|讨论|结论|参考文献|附录|标题|题目|关键词|声明|致谢|'
            r'\bpages?\s*\d+|\bp\.\s*\d+|\blines?\s*\d+|\bSection\s*\d+(?:\.\d+)*|\bTable\s*\d+|\bFigure\s*\d+|\bAppendix|'
            r'\bAbstract|\bIntroduction|\bMethods?\b|\bResults\b|\bDiscussion|\bConclusions?\b|\bLimitations|'
            r'修改说明|修改稿|标红|全文|全稿|通篇|致编辑|\bthroughout\b|response letter|cover letter|red-marked)')
KEEP = r'(?:未作修改|未修改|不作修改|保留原|维持原|not (?:been )?changed|no change|did not change|retain(?:ed)? the original)'
DEFENSIVE = (r'(?:误解|误读|没有注意到|未注意到|没有看到|显然|众所周知|毫无疑问|审稿(?:专家|人)(?:可能)?(?:并)?不了解|'
             r'misunderst|misread|\bmissed\b|obviously|\bclearly,|it is clear that|as is well known|failed to notice|did not notice)')
DEFER = r'(?:未来研究|后续研究|今后的研究|将来的研究|future (?:research|work|stud)|beyond the scope)'
NUM = r'(?<![A-Za-z0-9_.])(\d+(?:\.\d+)?|\.\d+)(?![A-Za-z0-9_])'
QUOTE = r'[“"「『]([^”"」』]{10,})[”"」』]'
DELETED = r'(?:删去|删除|去掉|原文|原稿|原句|原来|原为|改为前|deleted|removed|original(?:ly)?|previously)[^。.；;]{0,12}$'
PLACEHOLDER = r'(?:【待[补核填][^】]*】|〔待[补核填][^〕]*〕|\[待[补核填][^\]]*\]|\bTBD\b|\bTODO\b|_{3,}|＿{2,}|X\.XX)'
DONE = r'(?:已补充|已增加|已增补|已重新|已报告|已改用|已采用|已完成|have added|added|re-?ran|now report)'


def read(path):
    """Text of a .docx in document order (headings as '#' lines, table rows as '| a | b |') or of a text file."""
    if path.suffix.lower() == '.docx':
        import docx
        from docx.oxml.ns import qn
        d = docx.Document(path)
        names = {s.style_id: s.name for s in d.styles}
        parts = []
        for child in d.element.body.iterchildren():
            if child.tag == qn('w:p'):
                text = ''.join(t.text or '' for t in child.iter(qn('w:t')))
                style = child.find(qn('w:pPr') + '/' + qn('w:pStyle'))
                name = names.get(style.get(qn('w:val')), '') if style is not None else ''
                level = 1 if name == 'Title' else int(name[8:]) if re.fullmatch(r'Heading \d', name) else 0
                parts.append('#' * level + ' ' + text if level and text.strip() else text)
            elif child.tag == qn('w:tbl'):
                for tr in child.iter(qn('w:tr')):
                    cells = [''.join(t.text or '' for t in tc.iter(qn('w:t'))) for tc in tr.iter(qn('w:tc'))]
                    parts.append('| ' + ' | '.join(cells) + ' |')
        return '\n'.join(parts)
    return path.read_text(encoding='utf-8', errors='ignore')


def squash(s):
    """Compare text ignoring spaces, line breaks and full-/half-width punctuation differences."""
    table = str.maketrans('，。；：！？（）【】“”‘’', ',.;:!?()[]""\'\'')
    return re.sub(r'[\s*_>#|]+', '', s.translate(table))


def variants(cid):
    m = re.fullmatch(r'R(\d+)-(\w+)', cid, re.I)
    if m:
        r, k = m.groups()
        return [f'R{r}-{k}', f'R{r}.{k}', f'{r}-{k}', f'{r}.{k}']
    m = re.fullmatch(r'E-?(\w+)', cid, re.I)
    if m:
        return [f'E-{m.group(1)}', f'E.{m.group(1)}', f'E{m.group(1)}']
    return [cid]


def blocks(text, ids):
    """Map each id to the response text between its first marker and the next marker of any id."""
    marks = []
    for cid in ids:
        alt = '|'.join(re.escape(v) for v in sorted(variants(cid), key=len, reverse=True))
        pat = rf'(?:^|\|)\s*(?:\*\*)?\s*(?:Comment\s*|意见\s*|审稿意见\s*|序号\s*)?(?:{alt})(?![\d.\-]*\d)(?![a-z])'
        m = re.search(pat, text, re.M | re.I)
        if m:
            marks.append((m.start(), cid))
    marks.sort()
    out = {}
    for i, (start, cid) in enumerate(marks):
        end = marks[i + 1][0] if i + 1 < len(marks) else len(text)
        heading = re.search(r'^#{1,3}\s', text[start + 1:end], re.M)  # a following section heading ends the reply
        out[cid] = text[start:start + 1 + heading.start()] if heading else text[start:end]
    return out


def numbers(s):
    return [m.group(1) for m in re.finditer(NUM, s)]


def norm(n):
    if '.' in n:
        n = n.rstrip('0').rstrip('.') or '0'
    return n[1:] if n.startswith('0.') else n


def check(comments, response, revised):
    ids = [c['id'].strip() for c in comments]
    found = blocks(response, ids)
    rev_flat = squash(revised)
    rev_values = {norm(n) for n in numbers(revised)}
    rev_floats = [float(v) for v in rev_values if v]
    report = []
    for c in comments:
        cid = c['id'].strip()
        item = {'id': cid, 'issues': []}
        block = found.get(cid)
        if block is None:
            item['issues'].append('no response found')
            report.append(item)
            continue
        own = block
        comment = c.get('comment', '').strip()
        if comment:
            own = own.replace(comment, ' ')
            flat_comment = squash(comment)
        else:
            flat_comment = ''
        if len(squash(own)) < 25:
            item['issues'].append('very short response')
        if not re.search(LOCATION, own, re.I) and not re.search(KEEP, own, re.I):
            item['issues'].append('no location in the revised manuscript')
        for m in re.finditer(QUOTE, own):
            q, fq = m.group(1), squash(m.group(1))
            if re.search(DELETED, own[max(0, m.start() - 20):m.start()], re.I) or re.search(PLACEHOLDER, q):
                continue  # quoting the deleted wording, or a pending placeholder reported below
            if fq and fq not in rev_flat and fq not in flat_comment:
                item['issues'].append(f'quoted text not found in the revised manuscript: "{q[:60]}"')
        if re.search(PLACEHOLDER, own):
            done = ' that claims the change is done' if re.search(DONE, own, re.I) else ''
            item['issues'].append('unfinished placeholder in the reply' + done)
        stripped = re.sub(LOCATION, ' ', own, flags=re.I)
        stripped = re.sub(r'(?:^|\|)\s*(?:\*\*)?\s*(?:Comment\s*)?[RE]?\d+[.\-]\w+', ' ', stripped, flags=re.M | re.I)
        for n in numbers(stripped):
            v = norm(n)
            if v in rev_values or re.fullmatch(r'\d', v) or re.fullmatch(r'(19|20)\d\d', v) or v in flat_comment:
                continue
            d = len(v.split('.')[1]) if '.' in v else 0
            target = float('0' + v if v.startswith('.') else v)
            if any(abs(round(f * k, d) - target) < 1e-9 for f in rev_floats for k in (1, 100, 0.01)):
                continue
            item['issues'].append(f'number not found in the revised manuscript: {n}')
        for pattern, label in [(DEFENSIVE, 'defensive wording'),
                               (DEFER, 'deferred to future research (fine only for requests this study cannot meet)')]:
            m = re.search(pattern, own, re.I)
            if m:
                item['issues'].append(f'{label}: "{m.group(0)}"')
        report.append(item)
    return report


def main():
    sys.stdout.reconfigure(encoding='utf-8')
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--comments', required=True, type=Path)
    ap.add_argument('--response', required=True, type=Path)
    ap.add_argument('--revised', required=True, type=Path)
    ap.add_argument('--json', action='store_true')
    a = ap.parse_args()
    with open(a.comments, encoding='utf-8-sig', newline='') as f:
        comments = [r for r in csv.DictReader(f) if r.get('id', '').strip()]
    if not comments:
        sys.exit('comments.csv has no rows with an id')
    response, revised = read(a.response), read(a.revised)
    report = check(comments, response, revised)
    pending = {'response': len(re.findall(PLACEHOLDER, response)), 'revised': len(re.findall(PLACEHOLDER, revised))}
    if a.json:
        print(json.dumps({'items': report, 'placeholders': pending}, ensure_ascii=False, indent=1))
        return
    answered = sum('no response found' not in r['issues'] for r in report)
    print(f'Comments: {len(report)} | answered: {answered} | with prompts to check: {sum(bool(r["issues"]) for r in report)}')
    if any(pending.values()):
        print(f'NOT READY TO SUBMIT: placeholders left in the response {pending["response"]}, '
              f'in the revised manuscript {pending["revised"]}')
    for r in report:
        if r['issues']:
            print(f'  {r["id"]}: ' + '; '.join(r['issues']))


if __name__ == '__main__':
    main()
