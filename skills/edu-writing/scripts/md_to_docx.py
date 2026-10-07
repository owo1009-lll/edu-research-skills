"""Convert a Markdown draft (results section, report, grant application, coding report) into Word.

    python md_to_docx.py --input results.md --output results.docx [--lang zh|en] [--superscript-citations]

Supports headings (#-####), paragraphs, **bold**, pipe tables (three-line tables for Chinese), images
written as ![caption](path) on their own line, and "- " bullet lines. Chinese output uses SimSun body text
and SimHei headings; --superscript-citations turns GB/T 7714 markers such as [1] or [2-4] into superscripts (not in the reference list under a 参考文献/References heading).
Full papers with verified references and numbered displays use manuscript_package.py instead.
"""
from pathlib import Path
import argparse
import re

CITE = re.compile(r'(\[\d+(?:[,–-]\d+)*\])')
BOLD = re.compile(r'(\*\*[^*]+\*\*)')


def three_line(table, OxmlElement, qn):
    last = len(table.rows) - 1
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            edges = {}
            if r == 0:
                edges.update(top=12, bottom=6)
            if r == last:
                edges['bottom'] = 12
            box = OxmlElement('w:tcBorders')
            for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                e = OxmlElement('w:' + edge)
                e.set(qn('w:val'), 'single' if edge in edges else 'nil')
                if edge in edges:
                    e.set(qn('w:sz'), str(edges[edge])); e.set(qn('w:color'), '000000')
                box.append(e)
            cell._tc.get_or_add_tcPr().append(box)


def add_runs(paragraph, text, superscript):
    for part in BOLD.split(text):
        if not part:
            continue
        bold = part.startswith('**') and part.endswith('**')
        part = part[2:-2] if bold else part
        for piece in (CITE.split(part) if superscript else [part]):
            if piece:
                run = paragraph.add_run(piece)
                run.bold = bold
                run.font.superscript = superscript and bool(CITE.fullmatch(piece))


def convert(src, out, lang='zh', superscript=False):
    from docx import Document
    from docx.shared import Pt, Inches, Mm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(10.5 if lang == 'zh' else 11)
    normal.paragraph_format.line_spacing = 1.5 if lang == 'zh' else 1.15
    if lang == 'zh':
        for name, font in [('Normal', '宋体'), ('Heading 1', '黑体'), ('Heading 2', '黑体'), ('Heading 3', '黑体'), ('Title', '黑体')]:
            doc.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font)
    lines = Path(src).read_text(encoding='utf-8').splitlines()
    i, in_refs = 0, False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1; continue
        m = re.match(r'^(#{1,4})\s+(.*)$', line)
        img = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', line.strip())
        if m:
            level = len(m.group(1))
            in_refs = bool(re.fullmatch(r'(参考文献|references|bibliography)', m.group(2).strip(), re.I))
            doc.add_heading(m.group(2).strip(), 0 if level == 1 and i == 0 else min(level, 3))
            i += 1
        elif img:
            path = (Path(src).parent / img.group(2)).resolve()
            if path.is_file():
                doc.add_picture(str(path), width=Inches(6))
            if img.group(1):
                p = doc.add_paragraph(); add_runs(p, img.group(1), superscript); p.alignment = 1
            i += 1
        elif line.lstrip().startswith('|'):
            block = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                block.append(lines[i].strip()); i += 1
            rows = [[c.strip() for c in r.strip('|').split('|')] for r in block if not re.fullmatch(r'\|?[\s:|-]+\|?', r)]
            if not rows:
                continue
            table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
            if lang != 'zh':
                table.style = 'Table Grid'
            for r, values in enumerate(rows):
                for c, value in enumerate(values):
                    cell = table.rows[r].cells[c]; cell.text = ''
                    add_runs(cell.paragraphs[0], value, superscript)
                    for run in cell.paragraphs[0].runs:
                        run.font.size = Pt(9); run.bold = run.bold or r == 0 and lang != 'zh'
            if lang == 'zh':
                three_line(table, OxmlElement, qn)
        elif re.match(r'^\s*[-*]\s+', line):
            p = doc.add_paragraph(style='List Bullet'); add_runs(p, re.sub(r'^\s*[-*]\s+', '', line), superscript); i += 1
        else:
            # the reference list keeps its [n] labels on the line, not superscript
            p = doc.add_paragraph(); add_runs(p, line.strip(), superscript and not in_refs); i += 1
    Path(out).parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--input', required=True, type=Path)
    ap.add_argument('--output', required=True, type=Path)
    ap.add_argument('--lang', choices=['zh', 'en'], default='zh')
    ap.add_argument('--superscript-citations', action='store_true')
    a = ap.parse_args()
    print(convert(a.input, a.output, a.lang, a.superscript_citations))


if __name__ == '__main__':
    main()
