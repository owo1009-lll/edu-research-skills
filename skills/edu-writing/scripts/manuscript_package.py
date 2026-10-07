"""Assemble source-checked prose, editable tables and reproducible figures into Word."""
import argparse
import csv
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
from datetime import datetime, timezone
from collections import defaultdict

# Executing an installed tool must not create unowned caches in its read-only Skill tree.
sys.dont_write_bytecode = True
from make_figures import render, validate_spec

CITE = re.compile(r'\[@([A-Za-z0-9_-]+)\]')
CROSSREF = re.compile(r'\{\{(table|figure):([A-Za-z0-9_-]+)\}\}')


def save_json(path, data):
    Path(path).write_bytes((json.dumps(data, ensure_ascii=False, indent=2) + '\n').encode('utf-8'))


def author_label(metadata):
    authors = metadata['authors']
    if not authors:
        return metadata.get('short_title', metadata['title'].split(':')[0])
    families = [a.get('family', a.get('name', '')) for a in authors]
    if len(families) == 1:
        return families[0]
    if len(families) == 2:
        return families[0] + ' & ' + families[1]
    return families[0] + ' et al.'


def citation_labels(references):
    groups = defaultdict(list)
    for ref in references:
        m = ref['metadata']
        groups[(author_label(m), m['year'])].append(ref)
    labels = {}
    for (author, year), items in groups.items():
        for i, ref in enumerate(sorted(items, key=lambda r: r['metadata']['title'].casefold())):
            suffix = chr(97 + i) if len(items) > 1 else ''
            labels[ref['id']] = {'citation': f'({author}, {year}{suffix})', 'year': f'{year}{suffix}'}
    return labels


def bibliography(ref, label):
    m = ref['metadata']
    def initials(given):
        return ' '.join(part[0] + '.' for part in given.replace('-', ' ').split() if part)
    authors = [a.get('family', a.get('name', '')) + (', ' + initials(a['given']) if a.get('given') else '')
               for a in m['authors']]
    if len(authors) > 20:
        authors = authors[:19] + ['…', authors[-1]]
    author_text = ', '.join(authors[:-1]) + ', & ' + authors[-1] if len(authors) > 1 else ''.join(authors)
    title = m['title']
    start = f"{author_text} ({label['year']}). {title}." if authors else f"{title}. ({label['year']})."
    volume = m.get('volume', '')
    if m.get('issue'):
        volume += '(' + m['issue'] + ')'
    pages = ', ' + m['pages'].replace('-', '–') if m.get('pages') else ''
    identifier = 'https://doi.org/' + m['doi'] if m.get('doi') else m['url']
    details = volume + pages if volume else pages.lstrip(', ')
    publication = m['venue'] + (', ' + details if details else '')
    return f"{start} {publication}. {identifier}"


def blocks(manuscript):
    yield {'id': 'abstract', 'type': 'paragraph', 'text': manuscript['abstract']}
    for section in manuscript['sections']:
        yield from section['blocks']


def text_of(block):
    if block['type'] == 'paragraph':
        return block['text']
    text = block['title'] + ' ' + block.get('note', block.get('caption', ''))
    if block['type'] == 'table':
        text += ' ' + ' '.join(str(c) for c in block.get('columns', []))
        text += ' ' + ' '.join(str(c) for row in block.get('rows', []) for c in row)
    return text


def check(manuscript, verification):
    errors = []
    refs = {r['id']: r for r in verification['identity_results']}
    requested = manuscript['reference_ids']
    if len(requested) != len(set(requested)):
        errors.append('Duplicate bibliography identifiers')
    selected = [refs[k] for k in requested if k in refs]
    if len(selected) != len(requested):
        errors.append('Unknown bibliography identifier')
    seen_doi = set()
    for ref in selected:
        if ref['status'] != 'verified':
            errors.append('Unverified or mismatched reference: ' + ref['id'])
        doi = (ref.get('doi') or ref.get('lookup_url', '')).lower()
        if doi in seen_doi:
            errors.append('Duplicate DOI in bibliography')
        seen_doi.add(doi)
        if ref.get('reading', {}).get('status') not in ('full_text', 'source_excerpt', 'abstract', 'secondary_citation'):
            errors.append('Unread source: ' + ref['id'])
        if ref.get('notice_check', {}).get('status') in (None, '', 'not_checked', 'pending', 'unresolved', 'conflicting'):
            errors.append('Notice/version check missing: ' + ref['id'])
        for required in ref.get('notice_check', {}).get('required_reference_ids', []):
            if required not in requested:
                errors.append('Known correction not included: ' + required)
    all_blocks = list(blocks(manuscript))
    ids = [b['id'] for b in all_blocks]
    if len(ids) != len(set(ids)):
        errors.append('Duplicate block identifiers')
    labels, counts = {}, {'table': 0, 'figure': 0}
    for block in all_blocks:
        kind = block['type']
        if kind not in ('paragraph', 'table', 'figure'):
            errors.append('Unsupported manuscript block: ' + kind)
        if kind in counts:
            counts[kind] += 1
            labels[(kind, block['id'])] = (('表' if kind == 'table' else '图') + str(counts[kind])) if manuscript.get('language') == 'zh' else kind.title() + ' ' + str(counts[kind])
            if not block.get('title') or not block.get('note', block.get('caption')):
                errors.append('Missing display title/note: ' + block['id'])
            if kind == 'table':
                if not block.get('columns') or not block.get('rows') or any(len(r) != len(block['columns']) for r in block['rows']):
                    errors.append('Invalid editable table: ' + block['id'])
            else:
                try:
                    validate_spec(block['spec'])
                except (ValueError, KeyError) as exc:
                    errors.append('Invalid figure ' + block['id'] + ': ' + str(exc))
    cited, displayed = set(), set()
    semantic = verification.get('semantic_review', [])
    for block in all_blocks:
        text = text_of(block)
        for key in CITE.findall(text):
            cited.add(key)
            matches = [c for c in semantic if c.get('block_id') == block['id'] and c.get('reference_id') == key]
            if not matches or any(c.get('status') != 'supported' or not c.get('locator') or not c.get('support') or not c.get('reader') for c in matches):
                errors.append(f'Missing completed source-support judgment: {block["id"]}/{key}')
            if key not in requested:
                errors.append('Citation absent from bibliography: ' + key)
        for kind, key in CROSSREF.findall(text):
            if (kind, key) not in labels:
                errors.append('Unknown display reference: ' + kind + ':' + key)
            elif block['type'] == 'paragraph':
                displayed.add((kind, key))
    if set(requested) != cited:
        errors.append('Uncited bibliography entries: ' + ','.join(sorted(set(requested) - cited)))
    for key in labels:
        if key not in displayed:
            errors.append('Table/figure not cited in body: ' + str(key))
    if manuscript.get('scope') != 'full_paper':
        errors.append('Full-paper exporter requires full_paper scope; use direct output for short edits')
    if not manuscript.get('title') or not manuscript.get('keywords') or not manuscript.get('sections'):
        errors.append('Missing title/keywords/sections')
    research_type = manuscript.get('research_type')
    if research_type not in ('quantitative', 'qualitative', 'mixed', 'theoretical', 'narrative_review'):
        errors.append('Choose the actual supported research type')
    roles = {s.get('role', s['heading'].casefold()) for s in manuscript['sections']}
    needed = {'introduction', 'conclusion'}
    if research_type in ('quantitative', 'qualitative', 'mixed'):
        needed |= {'methods', 'results', 'discussion'}
    if not needed <= roles:
        errors.append('Incomplete full-paper sections: ' + ','.join(sorted(needed - roles)))
    # This verifies a recorded reader judgment, not the truth or quality of that judgment.
    return {'status': 'failed' if errors else 'passed', 'errors': errors,
            'cited_reference_ids': sorted(cited), 'display_numbers': {f'{k}:{v}': label for (k, v), label in labels.items()},
            'table_count': counts['table'], 'figure_count': counts['figure'],
            'semantic_limit': 'Source-support judgments require actual reading and AI/human review; scripts only check records and consistency'}


CJK = re.compile(r'[一-鿿]')
NUMBERED = re.compile(r'(\[\d+(?:[,–-]\d+)*\])')


def compress_numbers(nums):
    """[1, 2, 3, 5] -> '1-3,5' as GB/T 7714 writes consecutive citations."""
    nums, parts, i = sorted(set(nums)), [], 0
    while i < len(nums):
        j = i
        while j + 1 < len(nums) and nums[j + 1] == nums[j] + 1:
            j += 1
        parts.append(str(nums[i]) if j == i else (f'{nums[i]},{nums[j]}' if j == i + 1 else f'{nums[i]}-{nums[j]}'))
        i = j + 1
    return ','.join(parts)


def gbt_numbers(manuscript):
    order = []
    for block in blocks(manuscript):
        for key in CITE.findall(text_of(block)):
            if key not in order:
                order.append(key)
    return {key: i + 1 for i, key in enumerate(order)}


def gbt_author(a):
    family, given = a.get('family', a.get('name', '')), a.get('given', '')
    if CJK.search(family + given):
        return family + given
    initials = ' '.join(part[0].upper() for part in given.replace('-', ' ').replace('.', ' ').split() if part)
    return (family + ' ' + initials).strip()


def gbt_entry(ref, number):
    """One GB/T 7714-2015 reference-list entry (sequential numbering system)."""
    m = ref['metadata']
    kind = m.get('type', 'J')
    names = [gbt_author(a) for a in m.get('authors', [])]
    zh = bool(CJK.search(m['title']))
    if len(names) > 3:
        names = names[:3] + ['等' if zh else 'et al']
    authors = ', '.join(names)
    head = f'[{number}] ' + (authors + '. ' if authors else '') + f"{m['title']}[{kind}]. "
    pages = (m.get('pages') or '').replace('–', '-')
    if kind == 'J':
        vol = (m.get('volume') or '') + (f"({m['issue']})" if m.get('issue') else '')
        body = f"{m['venue']}, {m['year']}" + (f', {vol}' if vol else '') + (f': {pages}' if pages else '') + '.'
    elif kind in ('M', 'C', 'R'):
        place = m.get('place', '')
        publisher = m.get('publisher', m.get('venue', ''))
        body = (f'{place}: ' if place else '') + f"{publisher}, {m['year']}" + (f': {pages}' if pages else '') + '.'
    elif kind == 'D':
        body = (f"{m['place']}: " if m.get('place') else '') + f"{m.get('publisher', m.get('venue', ''))}, {m['year']}."
    else:  # EB/OL and other online sources
        body = (f"({m['published']})" if m.get('published') else '') + (f"[{m['accessed']}]" if m.get('accessed') else '') + f". {m.get('url', '')}"
    tail = f" DOI: {m['doi']}." if m.get('doi') and kind == 'J' else ''
    return head + body + tail


def build(manuscript, verification, output, output_format='docx', style=None):
    report = check(manuscript, verification)
    if report['errors']:
        raise ValueError('\n'.join(report['errors']))
    output = Path(output).resolve()
    output.mkdir(parents=True, exist_ok=False)
    save_json(output / 'input_manuscript.json', manuscript)
    save_json(output / 'verification.json', verification)
    selected = [r for r in verification['identity_results'] if r['id'] in manuscript['reference_ids']]
    zh = manuscript.get('language') == 'zh'
    style = style or manuscript.get('citation_style') or ('gbt7714' if zh else 'apa')
    labels = citation_labels(selected)
    numbers = gbt_numbers(manuscript) if style == 'gbt7714' else {}
    display = report['display_numbers']
    words = ({'abstract': '摘要', 'keywords': '关键词：', 'sep': '；', 'references': '参考文献', 'dot': ' ', 'abbr': ' 缩写：'} if zh else
             {'abstract': 'Abstract', 'keywords': 'Keywords: ', 'sep': '; ', 'references': 'References', 'dot': '. ', 'abbr': ' Abbreviations: '})
    def cite_group(match):
        return '[' + compress_numbers([numbers[k] for k in CITE.findall(match.group(0))]) + ']'
    def substitute(text):
        if style == 'gbt7714':
            text = re.sub(r'(?:\[@[A-Za-z0-9_-]+\])+', cite_group, text)
        else:
            text = CITE.sub(lambda m: labels[m[1]]['citation'], text)
        return CROSSREF.sub(lambda m: display[m[1] + ':' + m[2]], text)
    def add_par(document, text):
        """Add a paragraph; GB/T 7714 citation numbers are superscript."""
        p = document.add_paragraph()
        for piece in (NUMBERED.split(text) if style == 'gbt7714' else [text]):
            if piece:
                run = p.add_run(piece)
                run.font.superscript = style == 'gbt7714' and bool(NUMBERED.fullmatch(piece))
        return p
    def display_note(block, field):
        note = substitute(block[field])
        abbreviations = block.get('abbreviations', {})
        if abbreviations:
            note += words['abbr'] + words['sep'].join(k + ', ' + v for k, v in abbreviations.items()) + ('。' if zh else '.')
        return note
    md = ['# ' + manuscript['title'], '']
    if manuscript.get('disclosure'):
        md += ['*' + substitute(manuscript['disclosure']) + '*', '']
    md += ['## ' + words['abstract'], '', substitute(manuscript['abstract']), '', words['keywords'] + words['sep'].join(manuscript['keywords']), '']
    if output_format == 'docx':
        from docx import Document
        from docx.shared import Inches, Pt, Mm
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        document = Document()
        document.core_properties.author = manuscript.get('author', '')
        document.core_properties.title = manuscript['title']
        section = document.sections[0]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin = section.bottom_margin = Inches(.85)
        section.left_margin = section.right_margin = Inches(.85)
        normal = document.styles['Normal']
        normal.font.name, normal.font.size = 'Times New Roman', Pt(11)
        normal.paragraph_format.line_spacing = 1.15
        normal.paragraph_format.space_after = Pt(6)
        for heading in ['Heading 1', 'Heading 2', 'Title']:
            document.styles[heading].font.name = 'Arial'
        if zh:  # East Asian fonts: SimSun body, SimHei headings
            for name, font in [('Normal', '宋体'), ('Heading 1', '黑体'), ('Heading 2', '黑体'), ('Title', '黑体')]:
                document.styles[name].element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), font)
        document.add_heading(manuscript['title'], 0)
        if manuscript.get('disclosure'):
            p = document.add_paragraph(substitute(manuscript['disclosure']))
            for run in p.runs:
                run.italic = True; run.font.size = Pt(9)
        document.add_heading(words['abstract'], 1)
        add_par(document, substitute(manuscript['abstract']))
        document.add_paragraph(words['keywords'] + words['sep'].join(manuscript['keywords']))
        footer = document.sections[0].footer.paragraphs[0]
        footer.alignment = 2
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    figure_receipts = []
    for section in manuscript['sections']:
        md += ['## ' + section['heading'], '']
        if output_format == 'docx':
            document.add_heading(section['heading'], 1)
        for block in section['blocks']:
            kind = block['type']
            if kind == 'paragraph':
                text = substitute(block['text']); md += [text, '']
                if output_format == 'docx': add_par(document, text)
            elif kind == 'table':
                number = display['table:' + block['id']]
                title = number + words['dot'] + substitute(block['title'])
                note = display_note(block, 'note')
                table_dir = output / 'tables'; table_dir.mkdir(exist_ok=True)
                with (table_dir / (block['id'] + '.csv')).open('w', encoding='utf-8-sig', newline='') as stream:
                    writer = csv.writer(stream); writer.writerow([substitute(str(c)) for c in block['columns']])
                    writer.writerows([[substitute(str(c)) for c in row] for row in block['rows']])
                md += [title, '', '| ' + ' | '.join(substitute(str(c)) for c in block['columns']) + ' |',
                       '| ' + ' | '.join(['---'] * len(block['columns'])) + ' |']
                md += ['| ' + ' | '.join(substitute(str(c)) for c in row) + ' |' for row in block['rows']]
                md += ['', note, '']
                if output_format == 'docx':
                    p = document.add_paragraph(title); p.paragraph_format.keep_with_next = True
                    p.runs[0].bold = True
                    table = document.add_table(rows=1, cols=len(block['columns'])); table.style = 'Table Grid' if zh else 'Light Shading Accent 1'
                    for cell, value in zip(table.rows[0].cells, block['columns']): cell.text = substitute(str(value))
                    repeat = OxmlElement('w:tblHeader'); table.rows[0]._tr.get_or_add_trPr().append(repeat)
                    for values in block['rows']:
                        row = table.add_row()
                        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                        for cell, value in zip(row.cells, values): cell.text = substitute(str(value))
                    # Keep modest short tables and their note together. Large tables retain
                    # repeated headers and unbroken rows without forcing an oversized page.
                    compact = len(block['rows']) <= 12 and sum(len(str(c)) for row in block['rows'] for c in row) <= 1800
                    for row in table.rows:
                        for cell in row.cells:
                            for p in cell.paragraphs:
                                p.paragraph_format.keep_with_next = compact
                                for run in p.runs: run.font.size = Pt(9)
                    if zh:  # three-line table: thick top and bottom rules, thin rule under the header row
                        last = len(table.rows) - 1
                        for r, row in enumerate(table.rows):
                            for cell in row.cells:
                                edges = {}
                                if r == 0:
                                    edges.update(top=12, bottom=6)
                                if r == last:
                                    edges['bottom'] = 12
                                box = OxmlElement('w:tcBorders')
                                for edge in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                                    e = OxmlElement('w:' + edge)
                                    e.set(qn('w:val'), 'single' if edge in edges else 'nil')
                                    if edge in edges:
                                        e.set(qn('w:sz'), str(edges[edge])); e.set(qn('w:color'), '000000')
                                    box.append(e)
                                cell._tc.get_or_add_tcPr().append(box)
                    document.add_paragraph(note)
            else:
                number = display['figure:' + block['id']]
                destination = output / 'figures' / block['id']
                receipt = render(block['spec'], destination)
                receipt.update(id=block['id'], number=number); figure_receipts.append(receipt)
                caption = number + words['dot'] + substitute(block['title']).rstrip('. ') + ('。' if zh else '. ') + display_note(block, 'caption')
                md += [f'![{number}](figures/{block["id"]}/figure.png)', '', caption, '']
                if output_format == 'docx':
                    p = document.add_paragraph(); p.paragraph_format.keep_with_next = True
                    p.add_run().add_picture(str(destination / 'figure.png'), width=Inches(6.1))
                    p = document.add_paragraph(caption)
                    for run in p.runs: run.font.size = Pt(9)
    md += ['## ' + words['references'], '']
    if output_format == 'docx': document.add_heading(words['references'], 1)
    ordered = (sorted(selected, key=lambda r: numbers[r['id']]) if style == 'gbt7714' else
               sorted(selected, key=lambda r: (author_label(r['metadata']).casefold(), labels[r['id']]['year'], r['metadata']['title'])))
    for ref in ordered:
        line = gbt_entry(ref, numbers[ref['id']]) if style == 'gbt7714' else bibliography(ref, labels[ref['id']]); md += [line, '']
        if output_format == 'docx':
            p = document.add_paragraph(line)
            p.paragraph_format.left_indent = Inches(.25); p.paragraph_format.first_line_indent = Inches(-.25)
    (output / 'manuscript.md').write_bytes(('\n'.join(md) + '\n').encode('utf-8'))
    if output_format == 'docx': document.save(output / 'manuscript.docx')
    # Editable independent Word tables are also delivered, not rasterized screenshots.
    if output_format == 'docx' and report['table_count']:
        table_doc = Document()
        for section in manuscript['sections']:
            for block in section['blocks']:
                if block['type'] != 'table': continue
                table_doc.add_heading(display['table:' + block['id']] + '. ' + substitute(block['title']), 1)
                table = table_doc.add_table(rows=1, cols=len(block['columns'])); table.style = 'Table Grid'
                for cell, value in zip(table.rows[0].cells, block['columns']): cell.text = substitute(str(value))
                for row in block['rows']:
                    for cell, value in zip(table.add_row().cells, row): cell.text = substitute(str(value))
                table_doc.add_paragraph(display_note(block, 'note'))
        table_doc.save(output / 'editable_tables.docx')
    for script in ['make_figures.py', 'manuscript_package.py', 'verify_references.py']:
        folder = output / 'code'; folder.mkdir(exist_ok=True)
        shutil.copy2(Path(__file__).with_name(script), folder / script)
    (output / 'PENDING.md').write_bytes(('# Pending and scope notes\n\n' + '\n'.join('- ' + s for s in manuscript.get('pending', [])) + '\n').encode('utf-8'))
    with (output / 'citation_check.csv').open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream); writer.writerow(['id', 'doi', 'identity', 'reading', 'version', 'notice_check', 'lookup_url', 'claim_support_rows'])
        for ref in selected:
            writer.writerow([ref['id'], ref['doi'], ref['status'], ref['reading']['status'], ref['reading'].get('version', ''),
                             ref['notice_check']['status'], ref['lookup_url'],
                             sum(c.get('reference_id') == ref['id'] for c in verification['semantic_review'])])
    report.update(created_utc=datetime.now(timezone.utc).isoformat(), format=output_format,
                  extra_analysis=False, figure_receipts=figure_receipts,
                  artifacts={p.relative_to(output).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
                             for p in output.rglob('*') if p.is_file()})
    save_json(output / 'package_receipt.json', report)
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--manuscript', type=Path, required=True)
    parser.add_argument('--verification', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--format', choices=['docx', 'markdown'], default='docx')
    parser.add_argument('--style', choices=['apa', 'gbt7714'], help='default: gbt7714 when the manuscript language is zh, else apa')
    args = parser.parse_args()
    receipt = build(json.loads(args.manuscript.read_text(encoding='utf-8')),
                    json.loads(args.verification.read_text(encoding='utf-8')), args.output, args.format, args.style)
    print(json.dumps({k: receipt[k] for k in ['status', 'format', 'table_count', 'figure_count', 'cited_reference_ids']}, indent=2))
