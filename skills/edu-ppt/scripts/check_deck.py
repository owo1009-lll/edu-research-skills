"""Check plan/object consistency and preview freshness, separately from source facts."""
from pathlib import Path
from collections import Counter
import argparse
import hashlib
import json
from pptx import Presentation
from plan import load_plan
from fidelity import check_objects


def all_text(shape):
    if shape.has_text_frame:
        return shape.text
    if shape.has_table:
        return '\n'.join(cell.text for row in shape.table.rows for cell in row.cells)
    return ''


def check(source, plan_path, render_dir=None):
    plan = load_plan(plan_path)
    prs = Presentation(source)
    errors, warnings, inventory = [], [], []
    issues = check_objects(prs, plan, plan_path, source)
    errors.extend(f"slide {x['slide']}, object {x['object']}, {x['field']}: expected {x['expected']!r}; actual {x['actual']!r}" for x in issues)
    if len(prs.slides) != len(plan['slides']):
        errors.append('Slide count differs from plan')
    for i, (slide, page) in enumerate(zip(prs.slides, plan['slides']), 1):
        counts = Counter()
        shapes = []
        for shape in slide.shapes:
            text = all_text(shape)
            if shape.has_table:
                counts['native_tables'] += 1
            elif shape.has_chart:
                counts['native_charts'] += 1
            elif int(shape.shape_type) == 13:
                counts['images'] += 1
            elif shape.has_text_frame and text:
                counts['editable_text'] += 1
            else:
                counts['shapes_or_lines'] += 1
            if shape.left < -9144 or shape.top < -9144 or shape.left + shape.width > prs.slide_width + 9144 or shape.top + shape.height > prs.slide_height + 9144:
                errors.append(f'slide {i}: outside canvas: {shape.name}')
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        if run.font.size and run.font.size.pt < 10:
                            warnings.append(f'slide {i}: tiny text in {shape.name}')
            # Intentional dots/lines within interval plots are excluded from overlap checks.
            if shape.name in ('edu:title', 'edu:message', 'edu:body', 'edu:source', 'edu:table', 'edu:chart', 'edu:source-image', 'edu:quote', 'edu:caption', 'edu:truncation-note'):
                shapes.append(shape)
        for j, a in enumerate(shapes):
            for b in shapes[j + 1:]:
                x = min(a.left + a.width, b.left + b.width) - max(a.left, b.left)
                y = min(a.top + a.height, b.top + b.height) - max(a.top, b.top)
                if x > 9144 and y > 9144:
                    errors.append(f'slide {i}: overlapping content: {a.name} / {b.name}')
        if counts['editable_text'] == 0:
            errors.append(f'slide {i}: no editable text')
        inventory.append({'slide': i, **dict(counts)})
    render = {'status': 'not_checked', 'visual_review': 'not_performed'}
    if render_dir:
        path = Path(render_dir) / 'render.json'
        if not path.exists():
            errors.append('No render record')
        else:
            render = json.loads(path.read_text(encoding='utf-8'))
            if render.get('status') != 'rendered':
                errors.append('Rendering did not complete')
            if render.get('pptx_sha256') != hashlib.sha256(Path(source).read_bytes()).hexdigest():
                errors.append('Previews are stale: PPTX changed after rendering')
            if len(render.get('images', {})) != len(prs.slides):
                errors.append('Preview page count differs from PPTX')
            for name, expected in render.get('images', {}).items():
                image = Path(render_dir) / name
                if not image.exists() or hashlib.sha256(image.read_bytes()).hexdigest() != expected:
                    errors.append('Preview image missing or changed: ' + name)
    return {'status': 'passed' if not errors else 'failed', 'errors': errors, 'issues': issues, 'warnings': warnings,
            'inventory': inventory, 'render_status': render.get('status'),
            'plan_consistency': 'failed' if issues or len(prs.slides) != len(plan['slides']) else 'passed',
            'source_fact_check': 'not_performed',
            'visual_review': render.get('visual_review', 'not_performed'),
            'scope': 'native object consistency with supplied verified plan, geometry and preview freshness; original-source correctness and visual quality require separate review'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--pptx', type=Path, required=True)
    parser.add_argument('--plan', type=Path, required=True)
    parser.add_argument('--render', type=Path)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    report = check(args.pptx, args.plan, args.render)
    args.out.write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    if report['status'] != 'passed':
        raise SystemExit(1)
