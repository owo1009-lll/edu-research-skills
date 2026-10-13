"""Inspect dimensions, masters, layouts and theme choices without changing a template."""
from pathlib import Path
import argparse
import json
from zipfile import ZipFile
from xml.etree import ElementTree as ET
from pptx import Presentation


def inspect(path):
    prs = Presentation(path)
    result = {'width_inches': prs.slide_width / 914400, 'height_inches': prs.slide_height / 914400,
              'sample_slides': len(prs.slides), 'layouts': [], 'themes': []}
    for master in prs.slide_masters:
        for layout in master.slide_layouts:
            result['layouts'].append({'name': layout.name, 'placeholders': [
                {'name': s.name, 'box_inches': [v / 914400 for v in (s.left, s.top, s.width, s.height)]}
                for s in layout.placeholders]})
    ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main'}
    with ZipFile(path) as archive:
        for name in archive.namelist():
            if name.startswith('ppt/theme/theme') and name.endswith('.xml'):
                root = ET.fromstring(archive.read(name))
                colors = root.find('.//a:clrScheme', ns)
                result['themes'].append({'part': name,
                    'colors': {e.tag.rsplit('}', 1)[-1]: list(e)[0].get('val', list(e)[0].get('lastClr')) for e in colors} if colors is not None else {},
                    'fonts': [e.attrib for e in root.findall('.//a:majorFont/*', ns) + root.findall('.//a:minorFont/*', ns)]})
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('template', type=Path)
    parser.add_argument('--out', type=Path)
    args = parser.parse_args()
    data = json.dumps(inspect(args.template), ensure_ascii=False, indent=2)
    if args.out:
        args.out.write_text(data + '\n', encoding='utf-8')
    else:
        print(data)
