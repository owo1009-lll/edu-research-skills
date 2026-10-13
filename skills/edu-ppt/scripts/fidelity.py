"""Read native objects by stable ID and compare with the supplied plan."""
from pathlib import Path
import hashlib
import json
from pptx.enum.chart import XL_CHART_TYPE
from objects import identity
from plan import reference_text, interval_text, display_number, chart_range_error
from chart_labels import check_labels


def check_objects(prs, plan, plan_path, source):
    issues = []
    manifest_path = Path(source).parent / 'build.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8')) if manifest_path.exists() else {}
    records = {r['id']: r for r in manifest.get('objects', [])}
    for i, (slide, page) in enumerate(zip(prs.slides, plan['slides']), 1):
        mapped, required = {}, set()

        def fail(key, field, expected, actual):
            oid = key if key.startswith(page['id'] + '/') else page['id'] + '/' + key
            issues.append(dict(slide=i, object=oid, field=field, expected=expected, actual=actual))

        def equal(key, field, expected, actual):
            if expected != actual:
                fail(key, field, expected, actual)

        for shape in slide.shapes:
            oid = identity(shape).get('id')
            if oid:
                if oid in mapped:
                    fail(oid, 'identity', 'one object', 'duplicate ID')
                mapped[oid] = shape
            elif shape.has_table or shape.has_chart or int(shape.shape_type) == 13:
                fail(shape.name, 'identity', 'stable plan object ID', 'untracked native visual')

        def get(key, kind=None):
            oid = page['id'] + '/' + key
            required.add(oid)
            shape = mapped.get(oid)
            if shape is None:
                fail(key, 'presence', kind or 'planned object', 'missing')
            elif kind and not {'table': shape.has_table, 'chart': shape.has_chart,
                              'image': int(shape.shape_type) == 13, 'text': shape.has_text_frame,
                              'connector': bool(shape._element.xpath('self::p:cxnSp'))}[kind]:
                fail(key, 'object type', kind, str(shape.shape_type))
                return None
            return shape

        def text_at(key, expected):
            shape = get(key, 'text')
            if shape is not None:
                equal(key, 'planned text missing or changed', str(expected), shape.text)
            return shape

        def metadata(key, shape, **expected):
            if shape is not None:
                actual = identity(shape)
                for field, value in expected.items():
                    equal(key, 'record.' + field, value, actual.get(field))

        def position(key, label, expected, actual):
            if abs(expected - actual) > 1829:  # 0.002 inch, EMU rounding tolerance
                fail(key, label + ' (EMU)', round(expected), actual)

        text_at('title', page['title'])
        text_at('message', page['message'])
        for j, body in enumerate(page.get('body', [])):
            text_at(f'body/{j}', body)
        footer = reference_text(plan, page)
        text_at('source', footer)
        text_at('number', str(i))
        notes = slide.notes_slide.notes_text_frame
        equal('notes', 'speaker notes differ from plan',
              page['notes'] + '\n\n来源：' + footer + f"\n预计讲解：{page['seconds']}秒",
              notes.text if notes is not None else None)
        data = page.get('visual', {})
        kind = data.get('type')
        if data.get('caption'):
            text_at('visual/caption', data['caption'])
        if kind == 'table':
            shape = get('visual/table', 'table')
            if shape is not None:
                expected, table = [data['columns'], *data['rows']], shape.table
                equal('visual/table', 'rows/columns', [len(expected), len(expected[0])], [len(table.rows), len(table.columns)])
                for r, row in enumerate(expected[:len(table.rows)]):
                    for c, value in enumerate(row[:len(table.columns)]):
                        equal('visual/table', f'cell[{r},{c}]', str(value), table.cell(r, c).text)
        elif kind == 'chart':
            key = 'visual/chart'
            shape = get(key, 'chart')
            if data.get('truncation_note'):
                text_at('visual/truncation-note', '坐标说明：' + data['truncation_note'])
            if shape is not None:
                chart = shape.chart
                types = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'line': XL_CHART_TYPE.LINE_MARKERS}
                equal(key, 'chart kind', types[data.get('kind', 'column')], chart.chart_type)
                equal(key, 'categories/order', data['categories'], [c.label for c in chart.plots[0].categories])
                equal(key, 'series names/order', [s['name'] for s in data['series']], [s.name for s in chart.series])
                equal(key, 'values/order', [s['values'] for s in data['series']], [list(s.values) for s in chart.series])
                axis = chart.value_axis
                equal(key, 'unit', data['unit'], axis.axis_title.text_frame.text if axis.has_title else None)
                for field, prop in [('minimum', 'minimum_scale'), ('maximum', 'maximum_scale'), ('major_unit', 'major_unit')]:
                    equal(key, 'axis.' + field, data.get(field), getattr(axis, prop))
                equal(key, 'axis.scale', [], [x.get('val') for x in chart._chartSpace.xpath('.//c:logBase')])
                equal(key, 'category direction', False, chart.category_axis.reverse_order)
                equal(key, 'value direction', False, axis.reverse_order)
                equal(key, 'axis display units', [], [x.tag for x in axis._element.xpath('./c:dispUnits')])
                equal(key, 'axis number format', 'General', axis.tick_labels.number_format)
                actual_range = dict(data)
                for field, prop in [('minimum', 'minimum_scale'), ('maximum', 'maximum_scale')]:
                    value = getattr(axis, prop)
                    if value is None:
                        actual_range.pop(field, None)
                    else:
                        actual_range[field] = value
                problem = chart_range_error(actual_range)
                if problem and not (data.get('allow_truncation') is True and data.get('truncation_note')):
                    fail(key, 'axis range', 'data visible without silent truncation', problem)
                equal(key, 'legend', len(data['series']) > 1, chart.has_legend)
                plot = chart.plots[0]
                equal(key, 'data labels', data.get('labels', True), plot.has_data_labels)
                if plot.has_data_labels:
                    equal(key, 'number format', data.get('number_format', 'General'), plot.data_labels.number_format)
                for scope, field, expected, actual in check_labels(chart, data):
                    fail(key + '/' + scope, field, expected, actual)
                actual_colors = []
                for n, series in enumerate(chart.series):
                    try:
                        rgb = series.format.line.color.rgb if data.get('kind') == 'line' else series.format.fill.fore_color.rgb
                        actual_colors.append(str(rgb))
                    except (AttributeError, TypeError):
                        actual_colors.append(None)
                    for override in series._element.xpath('./c:dPt/c:spPr/a:solidFill/a:srgbClr'):
                        equal(key, f'series[{n}] point color', actual_colors[-1], override.get('val'))
                expected_colors = identity(shape).get('series_colors', [])
                equal(key, 'series colors', [c.upper() for c in expected_colors], actual_colors)
                if None in actual_colors or len(set(actual_colors)) != len(actual_colors):
                    fail(key, 'series distinction', 'distinct explicit colors', actual_colors)
        elif kind == 'interval':
            axis = get('visual/axis-line', 'connector')
            metadata('visual/axis-line', axis, range=data['range'], scale='linear', unit=data['unit'])
            left, right = data['range']
            reference = data.get('reference', 0)
            text_at('visual/axis-label', f"{display_number(left)} ← {data['unit']} → {display_number(right)}")
            text_at('visual/reference-label', '参考值 = ' + display_number(reference))
            ref = get('visual/reference', 'connector')
            metadata('visual/reference', ref, value=reference)
            at = lambda v: axis.left + (v - left) / (right - left) * axis.width
            if axis is not None and ref is not None:
                position('visual/reference', 'reference x', at(reference), ref.left)
                position('visual/reference', 'vertical reference width', 0, ref.width)
            for j, row in enumerate(data['rows']):
                key = f'visual/row/{j}'
                label = text_at(key + '/label', row['label'])
                display = text_at(key + '/display', interval_text(data, row))
                bar = get(key + '/interval', 'connector')
                point = get(key + '/estimate')
                metadata(key + '/interval', bar, low=row['low'], high=row['high'])
                metadata(key + '/estimate', point, value=row['estimate'])
                if axis is not None and bar is not None and point is not None:
                    position(key, 'low x', at(row['low']), bar.left)
                    position(key, 'high x', at(row['high']), bar.left + bar.width)
                    position(key, 'estimate x', at(row['estimate']), point.left + point.width / 2)
                    position(key, 'estimate y', bar.top, point.top + point.height / 2)
                if label is not None and display is not None and bar is not None and point is not None:
                    # Relative to each row's text, independent of slide origin,
                    # row order and theme. Cap by nearest other row spacing so
                    # tall text boxes cannot make a neighbouring row acceptable.
                    center = lambda s: s.top + s.height / 2
                    others = [center(s) for oid, s in mapped.items()
                              if '/visual/row/' in oid and oid.endswith('/label') and s is not label]
                    tolerance = max(1829, min(label.height, display.height) * .18)
                    if others:
                        tolerance = min(tolerance, min(abs(center(label) - y) for y in others) * .25)
                    for part, actual_y in [('display', center(display)), ('interval', bar.top + bar.height / 2), ('estimate', center(point))]:
                        if abs(center(label) - actual_y) > tolerance:
                            fail(key + '/' + part, 'row correspondence y (EMU)',
                                 {'label_center': round(center(label)), 'tolerance': round(tolerance)}, round(actual_y))
        elif kind == 'quote':
            text_at('visual/quote', '“' + data['text'] + '”')
            text_at('visual/attribution', data['attribution'])
        elif kind == 'process':
            nodes = [text_at(f'visual/node/{j}', value) for j, value in enumerate(data['steps'])]
            for j, (first, second) in enumerate(zip(nodes, nodes[1:])):
                key = f'visual/edge/{j}'
                edge = get(key, 'connector')
                metadata(key, edge, start=f'visual/node/{j}', end=f'visual/node/{j + 1}')
                if edge is not None:
                    for tag, node in [('stCxn', first), ('endCxn', second)]:
                        actual = edge._element.xpath('.//a:' + tag)
                        if node is not None:
                            equal(key, tag + ' relation', str(node.shape_id), actual[0].get('id') if actual else None)
                    equal(key, 'direction', ['triangle'], [e.get('type') for e in edge._element.xpath('.//a:tailEnd')])
                    equal(key, 'reverse arrow', [], [e.get('type') for e in edge._element.xpath('.//a:headEnd') if e.get('type') != 'none'])
        elif kind == 'image':
            key = 'visual/image'
            shape = get(key, 'image')
            path = (Path(plan_path).resolve().parent / data['path']).resolve()
            if not path.is_file():
                fail(key, 'source asset', data['path'], 'unavailable; cannot verify bytes')
            elif shape is not None:
                expected = hashlib.sha256(path.read_bytes()).hexdigest()
                actual = hashlib.sha256(shape.image.blob).hexdigest()
                equal(key, 'embedded image SHA256', expected, actual)
                metadata(key, shape, asset_path=data['path'], asset_sha256=expected)
                record = records.get(page['id'] + '/' + key)
                if record:
                    equal(key, 'build asset SHA256', expected, record.get('asset_sha256'))
                else:
                    fail(key, 'generation record', 'matching object in adjacent build.json', 'missing')
        for extra in set(mapped) - required:
            fail(extra, 'identity', 'object declared in plan', 'unexpected object')
    return issues
