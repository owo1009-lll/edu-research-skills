"""Extract local PDF/Word/text or an accessible HTTP article with stable locators.

PDF text and embedded images are extraction aids, not OCR or a reading claim.
"""
from pathlib import Path
from html.parser import HTMLParser
from urllib.parse import urlparse
import argparse
import hashlib
import json
import re


class ArticleText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.skip = 0
        self.parts = []

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'noscript'):
            self.skip += 1
        if tag in ('p', 'div', 'section', 'h1', 'h2', 'h3', 'h4', 'li', 'tr', 'br'):
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'noscript'):
            self.skip = max(0, self.skip - 1)
        if tag in ('p', 'div', 'section', 'h1', 'h2', 'h3', 'h4', 'li', 'tr'):
            self.parts.append('\n')

    def handle_data(self, data):
        if not self.skip:
            self.parts.append(data)


def extract(source, out):
    out = Path(out).resolve()
    if out.exists() and any(out.iterdir()):
        raise ValueError('Use an empty extraction directory')
    out.mkdir(parents=True, exist_ok=True)
    origin = str(source)
    if re.match(r'^https?://', origin):
        import requests
        response = requests.get(origin, timeout=45, headers={'User-Agent': 'edu-ppt/0.11 academic source reading'})
        response.raise_for_status()
        data = response.content
        mime = response.headers.get('Content-Type', '')
        suffix = '.pdf' if data.startswith(b'%PDF') else '.html'
        path = out / ('download' + suffix)
        path.write_bytes(data)
        retrieval = {'requested_url': origin, 'resolved_url': response.url, 'content_type': mime}
    else:
        path = Path(source).resolve()
        data = path.read_bytes()
        suffix = path.suffix.lower()
        retrieval = {'local_file': str(path)}
    records, warnings, images = [], [], []
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        for i, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ''
            if len(text.strip()) < 50:
                warnings.append(f'p.{i}: little extracted text; inspect page image/OCR manually')
            records.append({'location': f'p.{i}', 'text': text})
            for j, picture in enumerate(page.images, 1):
                ext = Path(picture.name).suffix or '.png'
                name = f'images/page-{i:03d}-image-{j:02d}{ext}'
                target = out / name
                target.parent.mkdir(exist_ok=True)
                target.write_bytes(picture.data)
                images.append({'path': name, 'location': f'p.{i}, embedded image {j}', 'note': 'Verify complete figure, caption and legend against original page'})
    elif suffix == '.docx':
        from docx import Document
        document = Document(path)
        for i, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip():
                records.append({'location': f'P{i:04d}', 'text': paragraph.text})
        for i, table in enumerate(document.tables, 1):
            rows = [' | '.join(cell.text.replace('\n', ' / ') for cell in row.cells) for row in table.rows]
            records.append({'location': f'Table{i}', 'text': '\n'.join(rows)})
        for rel in document.part.rels.values():
            if rel.reltype.endswith('/image') and not rel.is_external:
                part = rel.target_part
                name = 'images/' + Path(str(part.partname)).name
                target = out / name
                target.parent.mkdir(exist_ok=True)
                target.write_bytes(part.blob)
                images.append({'path': name, 'location': str(part.partname), 'note': 'Locate caption in source Word before use'})
        warnings.append('Word locators are paragraph/table indices, not printed page numbers; inspect floating objects in original')
    elif suffix in ('.md', '.txt', '.html', '.htm'):
        text = data.decode('utf-8-sig')
        if suffix in ('.html', '.htm'):
            parser = ArticleText(); parser.feed(text)
            text = '\n'.join(line.strip() for line in ''.join(parser.parts).splitlines() if line.strip())
            warnings.append('HTML extraction includes navigation; verify article boundaries and full-text access')
        records = [{'location': f'L{i:04d}', 'text': line} for i, line in enumerate(text.splitlines(), 1) if line.strip()]
    else:
        raise ValueError('Supported sources: PDF, DOCX, Markdown, UTF-8 text, HTML or HTTP article')
    result = {'origin': origin, 'sha256': hashlib.sha256(data).hexdigest(), 'retrieval': retrieval,
              'records': records, 'images': images, 'warnings': warnings,
              'reading_status': 'extracted_not_yet_verified'}
    (out / 'source.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (out / 'source.md').write_text('\n\n'.join(f"[{r['location']}] {r['text']}" for r in records) + '\n', encoding='utf-8')
    return {'records': len(records), 'images': len(images), 'warnings': warnings}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source', required=True)
    parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args()
    print(json.dumps(extract(args.source, args.out), ensure_ascii=False, indent=2))
