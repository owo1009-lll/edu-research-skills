"""Render a verified host-authored plan into editable PPTX or content and notes.

This renderer never rewrites content, fetches sources or calls a model.
"""
from pathlib import Path
import argparse
import json
import math
import re
import sys
from plan import load_plan, export_text, reference_text, digest, interval_text, display_number
from objects import identify, identity, SERIES_COLORS
from chart_labels import configure_labels


DEFAULT = {'font': 'Microsoft YaHei', 'accent': '176B76', 'ink': '172D3A',
           'muted': '516472', 'light': 'EAF3F4', 'background': 'FFFFFF',
           'title_size': 28, 'body_size': 22, 'footer_size': 11,
           'canvas': [0.055, 0.235, 0.89, 0.64]}


def units(text):
    return sum(1 if ord(c) > 255 else .54 for c in str(text))


def make_pptx(plan, destination, base, template=None, theme=None):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
    from pptx.chart.data import CategoryChartData
    from pptx.oxml.xmlchemy import OxmlElement

    settings = dict(DEFAULT)
    if theme:
        settings.update(theme)
    if template:
        prs = Presentation(str(template))
        # Remove sample slides only; keep masters, layouts, theme and their logos.
        for item in list(prs.slides._sldIdLst):
            prs.part.drop_rel(item.rId)
            prs.slides._sldIdLst.remove(item)
    else:
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333333), Inches(7.5)
    width, height = prs.slide_width / 914400, prs.slide_height / 914400
    layouts = [layout for master in prs.slide_masters for layout in master.slide_layouts]
    if theme and 'layout' in theme:
        matches = [x for x in layouts if x.name == theme['layout']]
        if not matches:
            raise ValueError('Template layout not found: ' + theme['layout'])
        layout = matches[0]
    else:
        layout = min(layouts, key=lambda x: len(x.placeholders))
    color = lambda key: RGBColor.from_string(settings.get(key, key))
    active_id = ''

    def tag(shape, key, **record):
        return identify(shape, active_id + '/' + key, **record)

    def fit(text, w, h, size):
        usable_w, usable_h = w * 72 - 10, h * 72 - 6
        lines = sum(max(1, math.ceil(units(line) * size / max(1, usable_w))) for line in str(text).split('\n'))
        if lines * size * 1.22 > usable_h + 2:
            raise ValueError(f'{active_id}: text too dense at {size}pt; shorten/split this page: {str(text)[:55]}')

    def format_frame(frame, text, size, ink='ink', bold=False):
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = Inches(.035)
        frame.margin_top = frame.margin_bottom = Inches(.025)
        for i, line in enumerate(str(text).split('\n')):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.space_after = Pt(0)
            paragraph.line_spacing = 1.12
            run = paragraph.add_run()
            run.text = line
            run.font.name = settings['font']
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color(ink)
            props = run._r.get_or_add_rPr()
            for tag in ('a:ea', 'a:cs'):
                elem = OxmlElement(tag)
                elem.set('typeface', settings['font'])
                props.append(elem)

    def textbox(slide, text, x, y, w, h, size=None, ink='ink', bold=False, name='text', key=None):
        size = size or settings['body_size']
        fit(text, w, h, size)
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = 'edu:' + name
        format_frame(shape.text_frame, text, size, ink, bold)
        return tag(shape, key or name)

    def line(slide, x1, y1, x2, y2, ink='accent', thickness=2, key=None, **record):
        shape = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        shape.name = 'edu:line'
        shape.line.color.rgb = color(ink)
        shape.line.width = Pt(thickness)
        return tag(shape, key, **record) if key else shape

    def visual(slide, data, x, y, w, h):
        kind = data['type']
        caption = data.get('caption', '')
        note = data.get('truncation_note', '')
        ch = (.48 if caption else 0) + (.65 if note else 0)
        usable = h - ch
        if kind == 'table':
            rows = [data['columns'], *data['rows']]
            shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y), Inches(w), Inches(usable))
            shape.name = 'edu:table'
            tag(shape, 'visual/table')
            table = shape.table
            weights = data.get('widths', [1] * len(rows[0]))
            if len(weights) != len(rows[0]) or any(v <= 0 for v in weights):
                raise ValueError(active_id + ': invalid column widths')
            for j, weight in enumerate(weights):
                table.columns[j].width = Inches(w * weight / sum(weights))
            size = data.get('font_size', 19)
            if size < 18:
                raise ValueError(active_id + ': table text below 18pt; split the table')
            for i, row in enumerate(rows):
                for j, value in enumerate(row):
                    cell = table.cell(i, j)
                    fit(value, table.columns[j].width / 914400 - .08, usable / len(rows) - .08, size)
                    cell.margin_left = cell.margin_right = Inches(.07)
                    cell.margin_top = cell.margin_bottom = Inches(.04)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = color('accent' if i == 0 else ('light' if i % 2 else 'background'))
                    format_frame(cell.text_frame, value, size, 'FFFFFF' if i == 0 else 'ink', i == 0)
        elif kind == 'chart':
            chart_data = CategoryChartData()
            chart_data.categories = data['categories']
            for series in data['series']:
                chart_data.add_series(series['name'], series['values'])
            types = {'column': XL_CHART_TYPE.COLUMN_CLUSTERED, 'bar': XL_CHART_TYPE.BAR_CLUSTERED, 'line': XL_CHART_TYPE.LINE_MARKERS}
            shape = slide.shapes.add_chart(types[data.get('kind', 'column')], Inches(x), Inches(y), Inches(w), Inches(usable), chart_data)
            shape.name = 'edu:chart'
            palette = settings.get('series_colors', SERIES_COLORS)
            colors = palette[:len(data['series'])]
            if len(colors) != len(data['series']) or any(not isinstance(c, str) or not re.fullmatch('[0-9A-Fa-f]{6}', c) for c in colors) or len({c.upper() for c in colors}) != len(colors):
                raise ValueError(active_id + ': supply distinct series_colors for every series')
            tag(shape, 'visual/chart', series_colors=colors)
            chart = shape.chart
            chart.font.name, chart.font.size = settings['font'], Pt(17)
            chart.has_legend = len(data['series']) > 1
            if chart.has_legend:
                chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                chart.legend.include_in_layout = False
            chart.value_axis.has_title = True
            chart.value_axis.axis_title.text_frame.text = data['unit']
            if 'minimum' in data:
                chart.value_axis.minimum_scale = data['minimum']
            if 'maximum' in data:
                chart.value_axis.maximum_scale = data['maximum']
            if 'major_unit' in data:
                chart.value_axis.major_unit = data['major_unit']
            chart.plots[0].has_data_labels = data.get('labels', True)
            if chart.plots[0].has_data_labels:
                chart.plots[0].data_labels.font.size = Pt(16)
                chart.plots[0].data_labels.number_format = data.get('number_format', 'General')
                chart.plots[0].data_labels.position = XL_LABEL_POSITION.ABOVE if data.get('kind') == 'line' else XL_LABEL_POSITION.OUTSIDE_END
            configure_labels(chart, data)
            for i, series in enumerate(chart.series):
                rgb = RGBColor.from_string(colors[i])
                if data.get('kind', 'column') != 'line':
                    series.format.fill.solid()
                    series.format.fill.fore_color.rgb = rgb
                else:
                    from pptx.enum.chart import XL_MARKER_STYLE
                    series.format.line.color.rgb = rgb
                    series.marker.style = [XL_MARKER_STYLE.CIRCLE, XL_MARKER_STYLE.SQUARE, XL_MARKER_STYLE.DIAMOND,
                                           XL_MARKER_STYLE.TRIANGLE, XL_MARKER_STYLE.X, XL_MARKER_STYLE.STAR][i % 6]
                    series.marker.format.fill.solid()
                    series.marker.format.fill.fore_color.rgb = rgb
                    series.marker.format.line.color.rgb = rgb
        elif kind == 'interval':
            left, right = data['range']
            start, span = x + w * .29, w * .38
            row_h = min(.9, (usable - .6) / len(data['rows']))
            at = lambda value: start + (value - left) / (right - left) * span
            reference = data.get('reference', 0)
            line(slide, at(reference), y, at(reference), y + row_h * len(data['rows']), 'muted', 1,
                 key='visual/reference', value=reference)
            axis_y = y + row_h * len(data['rows']) + .02
            line(slide, start, axis_y, start + span, axis_y, 'muted', 1,
                 key='visual/axis-line', range=data['range'], scale='linear', unit=data['unit'])
            for i, row in enumerate(data['rows']):
                ry = y + i * row_h + row_h / 2
                key = f'visual/row/{i}'
                textbox(slide, row['label'], x, ry - .25, w * .28, .52, 18, name='interval-label', key=key + '/label')
                line(slide, at(row['low']), ry, at(row['high']), ry, thickness=3,
                     key=key + '/interval', low=row['low'], high=row['high'])
                dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(at(row['estimate']) - .055), Inches(ry - .055), Inches(.11), Inches(.11))
                dot.name = 'edu:point'
                tag(dot, key + '/estimate', value=row['estimate'])
                dot.fill.solid(); dot.fill.fore_color.rgb = color('accent'); dot.line.fill.background()
                textbox(slide, interval_text(data, row), x + w * .69, ry - .25, w * .31, .52, 17, name='interval-value', key=key + '/display')
            textbox(slide, f"{display_number(left)} ← {data['unit']} → {display_number(right)}", start, axis_y + .03, span, .45, 15, 'muted', name='axis', key='visual/axis-label')
            textbox(slide, '参考值 = ' + display_number(reference), x, axis_y + .03, w * .28, .45, 15, 'muted', name='reference-label', key='visual/reference-label')
        elif kind == 'image':
            from PIL import Image
            path = (base / data['path']).resolve()
            with Image.open(path) as img:
                iw, ih = img.size
            scale = min(w / iw, usable / ih)
            shape = slide.shapes.add_picture(str(path), Inches(x + (w - iw * scale) / 2), Inches(y + (usable - ih * scale) / 2), width=Inches(iw * scale), height=Inches(ih * scale))
            shape.name = 'edu:source-image'
            tag(shape, 'visual/image', asset_path=data['path'], asset_sha256=digest(path))
        elif kind == 'process':
            steps = data['steps']
            gap = .32
            box_h = min(.88, (usable - gap * (len(steps) - 1)) / len(steps))
            nodes = []
            for i, value in enumerate(steps):
                yy = y + i * (box_h + gap)
                shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(yy), Inches(w), Inches(box_h))
                shape.name = 'edu:process'
                tag(shape, f'visual/node/{i}')
                nodes.append(shape)
                shape.fill.solid(); shape.fill.fore_color.rgb = color('light'); shape.line.fill.background()
                fit(value, w - .18, box_h, 20)
                format_frame(shape.text_frame, value, 20)
                shape.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            for i, (first, second) in enumerate(zip(nodes, nodes[1:])):
                arrow = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, first.left + first.width // 2,
                                                   first.top + first.height, second.left + second.width // 2, second.top)
                arrow.name = 'edu:arrow'
                arrow.begin_connect(first, 2)
                arrow.end_connect(second, 0)
                arrow.line.color.rgb = color('accent')
                arrow.line.width = Pt(2)
                end = OxmlElement('a:tailEnd'); end.set('type', 'triangle')
                arrow._element.spPr.get_or_add_ln().append(end)
                tag(arrow, f'visual/edge/{i}', start=f'visual/node/{i}', end=f'visual/node/{i + 1}')
        elif kind == 'quote':
            textbox(slide, '“' + data['text'] + '”', x + .12, y + .2, w - .24, usable - 1.0, 25, name='quote', key='visual/quote')
            textbox(slide, data['attribution'], x + .12, y + usable - .65, w - .24, .5, 17, 'muted', name='attribution', key='visual/attribution')
        if note:
            textbox(slide, '坐标说明：' + note, x, y + usable + .02, w, .63, 16, 'muted', name='truncation-note', key='visual/truncation-note')
        if caption:
            textbox(slide, caption, x, y + h - .48 + .02, w, .48, 13, 'muted', name='caption', key='visual/caption')

    for index, page in enumerate(plan['slides'], 1):
        active_id = page['id']
        slide = prs.slides.add_slide(layout)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        if not template or (theme and 'background' in theme):
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = color('background')
        title_size = settings['title_size']
        textbox(slide, page['title'], width * .055, height * .045, width * .88, height * .092, title_size, bold=True, name='title')
        line(slide, width * .055, height * .155, width * .945, height * .155, thickness=2)
        textbox(slide, page['message'], width * .055, height * .166, width * .89, height * .061, 19, 'accent', name='message')
        x, y, w, h = [n * dim for n, dim in zip(settings['canvas'], [width, height, width, height])]
        body = page.get('body', [])
        data = page.get('visual')
        if data and body:
            body_w, visual_x, visual_w = w * .36, x + w * .40, w * .60
        else:
            body_w, visual_x, visual_w = w, x, w
        if body:
            gap = .16
            size = settings['body_size']
            heights = [max(.55, sum(max(1, math.ceil(units(line) * size / max(1, body_w * 72 - 14)))
                                    for line in t.split('\n')) * size * 1.22 / 72 + .1) for t in body]
            if sum(heights) + gap * (len(body) - 1) > h:
                raise ValueError(f'{active_id}: body too dense; split or shorten, do not shrink font')
            yy = y + .05
            for body_index, (text, bh) in enumerate(zip(body, heights)):
                textbox(slide, text, x, yy, body_w, bh, name='body', key=f'body/{body_index}')
                yy += bh + gap
        if data:
            visual(slide, data, visual_x, y + .05, visual_w, h - .05)
        footer = reference_text(plan, page)
        textbox(slide, footer, width * .055, height * .916, width * .81, height * .061, settings['footer_size'], 'muted', name='source')
        textbox(slide, str(index), width * .9, height * .921, width * .045, height * .045, 12, 'muted', name='number')
        frame = slide.notes_slide.notes_text_frame
        if frame is None:
            raise ValueError('Template notes master has no notes text placeholder')
        frame.text = page['notes'] + '\n\n来源：' + footer + f"\n预计讲解：{page['seconds']}秒"
    prs.core_properties.title = plan['title']
    prs.core_properties.subject = 'Education research presentation; source-linked plan'
    prs.save(str(destination))
    return [dict(slide=i, **identity(s)) for i, slide in enumerate(prs.slides, 1) for s in slide.shapes if identity(s)]


def build(plan_path, out, content_only=False, template=None, theme_path=None):
    plan_path, out = Path(plan_path).resolve(), Path(out).resolve()
    if out.exists() and any(out.iterdir()):
        raise ValueError('Output directory is not empty; use a fresh run directory')
    plan = load_plan(plan_path)
    theme = json.loads(Path(theme_path).read_text(encoding='utf-8-sig')) if theme_path else None
    export_text(plan, out)
    (out / 'plan.json').write_text(json.dumps(plan, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    objects = []
    if not content_only:
        objects = make_pptx(plan, out / 'slides.pptx', plan_path.parent, template, theme)
    total = sum(s['seconds'] for s in plan['slides'])
    warnings = []
    if abs(total - plan['minutes'] * 60) > max(30, plan['minutes'] * 6):
        warnings.append('Planned slide times differ from requested duration by more than 10%/30 seconds')
    manifest = {'plan_sha256': digest(plan_path), 'mode': 'content' if content_only else 'pptx',
                'source_plan': str(plan_path), 'slides': len(plan['slides']), 'planned_seconds': total,
                'requested_minutes': plan['minutes'], 'warnings': warnings, 'render_status': 'not_run',
                'template_sha256': digest(template) if template else None,
                'theme_sha256': digest(theme_path) if theme_path else None,
                'objects': objects,
                'outputs': {p.name: digest(p) for p in out.iterdir() if p.is_file()}}
    (out / 'build.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    return manifest


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--plan', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--content-only', action='store_true')
    parser.add_argument('--template', type=Path)
    parser.add_argument('--theme', type=Path)
    args = parser.parse_args()
    try:
        print(json.dumps(build(args.plan, args.out, args.content_only, args.template, args.theme), ensure_ascii=False, indent=2))
    except (ValueError, FileNotFoundError) as exc:
        sys.exit(str(exc))


if __name__ == '__main__':
    main()
