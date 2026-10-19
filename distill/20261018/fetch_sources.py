"""Fetch a registered development source set to a new local-only directory.

No source text is committed. HTTP success is not a claim that the paper was read.
Keep failed attempts and received bytes/metadata hashes for provenance.
"""
import argparse
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timezone
import hashlib
import io
import json
from pathlib import Path
import re

from bs4 import BeautifulSoup
from docx import Document
import fitz
import requests


def sha(data):
    return hashlib.sha256(data).hexdigest()


def extract(data, content_type):
    # Some publisher downloads prepend whitespace (LLT currently sends CRLF).
    # Keep the original bytes for hashing and let the PDF parser validate them.
    if data[:1024].lstrip(b' \t\r\n').startswith(b'%PDF-'):
        pdf = fitz.open(stream=data, filetype='pdf')
        return 'pdf', '\n\n'.join(
            f'## PDF page {i + 1}\n{p.get_text(sort=False)}' for i, p in enumerate(pdf)
        ), {'pages': len(pdf), 'order': 'PDF text stream; verify column order against page images before analysis'}
    if data.startswith(b'PK'):
        doc = Document(io.BytesIO(data))
        return 'docx', '\n\n'.join(f'P{i + 1}: {p.text}' for i, p in enumerate(doc.paragraphs)), {'paragraphs': len(doc.paragraphs)}
    if b'<article' in data[:2000] and b'<article-title' in data:
        soup = BeautifulSoup(data, 'xml')
        if soup.find('body') is None:
            raise ValueError('XML contains no full article body')
        nodes = [soup.find('front'), soup.find('body'), soup.find('back')]
        body = '\n\n'.join(n.get_text(' ', strip=True) for n in nodes if n is not None)
        # Retain the XML as the authoritative section/paragraph locator.
        return 'xml', body, {'sections': len(soup.find_all('sec')), 'locator': 'JATS section titles and paragraph order'}
    soup = BeautifulSoup(data, 'html.parser')
    # An abstract-only landing page or anti-bot page must not become full text.
    sections = soup.select('.c-article-section')
    if not sections:
        raise ValueError('No publisher article sections found; possible landing/challenge page')
    blocks = []
    for section in sections:
        heading = section.find(['h2', 'h3'])
        blocks.append('\n## ' + (heading.get_text(' ', strip=True) if heading else 'Untitled section'))
        for element in section.find_all(['h3', 'h4', 'p', 'table']):
            if element.find_parent('table') is not None:
                continue
            value = element.get_text(' ', strip=True)
            if value:
                blocks.append(value)
    text = '\n\n'.join(blocks)
    if len(text) < 12000 or not re.search(r'\b(Discussion|Conclusion|Findings)\b', text, re.I):
        raise ValueError('Article body too short or lacks expected substantive sections')
    return 'html', text, {'sections': len(sections), 'table_note': 'Linked tables may require separate reading; only inline content extracted'}


def fetch(source, root):
    folder = root / source['id']
    folder.mkdir()
    record = {'id': source['id'], 'doi': source['doi'], 'retrieved_utc': datetime.now(timezone.utc).isoformat(), 'attempts': [], 'reading_status': 'not_yet_read'}
    try:
        response = requests.get('https://api.crossref.org/works/' + source['doi'], timeout=45)
        response.raise_for_status()
        meta = response.json()['message']
        (folder / 'crossref.json').write_bytes(response.content)
        record['identity'] = {k: meta.get(k) for k in ['DOI', 'title', 'author', 'published', 'container-title', 'type', 'update-to', 'relation', 'license']}
        record['metadata_sha256'] = sha(response.content)
    except Exception as exc:
        record['metadata_error'] = str(exc)
    for index, url in enumerate(source['urls'], 1):
        attempt = {'url': url}
        record['attempts'].append(attempt)
        try:
            response = requests.get(url, timeout=60, headers={'User-Agent': 'edu-writing-source-audit/1.0'})
            attempt.update(status=response.status_code, final_url=response.url, content_type=response.headers.get('content-type'), bytes=len(response.content))
            response.raise_for_status()
            extension, body, details = extract(response.content, attempt['content_type'])
            # DOI may be represented with spaces/line breaks in PDF extraction.
            normalized = re.sub(r'\s+', '', body).lower()
            doi_found = source['doi'].lower() in normalized
            (folder / ('source.' + extension)).write_bytes(response.content)
            (folder / 'body.txt').write_text(body, encoding='utf-8', newline='\n')
            record.update(status='downloaded_for_reading', selected_url=url, format=extension, raw_sha256=sha(response.content), text_sha256=sha(body.encode('utf-8')), text_characters=len(body), extraction=details)
            record['identity_check'] = 'DOI present; title/authors still require reading' if doi_found else 'Manual title/author/version verification required: DOI absent'
            attempt['outcome'] = 'downloaded'
            break
        except Exception as exc:
            attempt.update(outcome='failed', error=str(exc))
    else:
        record['status'] = 'access_failed'
    (folder / 'acquisition.json').write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(source['id'], record['status'], flush=True)
    return record


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--manifest', type=Path, default=Path(__file__).with_name('sources.json'))
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--ids', nargs='+', help='Fetch only these registered source IDs')
    args = parser.parse_args()
    sources = json.loads(args.manifest.read_text(encoding='utf-8'))['sources']
    if args.ids:
        unknown = set(args.ids) - {source['id'] for source in sources}
        if unknown:
            raise SystemExit('Unregistered source IDs: ' + ', '.join(sorted(unknown)))
        sources = [source for source in sources if source['id'] in args.ids]
    if args.out.exists():
        raise SystemExit('Use a new output directory to retain previous acquisition records')
    args.out.mkdir(parents=True)
    with ThreadPoolExecutor(max_workers=3) as pool:
        records = list(pool.map(lambda source: fetch(source, args.out), sources))
    (args.out / 'acquisition.json').write_text(json.dumps(records, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


if __name__ == '__main__':
    main()
