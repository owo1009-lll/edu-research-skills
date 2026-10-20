"""Reproduce a wholly synthetic project outside the public checkout.

python evals/author-project-20261020/reproduce.py --out /path/to/new-example [--render]
No external sources, private manuscripts, statistical runs or user confirmations are fabricated:
all values and the example user's requests below are explicitly synthetic fixtures.
"""
import argparse
from copy import deepcopy
import json
from pathlib import Path
import sys

REPO = Path(__file__).resolve().parents[2]
for folder in ["edu-research", "edu-writing", "edu-ppt", "edu-shared"]:
    sys.path.insert(0, str(REPO / "skills" / folder / "scripts"))
from project_check import register, check, refresh, synchronize
from project_documents import read_document, replace_docx
from private_store import digest, private_path
from build_deck import build
from check_deck import check as check_deck
from author_preferences import operate


def save(path, obj):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def identity(statistic, **extra):
    value = dict(outcome="classroom engagement", comparison="intervention - comparison",
                 timepoint="post", sample="complete cases", statistic=statistic, unit="score",
                 model="baseline-adjusted", adjustments="pretest", analysis="primary",
                 uncertainty="95% confidence interval")
    value.update(extra)
    return value


def binding(key, file, locator, **extra):
    return dict(id=key, file=file, locator=locator,
                reason="Read the labeled synthetic outcome, model, sample and timepoint.", **extra)


def slide(key, title, message, visual):
    return dict(id=key, title=title, message=message, body=[], visual=visual,
                visual_plan="Editable native display from synthetic evidence.",
                notes=message + " These are synthetic validation data, not a real educational finding.",
                seconds=60, role="presenter_comment", status="not_applicable", section="Synthetic example",
                evidence=[dict(source="fixture", location=key, supports="All values are synthetic fixtures.")])


def run(out, render=False):
    from docx import Document
    out = private_path(out)
    if out.exists():
        raise ValueError("Use a new output directory.")
    out.mkdir(parents=True)
    requests = [
        "合成请求：以后改我的英文定量 Discussion，保留术语 instructional responsiveness；先解释发现，再与文献比较。",
        "合成请求：这是我确认的新结果，请核对正文、表格和 PPT 中的旧内容，先不要改稿。",
        "合成请求：同步修改已确认的数值和区间，保留原件。解释变化先列出来。",
        "合成请求：把能直接从新区间确定的文字描述也同步；不要扩写教学建议。",
    ]
    (out / "requests.md").write_text("# Synthetic user requests\n\n" + "\n\n".join(requests), encoding="utf-8")
    old = {"analysis_n": 390, "enrolled_n": 410, "intervention_post": 18, "comparison_post": 14,
           "difference": {"estimate": .009, "low": .004, "high": .014},
           "outcome": "classroom engagement", "unit": "score", "model": "baseline-adjusted",
           "unadjusted_beta": .21, "adjusted_beta": .21}
    new = deepcopy(old)
    new.update(analysis_n=382, intervention_post=16)
    new["difference"] = {"estimate": .006, "low": -.002, "high": .014}
    save(out / "analysis-old.json", old)
    save(out / "analysis.json", old)
    save(out / "analysis-new.json", new)
    word = Document()
    word.add_heading("Synthetic classroom study", 0)
    word.add_heading("Abstract", 1)
    word.add_paragraph("The analysis included 390 pupils; 410 were enrolled.")
    word.add_heading("Results", 1)
    word.add_paragraph("The adjusted difference was 0.009 [0.004, 0.014].")
    word.add_paragraph("The interval remained above zero.")
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Group", "Post score"
    for label, value in [("Intervention", "18"), ("Comparison", "14")]:
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    word.save(out / "paper.docx")
    plan = dict(title="Synthetic material-update example", route="own", stage="progress", language="en",
                minutes=3, sources=[dict(id="fixture", title="Synthetic validation data", short="Synthetic",
                                        access="author_material", year=2026)])
    plan["slides"] = [
        slide("GROUPS", "Groups at two measurement points", "Post scores: intervention 18, comparison 14.",
              dict(type="chart", kind="column", categories=["Pre", "Post"],
                   series=[dict(name="Intervention", values=[12, 18]), dict(name="Comparison", values=[12, 14])],
                   unit="score", minimum=0, maximum=20, major_unit=5, number_format="0",
                   caption="Synthetic values; no real intervention conclusion.")),
        slide("INTERVAL", "An estimate and its uncertainty", "Adjusted difference: 0.009 [0.004, 0.014].",
              dict(type="interval", rows=[dict(label="Adjusted difference", **old["difference"])],
                   range=[-.005, .02], reference=0, unit="score", caption="Synthetic interval.")),
        slide("SAMPLE", "Enrollment and analysis sample", "The enrolled and analyzed populations differ.",
              dict(type="table", columns=["Population", "N"], rows=[["Enrolled", "410"], ["Analysis", "390"]],
                   caption="Synthetic sample flow."))]
    save(out / "plan.json", plan)
    build(out / "plan.json", out / "before")
    inventory = {"paper": read_document(out / "paper.docx"),
                 "ppt": read_document(out / "before/slides.pptx")}
    save(out / "inventory-before.json", inventory)

    def loc(file, text=None, kind=None, suffix=None):
        matches = [u for u in inventory[file]["units"]
                   if (text is None or u["text"] == text) and (kind is None or u["kind"] == kind)
                   and (suffix is None or u["locator"].endswith(suffix))]
        if len(matches) != 1:
            raise ValueError(f"Ambiguous fixture lookup {file}/{text}/{kind}/{suffix}: {matches}")
        return matches[0]["locator"]

    files = [
        dict(id="analysis", path=str(out / "analysis.json"), role="confirmed_synthetic_results"),
        dict(id="paper", path=str(out / "paper.docx"), role="manuscript"),
        dict(id="plan", path=str(out / "plan.json"), role="presentation_plan"),
        dict(id="ppt", path=str(out / "before/slides.pptx"), role="presentation"),
    ]
    facts = []
    def fact(key, statistic, source_path, occurrences, **ident):
        source = binding(key + "-source", "analysis", "json:" + source_path,
                         guards=[dict(locator="json:/outcome", equals="classroom engagement"),
                                 dict(locator="json:/model", equals="baseline-adjusted"),
                                 dict(locator="json:/unit", equals="score")])
        facts.append(dict(id=key, identity=identity(statistic, **ident),
                          mapping_reason="Same labeled synthetic result; enrolled N stays a distinct fact.",
                          sources=[source], occurrences=occurrences,
                          authority=dict(source=source["id"], kind="user_confirmed",
                                         reason="The example request explicitly designates analysis.json.")))

    n_pattern = r"included (?P<value>\d+) pupils"
    fact("sample", "n", "/analysis_n", [
        binding("sample-word", "paper", loc("paper", text="The analysis included 390 pupils; 410 were enrolled."), pattern=n_pattern),
        binding("sample-plan", "plan", "json:/slides/2/visual/rows/1/1"),
        binding("sample-ppt", "ppt", loc("ppt", text="390", kind="cell"))], unit="people", uncertainty="not applicable")
    fact("post", "mean", "/intervention_post", [
        binding("post-word", "paper", loc("paper", text="18", kind="table_paragraph")),
        binding("post-plan", "plan", "json:/slides/0/visual/series/0/values/1"),
        binding("post-ppt", "ppt", loc("ppt", kind="chart_value", suffix=":series:0:value:1")),
        binding("post-message", "plan", "json:/slides/0/message", pattern=r"intervention (?P<value>\d+)"),
        binding("post-notes", "plan", "json:/slides/0/notes", pattern=r"intervention (?P<value>\d+)")],
         comparison="intervention group", model="descriptive", adjustments="none", uncertainty="not supplied")
    num = r"-?\d+(?:\.\d+)?"
    for name in ["estimate", "low", "high"]:
        # The scalar plan and native interval label are both inspected through existing fidelity checks.
        pattern = (rf"difference was (?P<value>{num})" if name == "estimate" else
                   rf"\[(?P<value>{num})," if name == "low" else rf", (?P<value>{num})\]")
        occurrences = [
            binding(name + "-word", "paper", loc("paper", text="The adjusted difference was 0.009 [0.004, 0.014]."),
                    pattern=pattern, display_decimals=3),
            binding(name + "-plan", "plan", "json:/slides/1/visual/rows/0/" + name),
        ]
        # Message/notes are mapped too: updating the native graphic alone would leave stale narration.
        for field in ("message", "notes"):
            pat = pattern.replace("difference was ", "difference: ")
            occurrences.append(binding(name + "-" + field, "plan", "json:/slides/1/" + field,
                                       pattern=pat, display_decimals=3))
        fact(name, "difference_" + name, "/difference/" + name, occurrences)
    spec = dict(files=files, facts=facts, decks=[dict(ppt="ppt", plan="plan")], dependencies=[
        dict(id="direction-reading", depends_on=["estimate", "low", "high"],
             target=binding("direction-word", "paper", loc("paper", text="The interval remained above zero.")),
             reason="Whether the interval excludes zero must be reconsidered.")])
    save(out / "spec.json", spec)
    baseline = register(out, spec)
    save(out / "baseline-check.json", baseline)
    hashes_before = {f["id"]: digest(Path(f["path"]).read_bytes()) for f in files if f["id"] != "analysis"}
    save(out / "analysis.json", new)
    changed = check(out)
    save(out / "changed-check.json", changed)
    assert hashes_before == {f["id"]: digest(Path(f["path"]).read_bytes()) for f in files if f["id"] != "analysis"}
    assert {x["fact_id"] for x in changed["findings"] if x["status"] == "stale"} == {"sample", "post", "estimate", "low"}
    assert any(x.get("dependency_id") == "direction-reading" and x["status"] == "review" for x in changed["findings"])
    # Combine all numeric replacements within the same paragraph/plan field.
    groups = {}
    for row in changed["findings"]:
        if row["status"] != "stale" or row["file_id"] not in {"paper", "plan"}:
            continue
        key = (row["file_id"], row["location"])
        before, after = groups.get(key, (row["current"], row["current"]))
        fact_id = row["fact_id"]
        if fact_id == "sample":
            after = after.replace("390", "382")
        elif fact_id == "post":
            after = after.replace("18", "16")
        elif fact_id == "estimate":
            after = after.replace("0.009", "0.006")
        elif fact_id == "low":
            after = after.replace("0.004", "-0.002")
        groups[key] = (before, after)
    edits = []
    for file, name in [("paper", "paper-updated.docx"), ("plan", "plan-updated.json")]:
        replacements = [dict(locator=locator, before=before, old=before, new=after,
                             reason="Explicitly authorized same-result numeric update.")
                        for (file_id, locator), (before, after) in groups.items() if file_id == file]
        edits.append(dict(file=file, output=str(out / name), replacements=replacements))
    request = dict(authorization=requests[2], report_fingerprint=changed["input_fingerprint"], edits=edits)
    save(out / "authorized-sync.json", request)
    sync = synchronize(out, request)
    save(out / "sync-result.json", sync)
    build(out / "plan-updated.json", out / "after")
    fidelity = check_deck(out / "after/slides.pptx", out / "plan-updated.json")
    assert fidelity["status"] == "passed", fidelity
    save(out / "deck-check.json", fidelity)
    paths = {"paper": str(out / "paper-updated.docx"), "plan": str(out / "plan-updated.json"),
             "ppt": str(out / "after/slides.pptx")}
    check(out, paths)
    for item in spec["facts"]:
        for occurrence in item["occurrences"]:
            occurrence["reviewed_sha256"] = digest(Path(paths.get(occurrence["file"],
                next(f["path"] for f in spec["files"] if f["id"] == occurrence["file"]))).read_bytes())
    numeric_final = refresh(out, dict(reader="Example runner", reason="Inspected exact numeric replacements and unchanged analysis labels.",
                              replace={"facts": spec["facts"]}))
    save(out / "numeric-only-check.json", numeric_final)
    assert not any(x["status"] == "stale" for x in numeric_final["findings"]), numeric_final
    # The interpretation remains explicitly pending, not silently rewritten or marked correct.
    assert any(x.get("dependency_id") == "direction-reading" and x["status"] == "review" for x in numeric_final["findings"])
    # A separate, explicitly authorized semantic edit on a new copy. This is not
    # an invented automatic text inference or a claim about educational efficacy.
    assert new["difference"]["low"] < 0 < new["difference"]["high"]
    interpretation_edit = dict(locator=spec["dependencies"][0]["target"]["locator"],
                               before="The interval remained above zero.", old="The interval remained above zero.",
                               new="The interval includes zero.",
                               reason="The new lower bound is -0.002 and upper bound 0.014.")
    save(out / "interpretation-edit.json", dict(authorization=requests[3], edit=interpretation_edit,
                                               limits="No intervention-benefit or practice recommendation is inferred."))
    replace_docx(out / "paper-updated.docx", [interpretation_edit], out / "paper-reviewed.docx")
    check(out, {"paper": str(out / "paper-reviewed.docx")})
    paper_sha = digest((out / "paper-reviewed.docx").read_bytes())
    for item in spec["facts"]:
        for occurrence in item["occurrences"]:
            if occurrence["file"] == "paper":
                occurrence["reviewed_sha256"] = paper_sha
    spec["dependencies"][0]["target"]["reviewed_sha256"] = paper_sha
    reviewed = refresh(out, dict(reader="Example runner", reason=interpretation_edit["reason"],
                                 replace={"facts": spec["facts"], "dependencies": spec["dependencies"]}))
    review_row = next(x for x in reviewed["findings"] if x.get("dependency_id") == "direction-reading")
    spec["dependencies"][0]["decision"] = dict(status="still_supported", reader="Example runner",
                                               reason=interpretation_edit["reason"],
                                               fingerprint=review_row["review_fingerprint"])
    final = refresh(out, dict(reader="Example runner", reason="Read the corrected interval description against both bounds.",
                              replace={"dependencies": spec["dependencies"]}))
    save(out / "final-check.json", final)
    assert not any(x["status"] in {"stale", "review"} for x in final["findings"]), final
    profile = out / "personal/author-preferences.json"
    preference = dict(category="expression", content="Preserve instructional responsiveness; explain findings before literature comparison.",
                      basis="explicit", status="active", scope="author", languages=["en"],
                      research_types=["quantitative"], sections=["discussion"], tasks=["revise"],
                      conditions=["Supported concept and judgment"], exceptions=["No causal strengthening"],
                      source=dict(kind="explicit", evidence=requests[0], time="2026-10-20",
                                  cross_project_authorization=requests[0]))
    operate(out, dict(action="add", scope="author", item=preference), profile)
    save(out / "loaded-preferences.json", operate(out, dict(action="load", context=dict(
        language="en", research_type="quantitative", section="discussion", task="revise")), profile))
    if render:
        from render_deck import render as render_ppt
        for version in ("before", "after"):
            result = render_ppt(out / version / "slides.pptx", out / version / "render")
            save(out / version / "render-result.json", result)
    summary = dict(synthetic=True, original_files_unchanged=True,
                   stale_positions_before_sync=sum(x["status"] == "stale" for x in changed["findings"]),
                   stale_positions_after_sync=sum(x["status"] == "stale" for x in final["findings"]),
                   intermediate_pending_interpretation="The interval remained above zero.",
                   reviewed_interpretation="The interval includes zero.",
                   native_deck_consistency=fidelity["status"],
                   visual_review="pending image inspection" if render else "not performed")
    save(out / "summary.json", summary)
    return summary


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--render", action="store_true")
    args = parser.parse_args()
    print(json.dumps(run(args.out, args.render), ensure_ascii=False, indent=2))
