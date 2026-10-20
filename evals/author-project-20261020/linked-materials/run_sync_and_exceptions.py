from forward_trial import *
from pptx import Presentation
def load(rel):
    return json.loads((ROOT/rel).read_text(encoding="utf-8"))
spec=load("spec.json")
report=load("reports/01-first-use-check.json")
inv=load("inventory-before.json")
new=load("materials/result-current.json")
oldrr=load("materials/result-old.json")["phase2"]["rr"]
newrr=new["phase2"]["rr"]

def edit(file,loc,transform,why):
    before=next(u["text"] for u in inv[file]["units"] if u["locator"]==loc)
    after=transform(before)
    assert before!=after
    return {"locator":loc,"before":before,"old":before,"new":after,"reason":why}

numeric_changes={
 "doc":[edit("doc","docx:p:2",lambda s:s.replace("干预组70.0%","干预组65.0%").replace(oldrr,newrr),"已确认正式试验主分析数值变化"),
        edit("doc","docx:p:9",lambda s:s.replace(oldrr,newrr),"未调整模型风险比及置信区间更新")],
 "md":[edit("md","line:2",lambda s:s.replace("70.0%","65.0%"),"已确认干预组比例"),
       edit("md","line:3",lambda s:s.replace(oldrr,newrr),"已确认未调整模型结果")],
 "plan":[edit("plan","json:/slides/1/body/0",lambda s:s.replace("70.0%","65.0%"),"确认干预组比例"),
         edit("plan","json:/slides/1/visual/series/0/values/0",lambda s:"65","确认百分数单位原生图表"),
         edit("plan","json:/slides/1/notes",lambda s:s.replace("70.0%","65.0%"),"确认讲稿对应比例"),
         edit("plan","json:/slides/2/body/0",lambda s:s.replace(oldrr,newrr),"确认未调整RR及区间")]}
outputs={"doc":"updated/manuscript.docx","md":"updated/discussion.md","plan":"updated/deck-plan.json"}
request={"authorization":load("requests.json")["sync"],"report_fingerprint":report["input_fingerprint"],
 "edits":[{"file":f,"output":str(ROOT/outputs[f]),"replacements":edits} for f,edits in numeric_changes.items()]}
write("sync-request.json",request)
r=command("02-sync",CHECK,"sync","--project",ROOT/"project","--request",ROOT/"sync-request.json")
if r.returncode:raise RuntimeError(r.stderr)

# A genuine dependency-only correction cannot be submitted through sync even
# after a reader has found the old claim; record the actual refusal.
interpretation_request={"authorization":request["authorization"],"report_fingerprint":report["input_fingerprint"],
 "edits":[{"file":"doc","output":str(ROOT/"probe/interpretation.docx"),
   "replacements":[edit("doc","docx:p:3",lambda s:s.replace("排除了无效应","未排除无效应"),"阅读新95%区间包含1，修正明确的统计解释，不替作者决定推广结论")]}]}
write("interpretation-sync-request.json",interpretation_request)
command("03-dependency-only-sync",CHECK,"sync","--project",ROOT/"project","--request",ROOT/"interpretation-sync-request.json")

# Authorized manual edits on NEW copies, recorded separately because sync does
# not implement CSV or interpretation-only writes.
with (ROOT/"materials/results.csv").open(encoding="utf-8",newline="") as f:rows=list(csv.reader(f))
rows[1][5]="65.0";rows[3][5]=newrr
with (ROOT/"updated/results.csv").open("w",encoding="utf-8",newline="") as f:csv.writer(f).writerows(rows)
doc=Document(ROOT/"updated/manuscript.docx")
for p in doc.paragraphs:
    if p.text=="完成率优势排除了无效应，说明课程适合推广。":
        p.text="完成率差异未排除无效应；原推广结论需作者重新判断。"
doc.save(ROOT/"updated/manuscript.docx")
text=(ROOT/"updated/discussion.md").read_text(encoding="utf-8")
text=text.replace("完成率优势排除了无效应，因此推荐立即普遍推广。",
                  "完成率差异未排除无效应；原立即普遍推广建议需作者重新判断。")
(ROOT/"updated/discussion.md").write_text(text,encoding="utf-8")
plan=load("updated/deck-plan.json")
plan["slides"][2]["message"]="区间包含无效应参考值1"
plan["slides"][2]["notes"]="未调整模型未排除无效应；调整模型仍不精确。"
write("updated/deck-plan.json",plan)
write("manual-copy-edits.json",{"authorization":request["authorization"],
 "csv":"仅更新正式试验干预组完成率和未调整RR行，保留70%先导满意率、196问卷样本、调整模型",
 "doc_md":"区间包含1是新材料直接确认的事实；推广建议明确留给作者重新判断",
 "ppt_plan":"标题与讲稿仅修正区间参考值关系；由原生计划重新构建",
 "untouched":"原件全部保留；损坏附件、Word页眉、未重算Excel公式仍列为未检查，子组歧义及敏感性精度问题仍保留"})
r=command("04-build-after",BUILD,"--plan",ROOT/"updated/deck-plan.json","--out",ROOT/"deck-after")
if r.returncode:raise RuntimeError(r.stderr)
command("05-check-deck",DECKCHECK,"--pptx",ROOT/"deck-after/slides.pptx","--plan",ROOT/"updated/deck-plan.json")
render=REPO/"skills/edu-ppt/scripts/render_deck.py"
command("06-render-before",render,"--pptx",ROOT/"deck-before/slides.pptx","--out",ROOT/"render-before")
command("07-render-after",render,"--pptx",ROOT/"deck-after/slides.pptx","--out",ROOT/"render-after")
command("08-check-deck-rendered",DECKCHECK,"--pptx",ROOT/"deck-after/slides.pptx","--plan",ROOT/"updated/deck-plan.json","--render",ROOT/"render-after")

reloc={f:str(ROOT/p) for f,p in {**outputs,"csv":"updated/results.csv","ppt":"deck-after/slides.pptx"}.items()}
write("relocations.json",reloc)
command("09-relocate",CHECK,"check","--project",ROOT/"project","--relocations",ROOT/"relocations.json")
afterfiles={x["id"]:reloc.get(x["id"],x["path"]) for x in spec["files"]}
afterinv={f:read_document(path) for f,path in afterfiles.items()}
write("inventory-after.json",afterinv)
# Actual reader inspected all changed prose: identities retained; no source/model
# swaps. Update reviewed hashes, then record only the still-supported explanation.
facts=deepcopy(spec["facts"]);deps=deepcopy(spec["dependencies"]);cites=deepcopy(spec["citations"])
for f in facts:
    for binding in f["sources"]+f["occurrences"]:
        if binding["file"] in reloc:binding["reviewed_sha256"]=afterinv[binding["file"]]["sha256"]
for d in deps:d["target"]["reviewed_sha256"]=afterinv[d["target"]["file"]]["sha256"]
for c in cites:c["claim"]["reviewed_sha256"]=afterinv[c["claim"]["file"]]["sha256"]
refresh={"reader":"independent forward-testing reader","reason":"已重新读取新文件实际段落、图表数据与讲稿，数值和各自阶段模型仍一致；推广结论保留未决",
 "replace":{"facts":facts,"dependencies":deps,"citations":cites}}
write("refresh-request.json",refresh)
command("10-refresh",CHECK,"refresh","--project",ROOT/"project","--request",ROOT/"refresh-request.json")
after=load("project/edu_output/project/check.json")
write("reports/02-after-sync-check.json",after)
shutil.copy2(ROOT/"project/edu_output/project/check.md",ROOT/"reports/02-after-sync-check.md")
write("reports/02-readback.json",{
 "originals_unchanged":{x["id"]:hashfile(x["path"])==load("original-hashes.json")[x["id"]] for x in spec["files"]},
 "chart_values":[list(s.values) for slide in Presentation(ROOT/"deck-after/slides.pptx").slides for shape in slide.shapes if shape.has_chart for s in shape.chart.series],
 "doc_paragraphs":[p.text for p in Document(ROOT/"updated/manuscript.docx").paragraphs],
 "csv_rows":rows,"discussion":text,
 "counts":{s:sum(x["status"]==s for x in after["findings"]) for s in ["stale","review","consistent","unchecked"]}})

# Isolated exception copies: the original checking and synced deliverables above
# stay byte-for-byte preserved.
exception_spec=deepcopy(spec)
for f in exception_spec["files"]:
    src=Path(f["path"]);dest=ROOT/"exceptions"/"inputs"/(f["id"]+src.suffix)
    dest.parent.mkdir(parents=True,exist_ok=True);shutil.copy2(src,dest);f["path"]=str(dest)
write("exceptions/spec.json",exception_spec)
command("11-exception-register",CHECK,"register","--project",ROOT/"exceptions/project","--spec",ROOT/"exceptions/spec.json")
exception_report=load("exceptions/project/edu_output/project/check.json")
exception_request=deepcopy(request)
exception_request["report_fingerprint"]=exception_report["input_fingerprint"]
exception_request["edits"]=[{"file":"doc","output":str(ROOT/"exceptions/should-not-write.docx"),"replacements":numeric_changes["doc"]}]
write("exceptions/stale-request.json",exception_request)
paths={f["id"]:Path(f["path"]) for f in exception_spec["files"]}
doc=Document(paths["doc"]);doc.paragraphs[0].insert_paragraph_before("作者在核对后新增的方法说明。");doc.save(paths["doc"])
command("12-original-reedited-stale-sync",CHECK,"sync","--project",ROOT/"exceptions/project","--request",ROOT/"exceptions/stale-request.json")
rpt=load("exceptions/project/edu_output/project/check.json")
write("reports/03-original-reedited.json",rpt)
write("exceptions/readedit-safety.json",{"old_request_output_exists":(ROOT/"exceptions/should-not-write.docx").exists(),
 "doc_primary_locations":[{"location":x["location"],"status":x["status"]} for x in rpt["findings"] if x.get("fact_id")=="primary-rr" and x["file_id"]=="doc"]})
md=paths["md"];expanded=md.read_text(encoding="utf-8").replace(
 "同伴反馈帮助受访学生解释作业要求。[SYN-01]",
 "同伴反馈帮助受访学生解释作业要求，并显著提高了课程完成率。[SYN-01]")
md.write_text(expanded,encoding="utf-8")
command("13-claim-expanded",CHECK,"check","--project",ROOT/"exceptions/project")
write("reports/04-claim-expanded.json",load("exceptions/project/edu_output/project/check.json"))
# Structural cell relabeling is a different model and must not remain stale.
with paths["csv"].open(encoding="utf-8",newline="") as f:rs=list(csv.reader(f))
rs[3][3]="基线协变量调整"
with paths["csv"].open("w",encoding="utf-8",newline="") as f:csv.writer(f).writerows(rs)
command("14-model-label-changed",CHECK,"check","--project",ROOT/"exceptions/project")
write("reports/05-model-label-changed.json",load("exceptions/project/edu_output/project/check.json"))

# Historical update case uses a separate, explicitly documented prior reading.
hist=deepcopy(spec)
prior=load("materials/result-old.json")
prior["phase2"]["sensitivity"]=new["phase2"]["sensitivity"]
write("history/previous-confirmed-source.json",prior)
next(x for x in hist["files"] if x["id"]=="new")["path"]=str(ROOT/"history/previous-confirmed-source.json")
write("history/spec.json",hist)
command("15-historical-register",CHECK,"register","--project",ROOT/"history/project","--spec",ROOT/"history/spec.json")
write("history/source-relocation.json",{"new":str(ROOT/"materials/result-current.json")})
command("16-historical-update",CHECK,"check","--project",ROOT/"history/project","--relocations",ROOT/"history/source-relocation.json")
histreport=load("history/project/edu_output/project/check.json")
write("reports/06-historical-update.json",histreport)
deps=deepcopy(hist["dependencies"])
review=next(x for x in histreport["findings"] if x.get("dependency_id")=="discussion-cautious")
next(x for x in deps if x["id"]=="discussion-cautious")["decision"]={
 "status":"still_supported","reader":"independent forward-testing reader",
 "reason":"新RR=1.083333仍高于1，区间较宽且跨1；原句只说点估计倾向干预组、需关注区间宽度，仍受支持",
 "fingerprint":review["review_fingerprint"]}
write("history/decision.json",{"reader":"independent forward-testing reader","reason":"实际阅读点估计和区间后保留有依据的谨慎句","replace":{"dependencies":deps}})
command("17-reviewed-still-supported",CHECK,"refresh","--project",ROOT/"history/project","--request",ROOT/"history/decision.json")
write("reports/07-still-supported.json",load("history/project/edu_output/project/check.json"))
print("Completed sync/readback/exceptions/history.")
