"""Independent generated-material trial. No repository tests/evals or user data."""
from pathlib import Path
import sys
sys.dont_write_bytecode = True
import json, math, csv, shutil, hashlib, subprocess, os
from copy import deepcopy
from docx import Document
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent
REPO = Path(r"@REPO@")
CHECK = REPO / "skills/edu-research/scripts/project_check.py"
BUILD = REPO / "skills/edu-ppt/scripts/build_deck.py"
DECKCHECK = REPO / "skills/edu-ppt/scripts/check_deck.py"
sys.path.insert(0, str(CHECK.parent))
import project_check as pc
from project_documents import read_document

def write(path, data):
    path = ROOT/path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2)+"\n", encoding="utf-8")
    return path

def command(name, *args):
    env = dict(os.environ, PYTHONDONTWRITEBYTECODE="1", PYTHONIOENCODING="utf-8")
    r = subprocess.run([sys.executable, *map(str,args)], capture_output=True, text=True, encoding="utf-8", env=env)
    write(f"logs/{name}.json", {"args":list(map(str,args)), "returncode":r.returncode,"stdout":r.stdout,"stderr":r.stderr})
    print(name, r.returncode, r.stderr[:250])
    return r

def hashfile(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()

def generate():
    (ROOT/"materials").mkdir(exist_ok=True)
    old = {"phase1":{"n":48,"survey_n":42,"completion":35/48},
           "phase2":{"n":240,"survey_n":196,"treated_n":120,"control_n":120,
                     "complete_treated":84,"complete_control":72,"rate_treated":.70,"rate_control":.60,
                     "outcome":"8周课程完成","sample":"正式试验全部随机分组学生",
                     "model":"未调整风险比","rr":"RR=1.166667; 95%CI=[1.001200,1.359400]",
                     "adjusted":"RR=1.120000; 95%CI=[0.930000,1.350000]"}}
    new=deepcopy(old)
    new["phase2"]["complete_treated"]=78
    new["phase2"]["rate_treated"]=.65
    rr=(78/120)/(72/120)
    se=math.sqrt(1/78-1/120+1/72-1/120)
    lo,hi=math.exp(math.log(rr)-1.96*se),math.exp(math.log(rr)+1.96*se)
    new["phase2"]["rr"]=f"RR={rr:.6f}; 95%CI=[{lo:.6f},{hi:.6f}]"
    write("materials/result-old.json",old)
    write("materials/result-current.json",new)
    with (ROOT/"materials/participants-recounted.csv").open("w",encoding="utf-8",newline="") as f:
        w=csv.writer(f);w.writerow(["student_id","phase","group","completion"])
        for group,comp in (("干预",78),("对照",72)):
            for i in range(120): w.writerow([group+str(i+1),2,group,int(i<comp)])
    write("materials/recount-record.json",{"label":"SYNTHETIC, not real research","formula":"Risk ratio; independent-binomial log-Wald CI",
            "deduplicated_old_completions":6,"n":240,"rr":rr,"low":lo,"high":hi,
            "old_ci":"Synthetic prior mixed-cluster model output, not independently recomputed",
            "authority":"Requester explicitly selects result-current.json for the formal-trial primary result only"})
    doc=Document()
    doc.add_heading("同伴反馈课程：两阶段研究（合成试用材料）",0)
    doc.add_paragraph("摘要：先导阶段纳入48人，其中42人完成问卷；正式试验纳入240人，问卷有效样本为196人。")
    doc.add_paragraph("正式试验8周课程完成率：干预组70.0%，对照组60.0%；未调整风险比RR=1.166667; 95%CI=[1.001200,1.359400]。")
    doc.add_paragraph("完成率优势排除了无效应，说明课程适合推广。")
    tab=doc.add_table(rows=1, cols=3)
    for c,t in zip(tab.rows[0].cells,["分析","统计量","结果"]):c.text=t
    for vals in [("正式试验","未调整RR",old["phase2"]["rr"]),("正式试验","调整RR",old["phase2"]["adjusted"]),("先导阶段","问卷样本n","42")]:
        for c,t in zip(tab.add_row().cells,vals):c.text=t
    doc.sections[0].header.paragraphs[0].text="摘要快览：干预组70.0%"
    doc.save(ROOT/"materials/manuscript.docx")
    with (ROOT/"materials/results.csv").open("w",encoding="utf-8",newline="") as f:
        w=csv.writer(f);w.writerow(["phase","sample","outcome","model","statistic","value"])
        for row in [
          ["正式","全部240人","8周课程完成","描述统计","干预组完成率","70.0"],
          ["正式","全部240人","8周课程完成","描述统计","对照组完成率","60.0"],
          ["正式","全部240人","8周课程完成","未调整风险比","RR",old["phase2"]["rr"]],
          ["正式","全部240人","8周课程完成","基线协变量调整","RR",old["phase2"]["adjusted"]],
          ["先导","问卷有效42人","满意度","描述统计","高满意比例","70.0"],
          ["正式","问卷有效196人","问卷","描述统计","有效样本n","196"]]:w.writerow(row)
    text = """# 讨论（合成材料）
正式试验干预组课程完成率70.0%，对照组为60.0%。
未调整风险比RR=1.166667; 95%CI=[1.001200,1.359400]。
完成率优势排除了无效应，因此推荐立即普遍推广。
先导阶段问卷有效样本为42人；正式试验问卷有效样本为196人。两阶段样本不可相加。
协变量调整模型RR=1.120000; 95%CI=[0.930000,1.350000]，提示估计不够精确。
某子组课程完成率70.0%，分母和组别仍待作者确认。
从新材料看，完成率的点估计仍向干预组倾斜，但需要关注区间宽度。
引用论断：同伴反馈帮助受访学生解释作业要求。[SYN-01]
"""
    (ROOT/"materials/discussion.md").write_text(text,encoding="utf-8")
    (ROOT/"materials/synthetic-reading.md").write_text(
        "# SYN-01，虚构研究，仅用于测试\n8名受访学生描述，同伴反馈帮助他们解释作业要求；本研究没有测量课程完成率。\n",encoding="utf-8")
    write("materials/citation-ledger.json",{"SYN-01":{"claim":"同伴反馈帮助受访学生解释作业要求","support":"supported","scope":"8名访谈学生，自陈，未测量完成率","reader":"independent trial"}})
    book=Workbook();sheet=book.active;sheet.title="辅助计算"
    sheet.append(["指标","完成人数","分母","比例公式"]);sheet.append(["干预组",78,120,"=B2/C2"])
    book.save(ROOT/"materials/formula-supplement.xlsx")
    (ROOT/"materials/broken-supplement.docx").write_bytes(b"Truncated author attachment: not a zip")
    plan={"route":"own","stage":"completed","title":"同伴反馈课程：两阶段研究（合成）","minutes":2,
          "sources":[{"id":"author","title":"作者提供的合成分析材料","access":"author_material","short":"合成材料"}],"slides":[]}
    for sid,title,message,body,visual,notes in [
      ("sample","两阶段样本","样本口径必须分开",["先导阶段 n=48；问卷 n=42","正式试验 n=240；问卷 n=196"],None,"两个阶段独立，不能把问卷有效人数作为随机样本数。"),
      ("completion","正式试验课程完成率","8周课程完成",["干预组70.0%；对照组60.0%"],
       {"type":"chart","kind":"column","unit":"%","categories":["干预组","对照组"],"series":[{"name":"完成率","values":[70,60]}],"minimum":0,"maximum":100,"number_format":"0.0"},
       "干预组70.0%，对照组60.0%。"),
      ("effect","未调整风险比","区间下限高于1",["RR=1.166667; 95%CI=[1.001200,1.359400]"],
       {"type":"table","columns":["模型","RR及95%CI"],"rows":[["调整模型","1.12 [0.93,1.35]"]]},
       "未调整模型排除了无效应；调整模型仍不精确。")]:
        page={"id":sid,"title":title,"message":message,"body":body,"notes":notes,"visual_plan":"原生图表或文字",
              "seconds":40,"role":"finding","status":"completed",
              "evidence":[{"source":"author","location":sid,"supports":message}]}
        if visual:page["visual"]=visual
        plan["slides"].append(page)
    write("materials/deck-plan.json",plan)
    r=command("build-before",BUILD,"--plan",ROOT/"materials/deck-plan.json","--out",ROOT/"deck-before")
    if r.returncode:raise RuntimeError(r.stderr)
    requests={"check":"我已重新计算课程完成率，请以我明确指定的新结果 result-current.json 中正式试验主分析为准，检查 Word 摘要、结果表、Markdown 讨论和 PPT 中还用了哪些旧内容；先别改稿。辅助计算表和损坏附件也属于本次提交。",
              "sync":"同步修改已确定的旧内容，保留原件；歧义不要替我决定。",
              "exceptions":"原件被重新编辑；引用论断扩写；部分文件无法解析。"}
    write("requests.json",requests)
    files=[("new","materials/result-current.json","明确指定的当前正式试验结果"),("old","materials/result-old.json","旧版本，非当前权威"),
           ("doc","materials/manuscript.docx","Word摘要及结果表"),("csv","materials/results.csv","结果表"),
           ("md","materials/discussion.md","讨论"),("plan","materials/deck-plan.json","已有关联PPT计划"),
           ("ppt","deck-before/slides.pptx","PPT"),("reading","materials/synthetic-reading.md","虚构来源阅读文件"),
           ("ledger","materials/citation-ledger.json","已有引用支持记录"),
           ("formula","materials/formula-supplement.xlsx","辅助计算"),("broken","materials/broken-supplement.docx","损坏附录")]
    spec={"files":[{"id":id,"path":str(ROOT/p),"role":role} for id,p,role in files],"facts":[],"dependencies":[],"citations":[]}
    write("scope.json",spec)
    inv={x["id"]:read_document(x["path"]) for x in spec["files"]}
    write("inventory-before.json",inv)
    write("original-hashes.json",{x["id"]:hashfile(x["path"]) for x in spec["files"]})
    for file,d in inv.items():
        print(file, d["status"],len(d["units"]), len(d["unchecked"]))
        if file in {"doc","ppt","md"}:
            for u in d["units"]:print(" ",u["locator"],repr(u["text"])[:220])

if __name__=="__main__":
    generate()
