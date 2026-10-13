"""Reproduce observable defects against a selected checkout; synthetic data only.

Run from repository root. --scripts can point to the scripts of an older checkout.
The report records observations, not a quality score. Outputs use a fresh directory.
"""
from pathlib import Path
import argparse
import copy
import json
import sys


def run(scripts, out):
    sys.path.insert(0, str(Path(scripts).resolve()))
    from build_deck import build
    from check_deck import check
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    root = Path(__file__).resolve().parents[2]
    value = json.loads((root / 'skills/edu-ppt/examples/progress-plan.json').read_text(encoding='utf-8'))
    out = Path(out)
    out.mkdir(parents=True, exist_ok=False)
    report = {}
    for case in ('normal', 'table_changed', 'table_deleted', 'categories_changed',
                 'series_changed', 'unit_changed', 'small_interval', 'clipped_axis', 'two_series'):
        plan = copy.deepcopy(value)
        if case == 'small_interval':
            visual = plan['slides'][-1]['visual']
            visual.update(rows=[{'label': 'Small difference', 'estimate': .009, 'low': .004, 'high': .014}], range=[-.001, .02])
        elif case == 'clipped_axis':
            plan['slides'][1]['visual']['maximum'] = 5
        elif case == 'two_series':
            plan['slides'][1]['visual']['series'].append({'name': 'Control', 'values': [11, 9]})
        path = out / (case + '.json')
        path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding='utf-8')
        deck = out / case
        try:
            build(path, deck)
        except ValueError as exc:
            report[case] = {'build': 'rejected', 'reason': str(exc)}
            continue
        pptx = deck / 'slides.pptx'
        prs = Presentation(pptx)
        table = next(s for s in prs.slides[3].shapes if s.has_table)
        chart = next(s.chart for s in prs.slides[1].shapes if s.has_chart)
        if case == 'table_changed':
            table.table.cell(3, 1).text = '显著提升，p < .001'
        elif case == 'table_deleted':
            table._element.getparent().remove(table._element)
        elif case in ('categories_changed', 'series_changed'):
            data = CategoryChartData()
            data.categories = ['Changed A', 'Changed B'] if case == 'categories_changed' else plan['slides'][1]['visual']['categories']
            data.add_series('Changed series' if case == 'series_changed' else chart.series[0].name, list(chart.series[0].values))
            chart.replace_data(data)
        elif case == 'unit_changed':
            chart.value_axis.axis_title.text_frame.text = 'Changed unit'
        prs.save(pptx)
        result = {'build': 'generated', 'check': check(pptx, path)}
        if case == 'small_interval':
            result['display'] = [s.text for s in prs.slides[-1].shapes if s.has_text_frame and '[' in s.text]
        if case == 'two_series':
            result['series_colors'] = [str(s.format.fill.fore_color.rgb) for s in chart.series]
        report[case] = result
    (out / 'observations.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps({k: {'build': v['build'], 'check': v.get('check', {}).get('status'),
                          'display': v.get('display'), 'colors': v.get('series_colors')} for k, v in report.items()}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--scripts', default='skills/edu-ppt/scripts')
    parser.add_argument('--out', required=True)
    args = parser.parse_args()
    run(args.scripts, args.out)
