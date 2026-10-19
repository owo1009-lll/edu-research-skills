"""Public synthetic manuscripts, not user data or real research.

Run from any directory: python create_fixture.py --output NEW_DIRECTORY
Outputs native DOCX/PDF annotation examples, internal plans and verification.
Word opening/rendering is performed separately and must not be inferred here.
"""
import argparse
from datetime import datetime, timezone
import importlib.util
import io
import json
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
SCRIPT = ROOT / 'skills/edu-review/scripts/annotate_manuscript.py'
spec = importlib.util.spec_from_file_location('annotator', SCRIPT)
annotator = importlib.util.module_from_spec(spec)
spec.loader.exec_module(annotator)


def illustration():
    from PIL import Image, ImageDraw
    im = Image.new('RGB', (420, 170), 'white')
    draw = ImageDraw.Draw(im)
    draw.text((12, 8), 'SYNTHETIC ILLUSTRATION - NO RESEARCH DATA', fill='black')
    draw.line((35, 35, 35, 140, 390, 140), fill='black', width=2)
    draw.rectangle((80, 70, 150, 139), fill='#4b80b8')
    draw.rectangle((230, 70, 300, 139), fill='#b87c38')
    draw.text((80, 145), 'Group A', fill='black')
    draw.text((230, 145), 'Group B', fill='black')
    buf = io.BytesIO()
    im.save(buf, format='PNG')
    return buf.getvalue()


def create_docx(path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_COLOR_INDEX
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Calibri'
    style.font.size = Pt(11)
    doc.add_heading('Feedback and participation: synthetic review fixture', 0)
    doc.add_paragraph('This is a public software test manuscript. It reports no real study or participant data.')
    p = doc.add_paragraph('The abstract reports ')
    p.add_run('52').bold = True
    p.add_run(' students. The methods and Table 1 report 48 students.')
    doc.add_heading('Methods and findings', 1)
    doc.add_paragraph('Feedback supports learning in important ways. The intended meaning of participation is not specified here.')
    doc.add_paragraph('The study included 48 students. We have not supplied a record of when the analysis plan was fixed.')
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = 'Group', 'n'
    for values in [('A', '24'), ('B', '24'), ('Total', '48')]:
        for c, text in zip(table.add_row().cells, values):
            c.text = text
    doc.add_paragraph('Table 1. Synthetic participant counts. Total: 48.')
    doc.add_paragraph('The table provides counts only; the basis for selecting the groups needs verification.')
    doc.add_picture(io.BytesIO(illustration()), width=Inches(4.2))
    doc.add_paragraph('Figure 1. Synthetic illustration; no measured effect is depicted.')
    p = doc.add_paragraph('Existing author note: ')
    marked = p.add_run('keep this original comment and green highlight')
    marked.font.highlight_color = WD_COLOR_INDEX.BRIGHT_GREEN
    doc.add_comment(marked, text='Original author comment: do not remove.', author='Original Author', initials='OA')
    p = doc.add_paragraph('Existing revision: ')
    deleted = OxmlElement('w:del')
    deleted.set(annotator.w('id'), '21')
    deleted.set(annotator.w('author'), 'Original Author')
    deleted.set(annotator.w('date'), '2026-10-01T00:00:00Z')
    r, t = OxmlElement('w:r'), OxmlElement('w:delText')
    t.text = 'old wording'
    r.append(t)
    deleted.append(r)
    inserted = OxmlElement('w:ins')
    for key in ('id', 'author', 'date'):
        inserted.set(annotator.w(key), '22' if key == 'id' else deleted.get(annotator.w(key)))
    r, t = OxmlElement('w:r'), OxmlElement('w:t')
    t.text = 'revised wording'
    r.append(t)
    inserted.append(r)
    p._p.extend([deleted, inserted])
    p = doc.add_paragraph('Citation field: ')
    field = OxmlElement('w:fldSimple')
    field.set(annotator.w('instr'), 'REF example_reference')
    r, t = OxmlElement('w:r'), OxmlElement('w:t')
    t.text = '[1]'
    r.append(t)
    field.append(r)
    p._p.append(field)
    doc.add_paragraph('[1] Demonstration placeholder, not a real bibliographic reference.')
    doc.save(path)


def create_pdf(path):
    import fitz
    doc = fitz.open()
    p = doc.new_page(width=595, height=842)
    lines = [
        'Feedback and participation: synthetic review fixture',
        'PUBLIC SOFTWARE TEST: no real study or participant data.',
        'The abstract reports 52 students; Table 1 reports 48.',
        'Feedback supports learning in important ways.',
        'The timing of the analysis plan requires verification.',
        'Table 1. Synthetic counts',
        'Group A: 24', 'Group B: 24', 'Total: 48',
        '[1] Demonstration citation, not a real source.',
    ]
    for i, text in enumerate(lines):
        p.insert_text((45, 50 + i * 27), text, fontsize=11)
    p.insert_image(fitz.Rect(45, 340, 465, 510), stream=illustration())
    note = p.add_text_annot((500, 110), 'Original PDF comment: do not remove.')
    note.set_info(title='Original Author')
    note.update()
    # Rotated text and repeated numbers expose ambiguous string-search anchoring.
    p.insert_text((530, 660), 'Rotated 48 students', rotate=90, fontsize=10)
    doc.save(path)
    doc.close()


def anchor(inv, phrase, *, contains=None, in_table=None):
    matches = [u for u in inv['units'] if phrase in u['text'] and
               (contains is None or contains in u['text']) and
               (in_table is None or u.get('in_table') == in_table)]
    if len(matches) != 1:
        raise ValueError(f'Ambiguous fixture anchor {phrase!r}: {len(matches)} units')
    u = matches[0]
    start = u['text'].index(phrase)
    return {k:v for k,v in u.items() if k not in ('text', 'in_table')} | {
        'start': start, 'end': start + len(phrase), 'quote': phrase}


def plan_for(path):
    inv = annotator.inventory(path)
    entries = [
        ('C1', 'error', '52', 'abstract reports', '摘要写52，方法和表1合计48。请核对原始人数并统一；这里是稿内矛盾，尚未判定哪一数值真实。'),
        ('C2', 'improvement', 'Feedback supports learning in important ways.', None, '请说明本文的反馈具体指什么，以及它与参与这一结局有什么关系；这会使研究问题更具体。'),
    ]
    if path.suffix == '.docx':
        entries.append(('C3', 'verify', 'when the analysis plan was fixed', None, '请核实分析计划形成时间并补充依据；缺少记录不等于未事先制定。'))
        entries.append(('C4', 'verify', '48', None, '表内总数与方法一致。请据原始记录核对摘要人数；参见摘要批注C1。'))
    else:
        entries.append(('C3', 'verify', 'timing of the analysis plan', None, '请核实计划形成时间；不要将缺少记录断言为事后设定。'))
    annotations = []
    for aid, kind, quote, contains, comment in entries:
        annotations.append({'id': aid, 'kind': kind, 'severity': 'major' if aid == 'C1' else 'minor',
                            'comment': comment, 'anchors': [anchor(inv, quote, contains=contains,
                            in_table=True if aid == 'C4' else None)]})
    return {'schema_version': 1, 'source_sha256': inv['source_sha256'],
            'created_utc': datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ'),
            'annotations': annotations}


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('--output', type=Path, required=True)
    args = ap.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    results = []
    for extension, create in [('docx', create_docx), ('pdf', create_pdf)]:
        source = args.output / ('synthetic_manuscript.' + extension)
        create(source)
        plan = plan_for(source)
        (args.output / (extension + '-plan.json')).write_text(json.dumps(plan, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
        result = annotator.apply(source, plan)
        result['output'] = Path(result['output']).name
        results.append(result)
    (args.output / 'verification.json').write_text(json.dumps(results, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(results, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
