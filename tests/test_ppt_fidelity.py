"""Adversarial native-object edits plus valid outputs; no Office/network needed."""
from pathlib import Path
import copy
import json
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'skills/edu-ppt/scripts'))
from build_deck import build
from check_deck import check
from plan import validate
from objects import identity
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.util import Inches


class FidelityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='edu-ppt-fidelity-')
        self.root = Path(self.tmp.name)
        self.base = json.loads((ROOT / 'skills/edu-ppt/examples/progress-plan.json').read_text(encoding='utf-8'))
        self.count = 0

    def tearDown(self):
        self.tmp.cleanup()

    def deck(self, index=1, visual=None, theme=None):
        plan = copy.deepcopy(self.base)
        plan['slides'] = [plan['slides'][index]]
        plan['slides'][0]['body'] = []
        if visual is not None:
            plan['slides'][0]['visual'] = visual
        self.count += 1
        path = self.root / f'plan-{self.count}.json'
        path.write_text(json.dumps(plan, ensure_ascii=False), encoding='utf-8')
        theme_path = None
        if theme is not None:
            theme_path = self.root / f'theme-{self.count}.json'
            theme_path.write_text(json.dumps(theme), encoding='utf-8')
        out = self.root / f'run-{self.count}'
        build(path, out, theme_path=theme_path)
        pptx = out / 'slides.pptx'
        self.assertEqual(check(pptx, path)['errors'], [])
        return path, pptx, Presentation(pptx)

    def obj(self, prs, suffix):
        return next(s for s in prs.slides[0].shapes if identity(s).get('id', '').endswith('/' + suffix))

    def rejected(self, path, pptx, prs, field):
        prs.save(pptx)
        report = check(pptx, path)
        self.assertEqual(report['status'], 'failed')
        matching = [i for i in report['issues'] if field in i['field']]
        self.assertTrue(matching, report['errors'])
        self.assertTrue(all(i['slide'] == 1 and i['object'] and 'expected' in i and 'actual' in i for i in matching))

    def test_table_cells_header_structure_and_removal(self):
        for mutation in ('cell', 'header', 'row', 'delete'):
            with self.subTest(mutation=mutation):
                path, pptx, prs = self.deck(3)
                shape = self.obj(prs, 'visual/table')
                if mutation == 'cell':
                    shape.table.cell(3, 1).text = '显著提升，p < .001'
                elif mutation == 'header':
                    shape.table.cell(0, 1).text = '教学成效'
                elif mutation == 'row':
                    shape.table._tbl.remove(shape.table._tbl.tr_lst[-1])
                else:
                    shape._element.getparent().remove(shape._element)
                self.rejected(path, pptx, prs, {'row': 'rows/columns', 'delete': 'presence'}.get(mutation, 'cell['))

    def test_chart_labels_series_order_values_units_and_axes(self):
        visual = copy.deepcopy(self.base['slides'][1]['visual'])
        visual['series'].append({'name': '对照组', 'values': [10, 7]})
        for mutation in ('categories', 'category_order', 'name', 'series_order', 'values', 'unit', 'maximum', 'format', 'axis_format', 'reverse', 'legend'):
            with self.subTest(mutation=mutation):
                path, pptx, prs = self.deck(visual=visual)
                chart = self.obj(prs, 'visual/chart').chart
                if mutation in ('categories', 'category_order', 'name', 'series_order', 'values'):
                    data = CategoryChartData()
                    categories = visual['categories']
                    data.categories = ['A', 'B'] if mutation == 'categories' else list(reversed(categories)) if mutation == 'category_order' else categories
                    series = copy.deepcopy(visual['series'])
                    if mutation == 'name':
                        series[0]['name'] = '实验组已证实有效'
                    if mutation == 'series_order':
                        series.reverse()
                    if mutation == 'values':
                        series[0]['values'][0] = 99
                    for s in series:
                        data.add_series(s['name'], s['values'])
                    chart.replace_data(data)
                elif mutation == 'unit':
                    chart.value_axis.axis_title.text_frame.text = '百分比'
                elif mutation == 'maximum':
                    chart.value_axis.maximum_scale = 5
                elif mutation == 'format':
                    chart.plots[0].data_labels.number_format = '0%'
                elif mutation == 'axis_format':
                    chart.value_axis.tick_labels.number_format = '0%'
                elif mutation == 'reverse':
                    chart.category_axis.reverse_order = True
                else:
                    chart.has_legend = False
                field = {'categories': 'categories', 'category_order': 'categories', 'name': 'series names',
                         'series_order': 'series names', 'values': 'values', 'unit': 'unit',
                         'maximum': 'axis.maximum', 'format': 'number format', 'axis_format': 'axis number format',
                         'reverse': 'category direction', 'legend': 'legend'}[mutation]
                self.rejected(path, pptx, prs, field)

    def test_explicit_axis_validation_and_intentional_truncation(self):
        for settings in ({'maximum': 5}, {'minimum': 10, 'maximum': 5}, {'maximum': float('inf')},
                         {'maximum': 5, 'allow_truncation': True}, {'minimum': 7}, {'major_unit': 0}):
            plan = copy.deepcopy(self.base)
            plan['slides'][1]['visual'].update(settings)
            with self.subTest(settings=settings), self.assertRaises(ValueError):
                validate(plan)
        visual = copy.deepcopy(self.base['slides'][1]['visual'])
        visual.update(maximum=5, allow_truncation=True, truncation_note='纵轴截于5；完整计数为12和8。')
        path, pptx, prs = self.deck(visual=visual)
        self.assertIn('纵轴截于5', self.obj(prs, 'visual/truncation-note').text)
        self.obj(prs, 'visual/truncation-note').text = ''
        self.rejected(path, pptx, prs, 'text')
        # A line chart need not start at zero when the points are within range.
        visual = copy.deepcopy(self.base['slides'][1]['visual'])
        visual.update(kind='line', minimum=7, maximum=13, major_unit=2)
        self.deck(visual=visual)

    def test_two_series_have_distinct_colors_and_legend(self):
        for kind in ('column', 'bar', 'line'):
            visual = {'type': 'chart', 'kind': kind, 'categories': ['前测', '后测'],
                      'series': [{'name': '实验组', 'values': [12, 18]}, {'name': '对照组', 'values': [12, 14]}], 'unit': '分（合成）'}
            with self.subTest(kind=kind):
                path, pptx, prs = self.deck(visual=visual)
                chart = self.obj(prs, 'visual/chart').chart
                self.assertTrue(chart.has_legend)
                if kind == 'line':
                    colors = [s.format.line.color.rgb for s in chart.series]
                    self.assertNotEqual(chart.series[0].marker.style, chart.series[1].marker.style)
                    chart.series[1].format.line.color.rgb = colors[0]
                else:
                    colors = [s.format.fill.fore_color.rgb for s in chart.series]
                    chart.series[1].format.fill.fore_color.rgb = colors[0]
                self.assertNotEqual(*colors)
                self.rejected(path, pptx, prs, 'series distinction')
        with self.assertRaisesRegex(ValueError, 'distinct series_colors'):
            self.deck(visual=visual, theme={'series_colors': ['176B76', '176b76']})

    def test_small_intervals_and_original_numeric_labels(self):
        visual = {'type': 'interval', 'rows': [{'label': '微小差值', 'estimate': .009, 'low': .004, 'high': .014}],
                  'range': [-.001, .02], 'unit': '差值', 'reference': 0}
        path, pptx, prs = self.deck(4, visual)
        self.assertEqual(self.obj(prs, 'visual/row/0/display').text, '0.009 [0.004, 0.014]')
        self.assertEqual(identity(self.obj(prs, 'visual/row/0/interval'))['low'], .004)
        self.obj(prs, 'visual/row/0/display').text = '0.01 [0.00, 0.01]'
        self.rejected(path, pptx, prs, 'text')
        visual['number_format'] = '.2f'
        with self.assertRaisesRegex(ValueError, 'precision'):
            self.deck(4, visual)
        visual['number_format'] = '.3f'
        self.deck(4, visual)
        visual.pop('number_format')
        visual['rows'][0]['display'] = {'estimate': '0.0090', 'low': '0.004', 'high': '0.014'}
        _, _, prs = self.deck(4, visual)
        self.assertEqual(self.obj(prs, 'visual/row/0/display').text, '0.0090 [0.004, 0.014]')

    def test_nonzero_reference_geometry_and_label_changes(self):
        visual = {'type': 'interval', 'rows': [{'label': '风险比（合成）', 'estimate': 1.2, 'low': .9, 'high': 1.5}],
                  'range': [.5, 2], 'unit': '风险比', 'reference': 1}
        for mutation in ('reference', 'point', 'bound', 'label', 'raw'):
            with self.subTest(mutation=mutation):
                path, pptx, prs = self.deck(4, visual)
                ref = self.obj(prs, 'visual/reference')
                axis = self.obj(prs, 'visual/axis-line')
                self.assertAlmostEqual((ref.left - axis.left) / axis.width, 1 / 3, places=5)
                if mutation == 'reference':
                    ref.left += Inches(.3)
                elif mutation == 'point':
                    self.obj(prs, 'visual/row/0/estimate').left += Inches(.3)
                elif mutation == 'bound':
                    self.obj(prs, 'visual/row/0/interval').width += Inches(.3)
                elif mutation == 'label':
                    self.obj(prs, 'visual/row/0/label').text = '另一指标'
                else:
                    from objects import identify
                    shape = self.obj(prs, 'visual/row/0/estimate')
                    identify(shape, identity(shape)['id'], value=2.1)
                field = {'reference': 'reference x', 'point': 'estimate x', 'bound': 'high x', 'label': 'text', 'raw': 'record.value'}[mutation]
                self.rejected(path, pptx, prs, field)
        visual['scale'] = 'log'
        with self.assertRaisesRegex(ValueError, 'linear'):
            self.deck(4, visual)

    def test_display_precision_near_zero_and_ratio_reference(self):
        for row, reference, bounds in [
            ({'label': 'negative', 'estimate': -.009, 'low': -.014, 'high': -.004}, 0, [-.02,.01]),
            ({'label': 'ratio', 'estimate': 1.009, 'low': 1.004, 'high': 1.014}, 1, [.99,1.02]),
        ]:
            visual = {'type':'interval','rows':[row],'range':bounds,'reference':reference,'unit':'synthetic'}
            self.deck(4, visual)
            visual['number_format'] = '.2f'
            with self.assertRaisesRegex(ValueError, 'precision'):
                self.deck(4, visual)

    def test_process_order_connections_and_quote_attribution(self):
        for mutation in ('node', 'connection', 'arrow'):
            path, pptx, prs = self.deck(2)
            edge = self.obj(prs, 'visual/edge/0')
            if mutation == 'node':
                self.obj(prs, 'visual/node/0').text = '改变步骤'
            elif mutation == 'connection':
                edge.end_connect(self.obj(prs, 'visual/node/2'), 0)
            else:
                edge._element.xpath('.//a:tailEnd')[0].set('type', 'none')
            self.rejected(path, pptx, prs, {'node': 'text', 'connection': 'relation', 'arrow': 'direction'}[mutation])
        for key in ('visual/quote', 'visual/attribution'):
            path, pptx, prs = self.deck(visual={'type': 'quote', 'text': '这是合成引文。', 'attribution': '合成访谈者A'})
            self.obj(prs, key).text = '不同的引文或归属'
            self.rejected(path, pptx, prs, 'text')

    def test_source_image_bytes_and_generation_record(self):
        from PIL import Image
        asset = self.root / 'asset.png'
        Image.new('RGB', (40, 30), 'navy').save(asset)
        visual = {'type': 'image', 'path': 'asset.png', 'caption': '合成图片'}
        for mutation in ('delete', 'bytes', 'record'):
            path, pptx, prs = self.deck(visual=visual)
            shape = self.obj(prs, 'visual/image')
            if mutation == 'delete':
                shape._element.getparent().remove(shape._element)
            elif mutation == 'bytes':
                other = self.root / 'other.png'
                Image.new('RGB', (40, 30), 'orange').save(other)
                from pptx.oxml.ns import qn
                shape.part.related_part(shape._pic.blipFill.blip.get(qn('r:embed')))._blob = other.read_bytes()
            else:
                manifest_path = pptx.parent / 'build.json'
                record = json.loads(manifest_path.read_text(encoding='utf-8'))
                record['objects'] = []
                manifest_path.write_text(json.dumps(record), encoding='utf-8')
            self.rejected(path, pptx, prs, {'delete': 'presence', 'bytes': 'embedded image', 'record': 'generation record'}[mutation])

    def test_preview_freshness_is_separate_from_plan_and_source_checks(self):
        path, pptx, prs = self.deck(3)
        preview = self.root / 'stale-preview'
        preview.mkdir()
        (preview / 'render.json').write_text(json.dumps({'status': 'rendered', 'pptx_sha256': 'stale', 'images': {}}), encoding='utf-8')
        report = check(pptx, path, preview)
        self.assertEqual(report['status'], 'failed')
        self.assertTrue(any('stale' in error for error in report['errors']))
        self.assertEqual(report['plan_consistency'], 'passed')
        self.assertEqual(report['source_fact_check'], 'not_performed')
        self.assertEqual(report['visual_review'], 'not_performed')

    def test_individual_label_text_format_and_visibility(self):
        from chart_labels import apply_spec, label_spec
        for mutation in ('text', 'point_format', 'series_format', 'hidden', 'category'):
            with self.subTest(mutation=mutation):
                visual = copy.deepcopy(self.base['slides'][1]['visual'])
                path, pptx, prs = self.deck(visual=visual)
                chart = self.obj(prs, 'visual/chart').chart
                series = chart.series[0]
                group = series._element.get_or_add_dLbls()
                apply_spec(group, label_spec(visual))
                point = series.points[0].data_label._get_or_add_dLbl()
                apply_spec(point, label_spec(visual), point=True)
                if mutation == 'text':
                    series.points[0].data_label.text_frame.text = '999'
                    self.assertEqual(series.values[0], 12)
                elif mutation == 'point_format':
                    point.xpath('./c:numFmt')[0].set('formatCode', '0%')
                elif mutation == 'series_format':
                    group.get_or_add_numFmt().set('formatCode', '0%')
                elif mutation == 'hidden':
                    apply_spec(point, dict(label_spec(visual), visible=False), point=True)
                else:
                    point.xpath('./c:showCatName')[0].set('val', '1')
                self.rejected(path, pptx, prs, {'text': 'custom label text', 'hidden': 'label visibility',
                                             'category': 'label content'}.get(mutation, 'number format'))

    def test_legal_label_components_percent_and_point_overrides(self):
        from chart_labels import expected_text, label_spec, apply_spec
        for settings, values in [
            ({'number_format': '0.0'}, [12, 8]),
            ({'number_format': '0.0%', 'unit': '原始比例0–1；标签百分比'}, [.12, .08]),
            ({'number_format': '0"%"', 'unit': '百分数'}, [12, 8]),
            ({'number_format': '#,##0.00'}, [1234, 987]),
            ({'label_content': ['category']}, [12, 8]),
            ({'label_content': ['series', 'category', 'value'], 'label_separator': ' / '}, [12, 8]),
            ({'label_overrides': [{'series':0, 'point':0, 'number_format':'0.00', 'content':['category','value'], 'separator':': '}]}, [12, 8]),
            ({'label_overrides': [{'series':0, 'point':1, 'visible':False}]}, [12, 8]),
            ({'labels':False, 'label_overrides': [{'series':0, 'point':0, 'visible':True}]}, [12, 8]),
        ]:
            with self.subTest(settings=settings):
                visual = copy.deepcopy(self.base['slides'][1]['visual'])
                visual.update(settings)
                visual['series'][0]['values'] = values
                path, pptx, prs = self.deck(visual=visual)
                chart = self.obj(prs, 'visual/chart').chart
                apply_spec(chart.series[0]._element.get_or_add_dLbls(), label_spec(visual))
                point = chart.series[0].points[0].data_label
                apply_spec(point._get_or_add_dLbl(), label_spec(visual, 0, 0), point=True)
                point.text_frame.text = expected_text(visual, 0, 0, label_spec(visual, 0, 0))
                prs.save(pptx)
                self.assertEqual(check(pptx, path)['errors'], [])
        # Automatic complex Excel formats can be checked by their declared
        # format; custom literals in those formats must not be guessed.
        visual = copy.deepcopy(self.base['slides'][1]['visual'])
        visual['number_format'] = '[>=10]0.0;0.00'
        path, pptx, prs = self.deck(visual=visual)
        self.obj(prs, 'visual/chart').chart.series[0].points[0].data_label.text_frame.text = '12.0'
        self.rejected(path, pptx, prs, 'unverified.custom label format')

    def test_unsupported_label_sources_are_reported_unverified(self):
        from pptx.oxml.xmlchemy import OxmlElement
        from chart_labels import apply_spec, label_spec
        for mutation in ('linked', 'reference', 'field', 'extension', 'index', 'series_reset'):
            with self.subTest(mutation=mutation):
                visual = copy.deepcopy(self.base['slides'][1]['visual'])
                path, pptx, prs = self.deck(visual=visual)
                series = self.obj(prs, 'visual/chart').chart.series[0]
                apply_spec(series._element.get_or_add_dLbls(), label_spec(visual))
                point = series.points[0].data_label
                node = point._get_or_add_dLbl()
                apply_spec(node, label_spec(visual), point=True)
                if mutation == 'linked':
                    node.xpath('./c:numFmt')[0].set('sourceLinked', '1')
                elif mutation == 'reference':
                    tx = OxmlElement('c:tx'); ref = OxmlElement('c:strRef'); tx.append(ref); node.append(tx)
                elif mutation == 'field':
                    point.text_frame.text = '12'
                    field = OxmlElement('a:fld'); field.set('id', 'dynamic'); field.set('type', 'VALUE')
                    node.xpath('./c:tx/c:rich/a:p')[0].append(field)
                elif mutation == 'extension':
                    node.append(OxmlElement('c:extLst'))
                elif mutation == 'series_reset':
                    group = series._element.get_or_add_dLbls()
                    for element in group.xpath('./c:numFmt|./c:dLbl'):
                        group.remove(element)
                else:
                    node.xpath('./c:idx')[0].set('val', '99')
                self.rejected(path, pptx, prs, 'unverified.')

    def test_interval_rows_keep_text_and_marks_together(self):
        for mutation in ('marks_swap', 'label_only', 'display_only', 'whole_swap', 'whole_move', 'all_move'):
            with self.subTest(mutation=mutation):
                path, pptx, prs = self.deck(4)
                parts = ('label', 'display', 'interval', 'estimate')
                if mutation in ('marks_swap', 'whole_swap'):
                    for part in (parts if mutation == 'whole_swap' else ('interval','estimate')):
                        a = self.obj(prs, 'visual/row/0/' + part)
                        b = self.obj(prs, 'visual/row/1/' + part)
                        a.top, b.top = b.top, a.top
                elif mutation == 'all_move':
                    for shape in prs.slides[0].shapes:
                        if '/visual/' in identity(shape).get('id', ''):
                            shape.top += Inches(.1)
                else:
                    chosen = parts if mutation == 'whole_move' else ('label',) if mutation == 'label_only' else ('display',)
                    for part in chosen:
                        self.obj(prs, 'visual/row/0/' + part).top += Inches(.2)
                if mutation in ('marks_swap','label_only','display_only'):
                    self.rejected(path, pptx, prs, 'row correspondence')
                else:
                    prs.save(pptx)
                    self.assertEqual(check(pptx, path)['errors'], [])

    def test_identity_survives_layout_and_theme_and_defeats_decoy_text(self):
        path, pptx, prs = self.deck(3, theme={'accent': '663399', 'canvas': [.08, .26, .84, .58]})
        table = self.obj(prs, 'visual/table')
        table.name = 'Renamed by author'
        prs.save(pptx)
        self.assertEqual(check(pptx, path)['errors'], [])
        text = table.table.cell(3, 1).text
        table.table.cell(3, 1).text = 'Significant improvement'
        prs.slides[0].shapes.add_textbox(Inches(1), Inches(1), Inches(3), Inches(.5)).text = text
        self.rejected(path, pptx, prs, 'cell[3,1]')


if __name__ == '__main__':
    unittest.main()
