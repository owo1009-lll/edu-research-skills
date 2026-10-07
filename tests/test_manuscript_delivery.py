"""Synthetic fixtures test file plumbing and known errors; never real paper records."""
import copy
import csv
import importlib.util
import json
import hashlib
import shutil
import subprocess
from pathlib import Path
import sys
import tempfile
import unittest

SCRIPTS = Path(__file__).resolve().parents[1] / 'skills/edu-writing/scripts'
sys.path.insert(0, str(SCRIPTS))
import manuscript_package as package
from make_figures import validate_spec, render
from verify_references import compare_identity, resolve


def fixture():
    # Explicitly simulated author, metadata, analysis output and paper. No true DOI claimed.
    metadata = {'doi': 'synthetic-doi-1', 'title': 'Synthetic plumbing fixture',
                'authors': [{'family': 'Fixture', 'given': 'Test'}], 'year': 2020,
                'registered_years': [2020], 'venue': 'SYNTHETIC TEST ONLY', 'volume': '1', 'pages': '1-2'}
    figure = {'kind': 'group_means', 'source_kind': 'existing_analysis',
              'source': 'SYNTHETIC mock analysis output; not real paper data', 'locator': 'SYNTHETIC fixture row1',
              'conclusion': 'SYNTHETIC output retains provided values', 'units': 'SYNTHETIC points', 'limits': [0, 100],
              'panels': [{'title': 'SYNTHETIC output', 'rows': [
                  {'group': 'A', 'condition': 'C1', 'mean': 20, 'n': 10},
                  {'group': 'A', 'condition': 'C2', 'mean': 35, 'n': 10}]}]}
    sections = []
    for role in ['introduction', 'methods', 'results', 'discussion', 'conclusion']:
        paragraph = {'type': 'paragraph', 'id': role, 'text': 'SYNTHETIC software test [@r1].'}
        sections.append({'role': role, 'heading': role.title(), 'blocks': [paragraph]})
    sections[2]['blocks'][0]['text'] += ' See {{table:t1}} and {{figure:f1}}.'
    sections[2]['blocks'] += [
        {'type': 'table', 'id': 't1', 'title': 'Synthetic editable table', 'columns': ['Label', 'Value'],
         'rows': [['A', '20']], 'note': 'SYNTHETIC test [@r1].'},
        {'type': 'figure', 'id': 'f1', 'title': 'Synthetic fixture', 'caption': 'Not real research data [@r1].', 'spec': figure}]
    manuscript = {'scope': 'full_paper', 'research_type': 'quantitative', 'title': 'SYNTHETIC software test',
                  'disclosure': 'All materials, authors, sources and numbers are synthetic test fixtures.',
                  'abstract': 'SYNTHETIC abstract.', 'keywords': ['synthetic'], 'reference_ids': ['r1'],
                  'sections': sections, 'pending': ['Synthetic test, not a deliverable research paper.']}
    verification = {'identity_results': [{'id': 'r1', 'doi': metadata['doi'], 'status': 'verified',
                     'metadata': metadata, 'reading': {'status': 'source_excerpt', 'version': 'synthetic'},
                     'notice_check': {'status': 'synthetic_test'}, 'lookup_url': 'https://example.invalid/synthetic'}],
                    'semantic_review': [{'block_id': key, 'reference_id': 'r1', 'status': 'supported',
                      'locator': 'SYNTHETIC fixture', 'support': 'Synthetic known mechanical-check record.',
                      'reader': 'SYNTHETIC test'} for key in ['introduction', 'methods', 'results', 'discussion', 'conclusion', 't1', 'f1']]}
    return manuscript, verification


class ManuscriptDelivery(unittest.TestCase):
    def test_real_docx_embedding_editable_table_and_exact_plot_data(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as directory:
            m, v = fixture(); out = Path(directory) / 'package'
            result = package.build(m, v, out)
            doc = Document(out / 'manuscript.docx')
            self.assertEqual(len(doc.tables), 1)
            self.assertEqual(doc.tables[0].rows[1].cells[1].text, '20')
            self.assertEqual(len(doc.inline_shapes), 1)
            body = '\n'.join(p.text for p in doc.paragraphs)
            self.assertIn('See Table 1 and Figure 1.', body)
            self.assertNotIn('[@', body); self.assertNotIn('{{', body)
            self.assertEqual(result['table_count'], 1)
            with (out / 'figures/f1/source_data.csv').open(encoding='utf-8-sig') as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual([r['mean'] for r in rows], ['20', '35'])
            svg = (out / 'figures/f1/figure.svg').read_text(encoding='utf-8')
            self.assertIn('<text', svg)
            self.assertTrue((out / 'editable_tables.docx').exists())
            with self.assertRaises(FileExistsError): package.build(m, v, out)

    def test_doi_title_author_and_year_known_errors(self):
        m, v = fixture(); meta = v['identity_results'][0]['metadata']
        ref = {'doi': meta['doi'], 'expected': {'title': meta['title'], 'venue': meta['venue'], 'year': 2020, 'first_author': 'Fixture'}}
        self.assertEqual(compare_identity(ref, meta), [])
        for field, wrong in [('doi', 'different-target'), ('title', 'different paper'), ('registered_years', [2019]), ('authors', [{'family': 'Wrong'}])]:
            changed = copy.deepcopy(meta); changed[field] = wrong
            self.assertTrue(compare_identity(ref, changed), field)

    def test_pending_reading_and_semantic_review_rejected(self):
        m, v = fixture(); v['identity_results'][0]['reading']['status'] = 'full_text_to_be_read_for_case'
        self.assertEqual(package.check(m, v)['status'], 'failed')
        m, v = fixture(); v['identity_results'][0]['notice_check']['status'] = 'pending'
        self.assertEqual(package.check(m, v)['status'], 'failed')

    def test_chinese_source_identities_are_not_empty(self):
        m, v = fixture(); meta = v['identity_results'][0]['metadata']
        meta.update(title='课堂反馈与学习', venue='教育研究', authors=[{'family': '张'}])
        ref = {'doi': meta['doi'], 'expected': {'title': meta['title'], 'venue': meta['venue'], 'year': 2020, 'first_author': '张'}}
        self.assertEqual(compare_identity(ref, meta), [])
        meta['title'] = '教师发展研究'
        self.assertIn('title differs', compare_identity(ref, meta))

    def test_doi_less_primary_record_and_changed_evidence(self):
        m, v = fixture(); meta = v['identity_results'][0]['metadata']; meta['doi'] = ''
        with tempfile.TemporaryDirectory() as directory:
            file = Path(directory) / 'synthetic-evidence.json'
            file.write_text('SYNTHETIC test evidence, not an actual source', encoding='utf-8')
            ref = {'id': 'synthetic', 'doi': '', 'expected': {'title': meta['title'], 'venue': meta['venue'], 'year': 2020, 'first_author': 'Fixture'},
                   'primary_record': {'evidence_path': str(file), 'evidence_sha256': hashlib.sha256(file.read_bytes()).hexdigest(),
                                      'url': 'https://example.invalid/synthetic', 'reader': 'SYNTHETIC fixture',
                                      'locator': 'SYNTHETIC fixture', 'checked_utc': 'SYNTHETIC', 'provider': 'SYNTHETIC', 'metadata': meta}}
            self.assertEqual(resolve(ref, Path(directory))['status'], 'verified')
            file.write_text('SYNTHETIC changed evidence', encoding='utf-8')
            self.assertEqual(resolve(ref, Path(directory))['status'], 'not_found_or_unavailable')

    def test_cli_preserves_installed_tree_and_compact_table_layout(self):
        from docx import Document
        m, v = fixture()
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory); skill = root / 'installed'; skill.mkdir()
            for name in ['make_figures.py', 'manuscript_package.py', 'verify_references.py']:
                shutil.copy2(SCRIPTS / name, skill / name)
            for filename, value in [('m.json', m), ('v.json', v)]:
                (root / filename).write_text(json.dumps(value), encoding='utf-8')
            result = subprocess.run([sys.executable, str(skill / 'manuscript_package.py'), '--manuscript', str(root/'m.json'),
                                     '--verification', str(root/'v.json'), '--output', str(root/'out')], capture_output=True, encoding='utf-8')
            self.assertEqual(result.returncode, 0, result.stderr)
            self.assertFalse((skill / '__pycache__').exists())
            doc = Document(root/'out/manuscript.docx')
            self.assertTrue(doc.tables[0].rows[1].cells[0].paragraphs[0].paragraph_format.keep_with_next)
            self.assertIn('Figure 1. Synthetic fixture. Not real', '\n'.join(p.text for p in doc.paragraphs))

    def test_provided_intervals_export_without_estimation(self):
        m, v = fixture(); spec = m['sections'][2]['blocks'][2]['spec']
        spec.update(kind='estimate_intervals', rows=[{'label': 'SYNTHETIC estimate', 'estimate': 3, 'low': 2, 'high': 4}],
                    interval_definition='SYNTHETIC supplied interval, not calculated', reference_line=0)
        with tempfile.TemporaryDirectory() as directory:
            render(spec, Path(directory)/'figure')
            with (Path(directory)/'figure/source_data.csv').open(encoding='utf-8-sig') as stream:
                row = next(csv.DictReader(stream))
            self.assertEqual([row[k] for k in ['low','estimate','high']], ['2','3','4'])
        m, v = fixture(); v['semantic_review'][0]['status'] = 'pending'
        self.assertEqual(package.check(m, v)['status'], 'failed')

    def test_unknown_and_uncited_references_rejected(self):
        m, v = fixture(); m['sections'][0]['blocks'][0]['text'] += ' [@ghost]'
        self.assertTrue(any('absent from bibliography' in e for e in package.check(m, v)['errors']))
        m, v = fixture(); second = copy.deepcopy(v['identity_results'][0]); second.update(id='r2', doi='synthetic-doi-2'); v['identity_results'].append(second); m['reference_ids'].append('r2')
        self.assertTrue(any('Uncited' in e for e in package.check(m, v)['errors']))

    def test_display_numbering_and_known_wrong_reference(self):
        m, v = fixture(); m['sections'][2]['blocks'][0]['text'] = 'SYNTHETIC [@r1] {{figure:missing}}'
        errors = package.check(m, v)['errors']
        self.assertTrue(any('Unknown display reference' in e for e in errors))
        self.assertTrue(any('not cited in body' in e for e in errors))

    def test_same_author_year_disambiguation(self):
        m, v = fixture(); a = v['identity_results'][0]; b = copy.deepcopy(a); b['id'] = 'r2'; b['metadata']['title'] = 'Another synthetic title'
        labels = package.citation_labels([a, b])
        self.assertEqual(labels['r2']['year'], '2020a'); self.assertEqual(labels['r1']['year'], '2020b')
        self.assertIn('2020a', package.bibliography(b, labels['r2']))

    def test_no_fabricated_trajectory_or_interval(self):
        m, v = fixture(); spec = m['sections'][2]['blocks'][2]['spec']
        spec['kind'] = 'individual_trajectory'
        with self.assertRaises(ValueError): validate_spec(spec)
        spec.update(kind='estimate_intervals', rows=[{'label': 'synthetic', 'estimate': 3, 'low': 4, 'high': 5}], interval_definition='synthetic')
        with self.assertRaises(ValueError): validate_spec(spec)

    def test_known_correction_and_full_structure_gates(self):
        m, v = fixture(); v['identity_results'][0]['notice_check']['required_reference_ids'] = ['correction']
        self.assertTrue(any('Known correction' in e for e in package.check(m, v)['errors']))
        m, v = fixture(); m['sections'] = [s for s in m['sections'] if s['role'] != 'methods']
        self.assertTrue(any('Incomplete full-paper' in e for e in package.check(m, v)['errors']))

    def test_theory_structure_does_not_require_fake_methods(self):
        m, v = fixture(); m['research_type'] = 'theoretical'
        m['sections'] = [s for s in m['sections'] if s['role'] in ('introduction', 'discussion', 'conclusion')]
        self.assertEqual(package.check(m, v)['status'], 'passed')

    def test_conceptual_and_theme_diagram_status_and_exports(self):
        for kind, status in [('flow', 'procedure'), ('framework', 'hypothesized'), ('themes', 'candidate_themes')]:
            spec = {'kind': kind, 'evidence_status': status, 'source': 'SYNTHETIC diagram fixture',
                    'locator': 'SYNTHETIC nodes', 'conclusion': 'SYNTHETIC relation test',
                    'nodes': [{'id': 'a', 'label': 'Synthetic A'}, {'id': 'b', 'label': 'Synthetic B'}],
                    'edges': [{'from': 'a', 'to': 'b', 'label': 'synthetic link'}]}
            with tempfile.TemporaryDirectory() as directory:
                result = render(spec, Path(directory) / 'figure')
                self.assertFalse(result['extra_analysis'])
                self.assertTrue((Path(directory) / 'figure/figure.svg').is_file())
        spec['evidence_status'] = 'tested_causal_effect'
        with self.assertRaises(ValueError): validate_spec(spec)


if __name__ == '__main__': unittest.main()


class ChineseJournalLayout(unittest.TestCase):
    def test_gbt7714_numbering_superscript_and_three_line_table(self):
        from docx import Document
        m, v = fixture()
        zh_meta = {'title': '生成式人工智能与学习投入', 'authors': [{'family': '张', 'given': '三'}, {'family': '李', 'given': '四'},
                   {'family': '王', 'given': '五'}, {'family': '赵', 'given': '六'}], 'year': 2025, 'venue': '电化教育研究',
                   'volume': '46', 'issue': '8', 'pages': '45-52', 'doi': 'synthetic-doi-2', 'registered_years': [2025]}
        v['identity_results'].append({'id': 'r2', 'doi': 'synthetic-doi-2', 'status': 'verified', 'metadata': zh_meta,
                                      'reading': {'status': 'source_excerpt', 'version': 'synthetic'}, 'notice_check': {'status': 'synthetic_test'}, 'lookup_url': 'https://example.invalid/synthetic2'})
        m['language'] = 'zh'; m['reference_ids'] = ['r1', 'r2']
        m['sections'][0]['blocks'][0]['text'] = '合成测试段落[@r2][@r1]。'
        v['semantic_review'].append({'block_id': 'introduction', 'reference_id': 'r2', 'status': 'supported',
                                     'locator': 'SYNTHETIC', 'support': 'Synthetic.', 'reader': 'SYNTHETIC test'})
        with tempfile.TemporaryDirectory() as directory:
            out = Path(directory) / 'package'
            package.build(m, v, out)
            md = (out / 'manuscript.md').read_text(encoding='utf-8')
            self.assertIn('合成测试段落[1,2]。', md)                      # numbered by first citation, adjacent merged
            self.assertIn('## 参考文献', md); self.assertIn('## 摘要', md); self.assertIn('表1', md)
            self.assertIn('[1] 张三, 李四, 王五, 等. 生成式人工智能与学习投入[J]. 电化教育研究, 2025, 46(8): 45-52.', md)
            self.assertIn('[2] Fixture T. Synthetic plumbing fixture[J]. SYNTHETIC TEST ONLY, 2020, 1: 1-2.', md)
            doc = Document(out / 'manuscript.docx')
            runs = [r for p in doc.paragraphs for r in p.runs if r.text == '[1,2]']
            self.assertTrue(runs and all(r.font.superscript for r in runs))
            borders = doc.tables[0].rows[0].cells[0]._tc.xml
            self.assertIn('w:top w:val="single"', borders); self.assertIn('w:left w:val="nil"', borders)
