"""Observable PPTX, notes, template and evidence-boundary behavior; no network/Office required."""
from pathlib import Path
import copy
import importlib.util
import json
import sys
import tempfile
import unittest
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / 'skills/edu-ppt/scripts'
sys.path.insert(0, str(SCRIPTS))
from build_deck import build
from plan import load_plan, validate, digest
from check_deck import check
from extract_source import extract
from pptx import Presentation
from pptx.util import Inches


class PptTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='edu-ppt-test-')
        self.root = Path(self.tmp.name)
        self.plan = ROOT / 'skills/edu-ppt/examples/progress-plan.json'
        self.value = load_plan(self.plan)

    def tearDown(self):
        self.tmp.cleanup()

    def write_plan(self, value, name='plan.json'):
        path = self.root / name
        path.write_text(json.dumps(value, ensure_ascii=False), encoding='utf-8')
        return path

    def test_content_mode_delivers_actual_prose_and_notes(self):
        out = self.root / 'content'
        build(self.plan, out, content_only=True)
        self.assertFalse((out / 'slides.pptx').exists())
        text = (out / 'content.md').read_text(encoding='utf-8')
        notes = (out / 'speaker_notes.md').read_text(encoding='utf-8')
        for slide in self.value['slides']:
            self.assertIn(slide['message'], text)
            self.assertIn(slide['notes'], notes)

    def test_native_objects_and_notes_survive_editing(self):
        out = self.root / 'deck'
        build(self.plan, out)
        pptx = out / 'slides.pptx'
        self.assertEqual(check(pptx, self.plan)['errors'], [])
        prs = Presentation(pptx)
        chart = next(s.chart for s in prs.slides[1].shapes if s.has_chart)
        self.assertEqual(list(chart.series[0].values), [12.0, 8.0])
        table = next(s.table for s in prs.slides[3].shapes if s.has_table)
        table.cell(1, 1).text = 'Edited status'
        prs.slides[0].shapes[0].text = 'Edited title'
        changed = self.root / 'edited.pptx'
        prs.save(changed)
        reread = Presentation(changed)
        self.assertEqual(reread.slides[0].shapes[0].text, 'Edited title')
        self.assertEqual(next(s.table for s in reread.slides[3].shapes if s.has_table).cell(1, 1).text, 'Edited status')
        self.assertTrue(reread.slides[0].notes_slide.notes_text_frame.text.startswith(self.value['slides'][0]['notes']))
        with ZipFile(changed) as archive:
            self.assertTrue(any(n.startswith('ppt/embeddings/') and n.endswith('.xlsx') for n in archive.namelist()))
        self.assertTrue(any('planned text missing' in e for e in check(changed, self.plan)['errors']))

    def test_template_geometry_preserved_without_sample_pages(self):
        template = self.root / 'template.pptx'
        prs = Presentation(); prs.slide_width = Inches(12); prs.slide_height = Inches(8)
        sample = prs.slides.add_slide(prs.slide_layouts[6])
        sample.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = 'Template sample only'
        prs.save(template)
        before = digest(template)
        out = self.root / 'template-out'
        build(self.plan, out, template=template)
        result = Presentation(out / 'slides.pptx')
        self.assertEqual((result.slide_width, result.slide_height), (prs.slide_width, prs.slide_height))
        self.assertEqual(len(result.slides), len(self.value['slides']))
        self.assertEqual(digest(template), before)
        self.assertNotIn('Template sample only', '\n'.join(s.text for slide in result.slides for s in slide.shapes if s.has_text_frame))

    def test_style_rebuild_retains_content_and_input(self):
        first, second = self.root / 'first', self.root / 'second'
        theme = self.root / 'theme.json'
        theme.write_text('{"accent":"663399"}', encoding='utf-8')
        one = build(self.plan, first)
        two = build(self.plan, second, theme_path=theme)
        self.assertEqual(one['plan_sha256'], two['plan_sha256'])
        self.assertEqual((first / 'content.md').read_bytes(), (second / 'content.md').read_bytes())
        self.assertEqual(check(second / 'slides.pptx', self.plan)['errors'], [])

    def test_invalid_evidence_and_intervals_are_rejected(self):
        for mutation in ('source', 'scope', 'interval'):
            plan = copy.deepcopy(self.value)
            if mutation == 'source':
                plan['slides'][0]['evidence'][0]['source'] = 'unknown'
            elif mutation == 'scope':
                plan['sources'][0]['access'] = 'title'
            else:
                plan['slides'][-1]['visual']['rows'][0]['low'] = .4
            with self.subTest(mutation=mutation), self.assertRaises(ValueError):
                validate(plan)

    def test_density_and_overwrite_fail_explicitly(self):
        plan = copy.deepcopy(self.value)
        plan['slides'][0]['body'] = ['Long text ' * 400]
        out = self.root / 'dense'
        with self.assertRaisesRegex(ValueError, 'dense'):
            build(self.write_plan(plan), out)
        self.assertFalse((out / 'slides.pptx').exists())
        with self.assertRaisesRegex(ValueError, 'not empty'):
            build(self.plan, out)

    def test_notes_tampering_detected(self):
        out = self.root / 'notes'
        build(self.plan, out)
        path = out / 'slides.pptx'
        prs = Presentation(path)
        prs.slides[0].notes_slide.notes_text_frame.text = 'Different speech'
        prs.save(path)
        self.assertTrue(any('notes differ' in e for e in check(path, self.plan)['errors']))

    def test_multiline_body_keeps_line_breaks_and_room_for_next_item(self):
        plan = copy.deepcopy(self.value)
        plan['slides'][0]['body'] = ['Praderio Gaias 等（2024）\nFrontiers in Education, 9:1472721', '下一项正文']
        path = self.write_plan(plan)
        out = self.root / 'multiline'
        build(path, out)
        prs = Presentation(out / 'slides.pptx')
        body = [s for s in prs.slides[0].shapes if s.name == 'edu:body']
        self.assertEqual(body[0].text, plan['slides'][0]['body'][0])
        self.assertGreater(body[1].top, body[0].top + body[0].height)
        self.assertEqual(check(out / 'slides.pptx', path)['errors'], [])

    def test_content_only_table_is_readable_without_json_or_empty_body(self):
        plan = copy.deepcopy(self.value)
        plan['slides'] = [plan['slides'][3]]
        plan['slides'][0]['body'] = []
        plan['slides'][0]['visual']['rows'][0][1] = '已完成\nA | B'
        out = self.root / 'readable-table'
        build(self.write_plan(plan), out, content_only=True)
        content = (out / 'content.md').read_text(encoding='utf-8')
        self.assertIn('已完成<br>A \\| B', content)
        self.assertNotIn('```json', content)
        self.assertNotIn('上屏正文：', content)
        self.assertNotIn('身份：', content)
        self.assertIn(plan['slides'][0]['notes'], content)

    def test_docx_extraction_retains_paragraph_and_table_locators(self):
        from docx import Document
        document = Document(); document.add_paragraph('A result with a minus: -0.20')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = 'Estimate'; table.cell(1, 0).text = '-0.20'
        source = self.root / 'source.docx'; document.save(source)
        extract(source, self.root / 'extracted')
        data = json.loads((self.root / 'extracted/source.json').read_text(encoding='utf-8'))
        self.assertEqual(data['records'][0]['location'], 'P0001')
        self.assertEqual(data['records'][1]['location'], 'Table1')
        self.assertIn('-0.20', data['records'][1]['text'])
        self.assertEqual(data['reading_status'], 'extracted_not_yet_verified')


if __name__ == '__main__':
    unittest.main()
