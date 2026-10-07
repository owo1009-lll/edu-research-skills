"""Regression tests for problems found in the 2026-10-07 usability run."""
from pathlib import Path
import csv
import json
import shutil
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
SKILLS = ROOT / 'skills'


def run(args, cwd):
    return subprocess.run([sys.executable, '-X', 'utf8', *map(str, args)], cwd=cwd, capture_output=True, text=True, encoding='utf-8')


class Folder(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()


class GroundedCoding(Folder):
    def test_commit_tables_and_saturation(self):
        private = self.dir / 'edu_output/private'; private.mkdir(parents=True)
        shutil.copy(SKILLS / 'edu-qual/examples/synthetic_interviews.csv', private)
        (private / 'r.json').write_text('{}', encoding='utf-8')  # nothing to redact
        p = run([SKILLS / 'edu-qual/scripts/theme_analysis.py', 'prepare', '--input', 'edu_output/private/synthetic_interviews.csv',
                 '--redactions', 'edu_output/private/r.json', '--output', 'edu_output/private/prep'], self.dir)
        self.assertEqual(p.returncode, 0, p.stderr)
        seg = {s['document_id'] + s['text'][:6]: s for s in json.loads((private / 'prep/segments.json').read_text(encoding='utf-8'))}
        t1 = next(s for k, s in seg.items() if k.startswith('T01以前')); t2 = next(s for k, s in seg.items() if k.startswith('T02我主要'))
        rows = [[t1['id'], '先让AI列一个提纲', 'a1 AI起草', 'A1 人机分工', '备课方式重构', 'AI_assistant'],
                [t2['id'], '我挑两三个最适合学生的', 'a2 教师筛选', 'A1 人机分工', '备课方式重构', 'AI_assistant']]
        with open(private / 'coding.csv', 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f); w.writerow(['segment_id', 'quote', 'concept', 'category', 'main_category', 'coder']); w.writerows(rows)
        script = SKILLS / 'edu-qual/scripts/grounded_coding.py'
        p = run([script, '--prepared', 'edu_output/private/prep', '--coding', 'edu_output/private/coding.csv', '--actor', 'AI_assistant',
                 '--order', 'T01,T02', '--holdout', 'T02'], self.dir)
        self.assertEqual(p.returncode, 0, p.stderr)
        out = Path(json.loads(p.stdout)['revision'])
        self.assertIn('未达到饱和', json.loads(p.stdout)['saturation'])  # the held-out interview adds a concept
        self.assertTrue((out / 'open_coding.csv').is_file() and (out / 'axial_coding.csv').is_file())
        rows[1][1] = '这句话不在原文里'
        with open(private / 'coding.csv', 'w', encoding='utf-8-sig', newline='') as f:
            w = csv.writer(f); w.writerow(['segment_id', 'quote', 'concept', 'category', 'main_category', 'coder']); w.writerows(rows)
        p = run([script, '--prepared', 'edu_output/private/prep', '--coding', 'edu_output/private/coding.csv', '--actor', 'AI_assistant',
                 '--parent', out.name], self.dir)
        self.assertNotEqual(p.returncode, 0); self.assertIn('exact excerpt', p.stderr)


class WordAndChecks(Folder):
    def test_md_to_docx_three_line_table_and_superscript(self):
        from docx import Document
        (self.dir / 'r.md').write_text('# 研究结果\n\n表1 描述统计\n\n| 变量 | M |\n|---|---|\n| 教师支持 | 3.52 |\n\n已有研究表明[1,2]，**结果**稳定。\n', encoding='utf-8')
        p = run([SKILLS / 'edu-writing/scripts/md_to_docx.py', '--input', 'r.md', '--output', 'r.docx', '--superscript-citations'], self.dir)
        self.assertEqual(p.returncode, 0, p.stderr)
        doc = Document(self.dir / 'r.docx')
        self.assertEqual(doc.tables[0].rows[1].cells[0].text, '教师支持')
        self.assertIn('w:left w:val="nil"', doc.tables[0].rows[0].cells[0]._tc.xml)
        self.assertTrue(any(r.font.superscript and r.text == '[1,2]' for p in doc.paragraphs for r in p.runs))

    def test_proposal_official_headings_and_section_limits(self):
        body = ['1. [选题说明]', '甲' * 320, '2. [选题依据]', '（一）国外研究', '乙' * 50, '6. [研究基础]', '本研究首创了新模式。']
        (self.dir / 'h.md').write_text('\n'.join(body), encoding='utf-8')
        p = run([SKILLS / 'edu-proposal/scripts/check_proposal.py', '--draft', 'h.md', '--preset', 'qgjk'], self.dir)
        self.assertIn('选题说明: 320 / 300 (OVER)', p.stdout)
        self.assertIn('选题依据: 57', p.stdout)  # sub-heading （一）国外研究 (7) counted toward its parent section
        self.assertIn('[首创]', p.stdout)

    def test_rounded_values_count_as_sourced(self):
        (self.dir / 'src.csv').write_text('est,share\n0.4183,0.4271\n', encoding='utf-8')
        (self.dir / 'd.md').write_text('# 讨论\n\n路径系数为0.42，间接效应占42.7%，另有0.77。\n', encoding='utf-8')
        p = run([SKILLS / 'edu-writing/scripts/check_voice.py', '--draft', 'd.md', '--source', 'src.csv'], self.dir)
        self.assertIn('invented details are not): 1', p.stdout)
        self.assertIn('0.77', p.stdout)


if __name__ == '__main__':
    unittest.main()
