"""Cross-format changes, reasonable differences and uncertain correspondences."""
from copy import deepcopy
import json
from pathlib import Path
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills/edu-research/scripts"))
from project_check import register, check, refresh, synchronize
from project_documents import read_document
from private_store import digest


def identity(**values):
    result = dict(outcome="engagement", comparison="treatment - control", timepoint="post",
                  sample="complete cases", statistic="beta", unit="standardized", model="adjusted",
                  adjustments="baseline", analysis="primary", uncertainty="95% CI")
    result.update(values)
    return result


def bind(key, file, locator, **extra):
    return dict(id=key, file=file, locator=locator, reason="Synthetic materials explicitly identify this analysis.", **extra)


class ProjectCheck(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.results = self.root / "results.json"
        self.paper = self.root / "paper.md"
        self.results.write_text(json.dumps({"beta": 0.21, "n": 390, "model": "adjusted"}), encoding="utf-8")
        self.paper.write_text("# Results\nAdjusted engagement beta = .210.\nAnalysis n = 390.\nThis warrants a qualified judgment.\n", encoding="utf-8")
        self.spec = {
            "files": [{"id": "results", "path": str(self.results), "role": "analysis_output"},
                      {"id": "paper", "path": str(self.paper), "role": "manuscript"}],
            "facts": [
                {"id": "beta", "identity": identity(), "mapping_reason": "Same adjusted primary analysis.",
                 "authority": {"source": "beta-source", "kind": "user_confirmed", "reason": "Synthetic test request designates results.json."},
                 "sources": [bind("beta-source", "results", "json:/beta",
                                  guards=[{"locator": "json:/model", "equals": "adjusted"}])],
                 "occurrences": [bind("beta-paper", "paper", "line:2", pattern=r"beta = (?P<value>[.\d-]+)\.",
                                      display_decimals=3)]},
                {"id": "n", "identity": identity(statistic="n", unit="people", uncertainty="not applicable"),
                 "mapping_reason": "Explicit analysis sample, distinct from total enrollment.",
                 "authority": {"source": "n-source", "kind": "user_confirmed", "reason": "New analysis sample is authoritative."},
                 "sources": [bind("n-source", "results", "json:/n")],
                 "occurrences": [bind("n-paper", "paper", "line:3", pattern=r"n = (?P<value>\d+)")]},
            ],
            "dependencies": [{"id": "interpretation", "depends_on": ["beta"],
                              "target": bind("judgment", "paper", "line:4", value_type="text"),
                              "reason": "Interpretation relies on direction and uncertainty of adjusted beta."}]
        }

    def update(self, **changes):
        values = json.loads(self.results.read_text())
        values.update(changes)
        self.results.write_text(json.dumps(values), encoding="utf-8")

    def test_baseline_new_results_and_no_unauthorized_writes(self):
        original = self.paper.read_bytes()
        first = register(self.root, self.spec)
        self.assertFalse([x for x in first["findings"] if x["status"] == "stale"])
        self.update(beta=-0.08, n=382)
        report = check(self.root)
        stale = [x for x in report["findings"] if x["status"] == "stale"]
        self.assertEqual({x["fact_id"] for x in stale}, {"beta", "n"})
        self.assertTrue(any(x.get("dependency_id") == "interpretation" and x["status"] == "review" for x in report["findings"]))
        self.assertEqual(self.paper.read_bytes(), original)
        report_file = self.root / "edu_output/project/check.json"
        before = report_file.stat().st_mtime_ns
        self.assertEqual(report, check(self.root))
        self.assertEqual(before, report_file.stat().st_mtime_ns)

    def test_different_statistics_samples_and_models_remain_separate(self):
        values = {"beta": .21, "n": 390, "model": "adjusted", "r": .21, "total": 410,
                  "unadjusted": .27, "sensitivity": .18}
        self.results.write_text(json.dumps(values), encoding="utf-8")
        self.paper.write_text(self.paper.read_text() + "r = .21; total N = 410; unadjusted = .27; sensitivity = .18.\n", encoding="utf-8")
        for key, stat, sample, model, analysis in [
            ("r", "r", "complete cases", "bivariate", "primary"),
            ("total", "n", "all enrolled", "not applicable", "recruitment"),
            ("unadjusted", "beta", "complete cases", "unadjusted", "primary"),
            ("sensitivity", "beta", "imputed sample", "adjusted", "sensitivity")]:
            fact = {"id": key, "identity": identity(statistic=stat, sample=sample, model=model, analysis=analysis),
                    "mapping_reason": "Separate explicitly labeled analysis.",
                    "authority": {"source": key + "-src", "kind": "user_confirmed", "reason": "Designated synthetic results."},
                    "sources": [bind(key + "-src", "results", "json:/" + key)],
                    "occurrences": [bind(key + "-occ", "paper", "line:5",
                                        pattern=(r"total N" if key == "total" else key) + r" = (?P<value>\d+(?:\.\d+)?|\.\d+)")] }
            self.spec["facts"].append(fact)
        register(self.root, self.spec)
        self.update(beta=.14, n=382)
        report = check(self.root)
        self.assertEqual({r["fact_id"] for r in report["findings"] if r["status"] == "stale"}, {"beta", "n"})
        self.assertTrue(all(r["status"] == "consistent" for r in report["findings"]
                            if r.get("fact_id") in {"r", "total", "unadjusted", "sensitivity"}))

    def test_rounding_units_and_interval_reference_boundary(self):
        self.update(beta=.214)
        fact = self.spec["facts"][0]
        fact["occurrences"][0]["display_decimals"] = 2
        self.paper.write_text(self.paper.read_text().replace(".210", ".21"), encoding="utf-8")
        report = register(self.root, self.spec)
        self.assertEqual(next(x for x in report["findings"] if x.get("fact_id") == "beta")["status"], "consistent")
        self.update(beta=.004)
        text = self.paper.read_text().replace(".21", ".00")
        self.paper.write_text(text, encoding="utf-8")
        fact["occurrences"][0]["reviewed_sha256"] = digest(self.paper.read_bytes())
        report = refresh(self.root, {"reader": "test reader", "reason": "Read changed manuscript.",
                                     "replace": {"facts": self.spec["facts"]}})
        self.assertEqual(next(x for x in report["findings"] if x.get("fact_id") == "beta")["status"], "review")

    def test_authority_unclear_and_context_changed_are_not_guessed(self):
        self.spec["facts"][1].pop("authority")
        register(self.root, self.spec)
        self.update(model="unadjusted")
        report = check(self.root)
        self.assertTrue(all(x["status"] == "review" for x in report["findings"] if x.get("fact_id") in {"n", "beta"}))
        self.assertFalse(any(x["status"] == "stale" for x in report["findings"]))

    def test_rename_relocate_modified_anchor_and_missing_file(self):
        register(self.root, self.spec)
        moved = self.root / "renamed.md"
        self.paper.rename(moved)
        report = check(self.root, {"paper": str(moved)})
        self.assertEqual(report["files"]["paper"]["baseline_state"], "unchanged")
        moved.write_text("New title\n" + moved.read_text(), encoding="utf-8")
        self.update(n=382)
        report = check(self.root)
        n = next(x for x in report["findings"] if x.get("fact_id") == "n")
        self.assertEqual(n["location"], "line:4")
        self.assertEqual(n["status"], "stale")
        moved.write_text(moved.read_text().replace("Analysis n", "Enrollment n"), encoding="utf-8")
        report = check(self.root)
        self.assertEqual(next(x for x in report["findings"] if x.get("fact_id") == "n")["status"], "review")
        moved.unlink()
        report = check(self.root)
        self.assertTrue(any(x["status"] == "unchecked" for x in report["findings"]))

    def test_claim_expansion_invalidates_support(self):
        source, ledger = self.root / "source.md", self.root / "ledger.json"
        source.write_text("The source describes engagement in one course.", encoding="utf-8")
        ledger.write_text('{"claims":[{"block_id":"C1","reference_id":"R1","status":"supported"}]}', encoding="utf-8")
        self.spec["files"] += [{"id": "source", "path": str(source), "role": "reference"},
                               {"id": "ledger", "path": str(ledger), "role": "citation_ledger"}]
        self.spec["citations"] = [{"id": "C1", "reference_id": "R1", "reader": "synthetic test reader",
                                   "source_file": "source", "ledger_file": "ledger",
                                   "claim": bind("claim1", "paper", "line:4"),
                                   "support": bind("support1", "source", "line:1")}]
        register(self.root, self.spec)
        self.paper.write_text(self.paper.read_text().replace("qualified judgment", "universal causal conclusion"),
                              encoding="utf-8")
        report = check(self.root)
        self.assertEqual(next(x for x in report["findings"] if x.get("citation_id") == "C1")["support_status"], "pending")

    def test_authorized_sync_preserves_original_and_refuses_stale_report(self):
        register(self.root, self.spec)
        self.update(n=382)
        report = check(self.root)
        before = self.paper.read_bytes()
        request = {"authorization": "Synchronize the reported analysis sample.",
                   "report_fingerprint": report["input_fingerprint"],
                   "edits": [{"file": "paper", "output": str(self.root / "updated.md"),
                              "replacements": [{"locator": "line:3", "before": "Analysis n = 390.",
                                                "old": "390", "new": "382", "reason": "Same confirmed analysis sample."}]}]}
        synchronize(self.root, request)
        self.assertIn("n = 382", (self.root / "updated.md").read_text())
        self.assertEqual(before, self.paper.read_bytes())
        self.update(n=380)
        request["edits"][0]["output"] = str(self.root / "another.md")
        with self.assertRaises(ValueError):
            synchronize(self.root, request)

    def test_first_use_mismatch_also_flags_dependent_interpretation(self):
        self.update(beta=-.08)
        report = register(self.root, self.spec)
        self.assertTrue(any(x.get("dependency_id") == "interpretation" and x["status"] == "review"
                            for x in report["findings"]))
        self.assertFalse(any(x.get("source_changed_since_baseline") for x in report["findings"]))

    def test_wrong_authorized_replacement_is_rejected_before_writing(self):
        register(self.root, self.spec)
        self.update(n=382)
        report = check(self.root)
        request = {"authorization": "Sync sample only.", "report_fingerprint": report["input_fingerprint"],
                   "edits": [{"file": "paper", "output": str(self.root / "wrong.md"),
                              "replacements": [{"locator": "line:3", "before": "Analysis n = 390.",
                                                "old": "390", "new": "999", "reason": "Erroneous proposed replacement."}]}]}
        with self.assertRaises(ValueError):
            synchronize(self.root, request)
        self.assertFalse((self.root / "wrong.md").exists())

    def test_mixed_certain_and_uncertain_same_paragraph_cannot_be_bulk_synced(self):
        self.paper.write_text(self.paper.read_text().replace("Analysis n = 390.", "Analysis n = 390; subgroup n = 390."),
                              encoding="utf-8")
        uncertain = deepcopy(self.spec["facts"][1])
        uncertain.update(id="subgroup", identity=identity(statistic="n", sample="unclear subgroup"))
        uncertain.pop("authority")
        uncertain["sources"][0]["id"] = "unclear-source"
        uncertain["occurrences"][0].update(id="unclear-target", pattern=r"subgroup n = (?P<value>\d+)")
        self.spec["facts"].append(uncertain)
        register(self.root, self.spec)
        self.update(n=382)
        report = check(self.root)
        before = "Analysis n = 390; subgroup n = 390."
        request = {"authorization": "Sync confirmed values; don't decide the subgroup.",
                   "report_fingerprint": report["input_fingerprint"],
                   "edits": [{"file": "paper", "output": str(self.root / "unsafe.md"),
                              "replacements": [{"locator": "line:3", "before": before, "old": before,
                                                "new": before.replace("390", "382"), "reason": "Erroneous bulk proposal."}]}]}
        with self.assertRaisesRegex(ValueError, "unresolved"):
            synchronize(self.root, request)
        self.assertFalse((self.root / "unsafe.md").exists())

    def test_percent_conversion_requires_explicit_same_quantity(self):
        self.results.write_text('{"beta":0.21,"n":390,"model":"adjusted"}', encoding="utf-8")
        self.paper.write_text(self.paper.read_text().replace("beta = .210", "proportion = 21%"), encoding="utf-8")
        fact = self.spec["facts"][0]
        fact["identity"] = identity(statistic="proportion", unit="fraction")
        target = fact["occurrences"][0]
        target.update(pattern=r"proportion = (?P<value>\d+)%", scale=.01)
        with self.assertRaises(ValueError):
            register(self.root, self.spec)
        target["conversion_reason"] = "The same proportion is displayed as a percentage."
        report = register(self.root, self.spec)
        self.assertEqual(next(x for x in report["findings"] if x.get("fact_id") == "beta")["status"], "consistent")

    def test_interval_only_change_and_nonzero_reference(self):
        self.results.write_text('{"beta":1.2,"n":390,"model":"adjusted","low":1.004,"high":1.4}', encoding="utf-8")
        self.paper.write_text(self.paper.read_text().replace("beta = .210.", "ratio = 1.20 [1.00, 1.40]."),
                              encoding="utf-8")
        fact = self.spec["facts"][0]
        fact["identity"] = identity(statistic="ratio", unit="ratio")
        fact["sources"] = [bind("ratio-source", "results", "json:/low", reference_values=[1])]
        fact["authority"]["source"] = "ratio-source"
        fact["occurrences"] = [bind("ratio-bound", "paper", "line:2",
                                   pattern=r"\[(?P<value>[\d.]+),", display_decimals=2)]
        report = register(self.root, self.spec)
        row = next(x for x in report["findings"] if x.get("fact_id") == "beta")
        self.assertEqual(row["status"], "review")
        self.assertIn("reference boundary", row["reason"])
        self.update(low=.98)
        report = check(self.root)
        self.assertTrue(any(x.get("dependency_id") == "interpretation" for x in report["findings"]))

    def test_unrelated_citation_entry_does_not_invalidate_read_support(self):
        source, ledger = self.root / "source.md", self.root / "ledger.json"
        source.write_text("One course supports a qualified judgment.", encoding="utf-8")
        records = {"claims": [{"block_id": "C1", "reference_id": "R1", "status": "supported",
                               "reader": "synthetic reader", "support": "Bounded to one course."}]}
        ledger.write_text(json.dumps(records), encoding="utf-8")
        self.spec["files"] += [{"id": "source", "path": str(source), "role": "reference"},
                               {"id": "ledger", "path": str(ledger), "role": "citation_ledger"}]
        self.spec["citations"] = [{"id": "C1", "reference_id": "R1", "reader": "synthetic reader",
                                   "source_file": "source", "ledger_file": "ledger",
                                   "claim": bind("claim1", "paper", "line:4"),
                                   "support": bind("support1", "source", "line:1")}]
        register(self.root, self.spec)
        records["claims"].append({"block_id": "C2", "reference_id": "R2", "status": "pending"})
        ledger.write_text(json.dumps(records), encoding="utf-8")
        report = check(self.root)
        self.assertEqual(next(x for x in report["findings"] if x.get("citation_id") == "C1")["status"], "consistent")

    def test_word_tables_csv_xlsx_pdf_and_parse_failure_are_visible(self):
        from docx import Document
        from openpyxl import Workbook
        import fitz
        word = self.root / "paper.docx"
        doc = Document()
        doc.add_paragraph("Analysis n = 390.")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text, table.cell(0, 1).text = "beta", ".21"
        doc.save(word)
        self.assertTrue(any(u["kind"] == "table_paragraph" for u in read_document(word)["units"]))
        csv = self.root / "results.csv"
        csv.write_text("outcome,beta\nengagement,.21\n", encoding="utf-8")
        self.assertEqual(read_document(csv)["units"][-1]["text"], ".21")
        xlsx = self.root / "results.xlsx"
        book = Workbook()
        book.active.append(["beta", .21, "=B1*100"])
        book.save(xlsx)
        self.assertTrue(read_document(xlsx)["unchecked"])
        pdf = self.root / "scan.pdf"
        p = fitz.open()
        p.new_page()
        p.save(pdf)
        p.close()
        self.assertTrue(read_document(pdf)["unchecked"])
        readable = self.root / "readable.pdf"
        p = fitz.open()
        p.new_page().insert_text((50, 50), "Analysis sample: 382")
        p.save(readable)
        p.close()
        self.assertTrue(any("382" in u["text"] for u in read_document(readable)["units"]))
        broken = self.root / "broken.docx"
        broken.write_text("not a docx")
        self.assertEqual(read_document(broken)["status"], "failed")

    def test_previously_confirmed_run_cannot_remain_authoritative_after_failure(self):
        status = self.root / "run_status.json"
        provenance = self.root / "provenance.json"
        status.write_text('{"status":"completed","returncode":0}', encoding="utf-8")
        provenance.write_text('{"script":{"sha256":"synthetic-fixture-only"},"inputs":{}}', encoding="utf-8")
        self.spec["files"] += [{"id": "run", "path": str(status), "role": "run_status"},
                               {"id": "provenance", "path": str(provenance), "role": "provenance"}]
        self.spec["facts"][0]["authority"].update(kind="confirmed_run", run_status="run", provenance="provenance")
        register(self.root, self.spec)
        status.write_text('{"status":"failed","returncode":1}', encoding="utf-8")
        report = check(self.root)
        self.assertTrue(all(x["status"] == "review" for x in report["findings"] if x.get("fact_id") == "beta"))

    def test_native_deck_pair_detects_labels_and_units_even_when_values_match(self):
        import subprocess
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        plan = ROOT / "skills/edu-ppt/examples/chart-fidelity-plan.json"
        build = ROOT / "skills/edu-ppt/scripts/build_deck.py"
        proc = subprocess.run([sys.executable, "-X", "utf8", str(build), "--plan", str(plan),
                               "--out", str(self.root / "deck")], capture_output=True)
        self.assertEqual(proc.returncode, 0, proc.stderr)
        deck = self.root / "deck/slides.pptx"
        self.spec["files"] += [{"id": "ppt", "path": str(deck), "role": "presentation"},
                               {"id": "plan", "path": str(plan), "role": "presentation_plan"}]
        self.spec["decks"] = [{"ppt": "ppt", "plan": "plan"}]
        report = register(self.root, self.spec)
        self.assertEqual(report["deck_checks"][0]["result"]["plan_consistency"], "passed")
        prs = Presentation(deck)
        chart_shape = next(x for x in prs.slides[0].shapes if x.has_chart)
        data = CategoryChartData()
        data.categories = ["Renamed baseline", "Renamed follow-up"]
        for i, series in enumerate(chart_shape.chart.series):
            data.add_series("Changed group " + str(i), series.values)
        chart_shape.chart.replace_data(data)
        chart_shape.chart.value_axis.has_title = True
        chart_shape.chart.value_axis.axis_title.text_frame.text = "wrong unit"
        prs.save(deck)
        report = check(self.root)
        errors = " ".join(report["deck_checks"][0]["result"]["errors"])
        self.assertEqual(report["deck_checks"][0]["result"]["plan_consistency"], "failed")
        self.assertIn("categor", errors)
        self.assertIn("series", errors)
        self.assertIn("unit", errors)


if __name__ == "__main__":
    unittest.main()
