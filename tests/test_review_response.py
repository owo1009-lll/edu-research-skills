"""Check scripts added in 0.5.0: manuscript consistency, statistical reporting, responses, redline, review reports.

Each test seeds known problems into a small document and asserts that the script reports them, and that clean
material passes.
"""
from pathlib import Path
import json
import subprocess
import sys
import tempfile
import unittest

ROOT = Path(__file__).resolve().parents[1]
SKILLS = ROOT / 'skills'

PAPER_ZH = """# 生成式人工智能使用与学习投入

## 摘要

基于对512名大学生的问卷调查，运用结构方程模型（SEM）检验AI使用与学习投入的关系。AI使用正向预测学习投入（β = 0.42），中介效应占总效应的38.5%。

关键词：生成式人工智能；学习投入

## 一、问题提出

已有研究关注AI工具与学业表现的关系[1]，但对学业投入的作用机制关注不足[3,4]。本研究采用SEM分析。

## 二、研究设计

共回收有效问卷498份。研究采用结构方程模型（SEM）。量表信度良好，见表2。

表1 描述统计

| 变量 | M | SD |
|---|---|---|
| 学习投入 | 3.52 | 0.418 |

## 三、研究结果

AI使用正向预测学习投入（β = 0.418, p = 0.000），中介效应显著（p < .001）。本文认为这一结果与已有研究一致[2]。如图1所示。

## 参考文献

[1] 张三. 论文一[J]. 电化教育研究, 2024(1): 1-10.
[2] 李四. 论文二[J]. 中国电化教育, 2024(2): 2-12.
[3] 王五. 论文三[J]. 开放教育研究, 2023(3): 3-13.
[4] 赵六. 论文四[J]. 现代远程教育研究, 2023(4): 4-14.
[5] 钱七. 论文五[J]. 远程教育杂志, 2022(5): 5-15.
"""


def run(script, args, cwd):
    return subprocess.run([sys.executable, '-X', 'utf8', str(script), *map(str, args)], cwd=cwd,
                          capture_output=True, text=True, encoding='utf-8')


class Folder(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def write(self, name, content):
        (self.dir / name).write_text(content, encoding='utf-8')
        return name


class Consistency(Folder):
    SCRIPT = SKILLS / 'edu-writing/scripts/check_consistency.py'

    def report(self, name, *extra):
        p = run(self.SCRIPT, ['--draft', name, '--json', *extra], self.dir)
        self.assertEqual(p.returncode, 0, p.stderr)
        return json.loads(p.stdout)

    def test_seeded_chinese_manuscript(self):
        r = self.report(self.write('p.md', PAPER_ZH), '--term-group', '学习投入=学习投入|学业投入')
        self.assertEqual([h['value'] for h in r['abstract_numbers_not_in_body']], ['512', '38.5'])
        self.assertEqual(set(r['sample_sizes']), {'512', '498'})
        self.assertIn({'value': 'β = 0.418', 'also_printed_as': '0.42'}, r['precision_drift'])
        self.assertTrue(r['p_equals_zero'] and r['p_value_styles'])
        self.assertEqual(r['term_groups'][0]['counts'], {'学习投入': 4, '学业投入': 1})
        self.assertEqual([(a['acronym'], a['issue']) for a in r['acronyms']], [('SEM', 'used in the body before its definition')])
        issues = {(d['display'], d['issue']) for d in r['displays']}
        self.assertIn(('table 2', 'mentioned but no caption found'), issues)
        self.assertIn(('table 1', 'captioned but never mentioned in the text'), issues)
        self.assertIn(('figure 1', 'mentioned but no caption found'), issues)
        cit = {c['issue']: c for c in r['citations']}
        self.assertEqual(cit['in the reference list but never cited']['numbers'], [5])
        self.assertEqual(cit['not numbered in order of first citation']['first_mismatch'], 3)

    def test_keyword_variants_found_without_term_group(self):
        r = self.report(self.write('p.md', PAPER_ZH))
        self.assertEqual(r['term_groups'][0]['group'], 'keyword 学习投入')
        self.assertEqual(r['term_groups'][0]['counts'], {'学习投入': 4, '学业投入': 1})

    def test_consistent_english_manuscript_passes(self):
        text = ('# Title\n\nAbstract: A survey of 312 students used structural equation modeling (SEM). AI use predicted '
                'engagement (beta = .31).\n\nKeywords: AI\n\n## Method\n\nParticipants were 312 students. Structural equation modeling (SEM) was '
                'estimated as shown in Table 1.\n\nTable 1. Estimates\n\n## Results\n\nAI use predicted engagement '
                '(beta = .31, 95% CI [.20, .42], p < .001).\n')
        r = self.report(self.write('e.md', text), '--fail-on-findings')
        self.assertEqual(r['abstract_numbers_not_in_body'], [])
        self.assertEqual(r['displays'] + r['acronyms'] + r['citations'] + r['precision_drift'], [])


class Reporting(Folder):
    def test_seeded_statements(self):
        text = '\n'.join([
            '本研究采用问卷调查，使用SPSS 26.0分析。',
            '实验组后测成绩显著高于对照组，t = 2.31，p < 0.05。',
            '性别与AI使用存在关联，χ² = 8.4，p = 0.004。',
            'AI使用正向预测学习投入（β = 0.42, p = 0.000）。',
            '自我效能感的中介效应显著，间接效应为0.15。',
            '教师支持显著提升了学生的学习投入。',
            '已有研究表明教师支持显著影响学生投入。',
            '事后比较采用LSD法。模型拟合良好，χ²/df = 2.31，CFI = 0.95，RMSEA = 0.05。',
            'AI use predicted engagement (B = 0.31, SE = 0.05, p < .001); the indirect effect was 0.12, 95% CI [0.05, 0.20].'])
        p = run(SKILLS / 'edu-audit/scripts/check_reporting.py', ['--draft', self.write('r.md', text), '--json'], self.dir)
        r = json.loads(p.stdout)
        kinds = [h['check'] for h in r['sentences']]
        for expected in ['t without degrees of freedom', 'chi-square without degrees of freedom', 'p printed as zero (write p < .001)',
                         'indirect effect without a (bootstrap) confidence interval', 'significance claimed without a statistic or table reference']:
            self.assertIn(expected, kinds)
        flagged = ' '.join(h['sentence'] for h in r['sentences'])
        self.assertNotIn('已有研究表明', flagged)  # restating the literature is not a reported result
        self.assertNotIn('SE = 0.05', flagged)
        missing = {d['analysis']: d['not_found'] for d in r['document']}
        self.assertIn('common-method bias check', missing['questionnaire'])
        self.assertIn('fit: SRMR', missing['sem'])
        self.assertIn('LSD does not correct for multiple comparisons', sum(missing.values(), []))


class Response(Folder):
    def test_response_checks_and_redline(self):
        self.write('orig.md', '# 题目\n\n## 二、研究设计\n\n共回收有效问卷498份。\n\n## 三、研究结果\n\nAI使用影响学习投入（β = 0.42）。\n')
        self.write('rev.md', '# 题目\n\n## 二、研究设计\n\n共回收有效问卷498份，采用Harman单因子检验，第一个因子解释了28.6%的方差。\n\n'
                             '## 三、研究结果\n\nAI使用正向预测学习投入（β = 0.42，95% CI [0.31, 0.53]）。\n\n本研究为横断面设计。\n')
        self.write('c.csv', 'id,reviewer,comment\nR1-1,1,未报告共同方法偏差检验。\nR1-2,1,"""影响""一词有因果含义。"\n'
                            'R2-1,2,请报告置信区间。\nR2-2,2,样本代表性不足。\n')
        self.write('resp.md', '| 序号 | 审稿意见 | 修改说明 | 修改位置 |\n|---|---|---|---|\n'
                              '| 1-1 | 未报告共同方法偏差检验。 | 已补充Harman单因子检验，第一个因子解释了31.2%的方差。 | 第二部分 |\n'
                              '| 1-2 | "影响"一词有因果含义。 | 审稿人误解了我们的意思，已改为“AI使用正向预测学习投入”。 | 第三部分 |\n'
                              '| 2-1 | 请报告置信区间。 | 已补充：“AI使用正向预测学习投入（β = 0.42，95% CI [0.31, 0.60]）”。 | 第三部分 |\n')
        p = run(SKILLS / 'edu-response/scripts/check_response.py', ['--comments', 'c.csv', '--response', 'resp.md', '--revised', 'rev.md', '--json'], self.dir)
        r = {i['id']: ' | '.join(i['issues']) for i in json.loads(p.stdout)['items']}
        self.assertIn('31.2', r['R1-1'])
        self.assertIn('defensive wording', r['R1-2'])
        self.assertNotIn('quoted text', r['R1-2'])  # the quoted revision exists in the revised manuscript
        self.assertIn('quoted text not found', r['R2-1'])
        self.assertEqual(r['R2-2'], 'no response found')

        from docx import Document
        p = run(SKILLS / 'edu-response/scripts/redline.py', ['--original', 'orig.md', '--revised', 'rev.md', '--output', 'red.docx', '--show-deletions'], self.dir)
        self.assertEqual(p.returncode, 0, p.stderr)
        runs = [(r.text, r.font.color.rgb is not None, bool(r.font.strike)) for para in Document(self.dir / 'red.docx').paragraphs for r in para.runs]
        self.assertIn(('影响', True, True), runs)
        self.assertIn(('正向预测', True, False), runs)
        self.assertIn(('本研究为横断面设计。', True, False), runs)
        self.assertIn(('共回收有效问卷498份', False, False), runs)
        self.assertIn('modified', (self.dir / 'red.docx.changes.csv').read_text(encoding='utf-8-sig'))


    def test_placeholders_deleted_quotes_and_procedural_items(self):
        self.write('rev.md', '# 题目\n\n## 二、研究设计\n\n采用 Harman 单因子检验。【待补①：第一个因子解释的方差比例】\n')
        self.write('c.csv', 'id,reviewer,comment\nE-1,编辑,请逐条回复并标红。\nR1-1,1,未报告共同方法偏差检验。\nR1-2,1,删去首次等说法。\n')
        self.write('resp.md', '| 序号 | 审稿意见 | 修改说明 | 修改位置 |\n|---|---|---|---|\n'
                              '| E-1 | 请逐条回复并标红。 | 已按要求逐条回复，修改处见修改稿（标红版）。 | 本修改说明 |\n'
                              '| R1-1 | 未报告共同方法偏差检验。 | 已补充 Harman 单因子检验，【待补①】。 | 第二部分 |\n'
                              '| R1-2 | 删去首次等说法。 | 已删去原文“本研究首次从自我效能感视角揭示了机制”。 | 第一部分 |\n'
                              '\n## 其他修改\n\n统一了术语写法。\n')
        p = run(SKILLS / 'edu-response/scripts/check_response.py', ['--comments', 'c.csv', '--response', 'resp.md', '--revised', 'rev.md', '--json'], self.dir)
        out = json.loads(p.stdout)
        r = {i['id']: ' | '.join(i['issues']) for i in out['items']}
        self.assertEqual(r['E-1'], '')  # procedural items point to the response or the red-marked copy
        self.assertIn('placeholder in the reply that claims the change is done', r['R1-1'])
        self.assertEqual(r['R1-2'], '')  # quoting the deleted wording is not a missing revision
        self.assertEqual(out['placeholders'], {'response': 1, 'revised': 1})

    def test_redline_keeps_tables_in_place_and_whole_numbers(self):
        from docx import Document
        from docx.oxml.ns import qn
        self.write('o.md', '# 题目\n\n第一段。\n\n| 变量 | M |\n|---|---|\n| 学习投入 | 3.52 |\n\n路径系数为0.42。\n\n## 参考文献\n')
        self.write('n.md', '# 题目\n\n第一段。\n\n| 变量 | M |\n|---|---|\n| 学习投入 | 3.61 |\n\n路径系数为0.418。\n\n## 参考文献\n')
        p = run(SKILLS / 'edu-response/scripts/redline.py', ['--original', 'o.md', '--revised', 'n.md', '--output', 'r.docx'], self.dir)
        self.assertEqual(p.returncode, 0, p.stderr)
        d = Document(self.dir / 'r.docx')
        order = ''.join('T' if c.tag == qn('w:tbl') else 'p' for c in d.element.body.iterchildren() if c.tag in (qn('w:tbl'), qn('w:p')))
        self.assertTrue(order.startswith('ppTp'), order)  # title, paragraph, table, then the following paragraph
        self.assertEqual(d.paragraphs[0].style.name, 'Title')
        red = [r.text for para in d.paragraphs for r in para.runs if r.font.color.rgb is not None]
        self.assertIn('0.418', red)
        cell_red = [r.text for r in d.tables[0].rows[1].cells[1].paragraphs[0].runs if r.font.color.rgb is not None]
        self.assertEqual(cell_red, ['3.61'])

    def test_reference_list_labels_not_superscript(self):
        from docx import Document
        self.write('a.md', '# 题目\n\n已有研究[1]。\n\n## 参考文献\n\n[1] 张三. 论文[J]. 刊物, 2024(1): 1-2.\n')
        run(SKILLS / 'edu-writing/scripts/md_to_docx.py', ['--input', 'a.md', '--output', 'a.docx', '--superscript-citations'], self.dir)
        paras = Document(self.dir / 'a.docx').paragraphs
        self.assertTrue(any(r.font.superscript for r in paras[1].runs))
        self.assertFalse(any(r.font.superscript for r in paras[-1].runs))


class RealManuscriptConventions(Folder):
    """Conventions of real Chinese manuscripts that broke the checks (full-width signs, Word headings, numbering)."""

    def test_fullwidth_signs_table_mentions_and_numbering(self):
        doc = ('# 题目\n\n## 摘要\n\n第一部分（n＝347）：结果显示d*＝0.36。\n\n关键词：甲；乙\n\n## 1 引言\n\n'
               '本文围绕两个研究问题展开。RQ1：……？RQ2：……？\n\n## 4 结果\n\n表1列出描述统计（RQ1）。\n\n表1 描述统计\n\n'
               '| 变量 | M (SD) |\n|---|---|\n| 甲 | 3.1 (1.2) |\n\n调整差为0.549，d*＝0.364，Holm p＜0.001。'
               '主计分模型CFI＝1.000，全条目模型CFI＝0.987。不同年级存在差异，F＝4.87，p＝0.000。\n')
        name = self.write('fw.md', doc)
        r = json.loads(run(SKILLS / 'edu-writing/scripts/check_consistency.py', ['--draft', name, '--json'], self.dir).stdout)
        self.assertIn('347', r['sample_sizes'])  # n＝347 with a full-width equals sign
        self.assertEqual(r['precision_drift'], [{'value': 'd* = 0.364', 'also_printed_as': '0.36'}])  # not CFI 0.987 / 1.000
        self.assertEqual(r['acronyms'], [])  # RQ1 and M (SD) are not abbreviation definitions
        self.assertEqual(r['displays'], [])  # 表1列出… at the start of a line is a mention
        rep = json.loads(run(SKILLS / 'edu-audit/scripts/check_reporting.py', ['--draft', name, '--json'], self.dir).stdout)
        kinds = [h['check'] for h in rep['sentences']]
        self.assertIn('F without degrees of freedom', kinds)
        self.assertIn('p printed as zero (write p < .001)', kinds)

    def test_voice_check_stops_at_back_matter_in_word(self):
        from docx import Document
        d = Document()
        d.add_heading('5 讨论', 1)
        d.add_paragraph('研究发现，练习与投入正相关。可能的原因在于练习提供了反馈。')
        d.add_heading('声明', 1)
        d.add_paragraph('本研究存在局限，样本量较小，未来研究可扩大样本。')
        d.add_heading('参考文献', 1)
        d.add_paragraph('张三. 推广性研究的局限[J]. 刊物, 2024.')
        d.save(self.dir / 'v.docx')
        out = run(SKILLS / 'edu-writing/scripts/check_voice.py', ['--draft', 'v.docx'], self.dir).stdout
        self.assertIn('Limitation sentences outside the limitations unit: 0', out)


class Prose(Folder):
    """check_prose.py flags AI-typical writing and leaves published-style prose alone."""
    INTRO = ('## 一、问题提出\n\n近年来，生成式人工智能进入大学课堂，学习方式随之变化[1]。已有研究主要从使用意愿和学业表现两方面展开[2,3]。'
             '然而，现有研究多聚焦使用频率，较少关注使用方式[4]。基于此，本研究以三所高校学生为对象，探讨使用方式与学习投入的关系，'
             '以期为课程设计提供依据[5]。\n')
    GOOD_DISCUSSION = ('## 五、讨论\n\n### （一）自主型使用与学习投入正相关\n\n'
                       '本研究发现，自主型使用与学习投入正相关。这与已有研究的结论一致[6]。可以从自我决定理论的视角理解：学习者保留了判断与选择。'
                       '这说明使用方式比使用频率更能解释投入的差异。教师可以在任务中保留学生的选择环节。\n')
    WEAK_DISCUSSION = ('## 五、讨论\n\n下面分三个方面讨论。两组均值分别为20.55分和19.05分，差值为1.50分，95%置信区间为-6.07至3.07分，'
                       'η²=0.32，两种方式孰优孰劣仍是开放的问题。这一比较可以在现有数据上完成。两种解释可以分别检验：前者预测A，后者预测B。'
                       '学习者的平均得分为3.60，标准差为0.82。本研究认为该结果具有重要意义。\n')

    def flags(self, text):
        name = self.write('d.md', '# 题目\n\n' + text)
        r = json.loads(run(SKILLS / 'edu-writing/scripts/check_prose.py', ['--draft', name, '--json'], self.dir).stdout)
        return ' | '.join([f for s in r['sections'].values() for f in s['flags']] + r['document'])

    def test_weak_discussion_is_flagged(self):
        out = self.flags(self.INTRO + self.WEAK_DISCUSSION)
        self.assertIn('reviewer-style sentences', out)
        self.assertIn('signposting', out)
        self.assertIn('decimal numbers per 1,000 characters', out)
        self.assertIn('neither cites nor compares', out)

    def test_published_style_passes(self):
        self.assertEqual(self.flags(self.INTRO + self.GOOD_DISCUSSION), '')

    def test_no_chinese_references(self):
        refs = '\n## 参考文献\n\n' + ''.join(f'[{i}] Author A. A title about learning number {i}[J]. Journal, 2020: 1-10.\n'
                                            for i in range(1, 12))
        self.assertIn('none of the 11 references is Chinese', self.flags(self.INTRO + self.GOOD_DISCUSSION + refs))


class ProseEnglish(Folder):
    """English drafts: check_prose.py detects the language and compares with the 31 SSCI articles."""
    INTRO = ('## 1. Introduction\n\nGenerative AI tools have become a common learning resource for university students (Kasneci et al., 2023). '
             'Early studies report gains in performance while students use the tools, but weaker performance once access is withdrawn '
             '(Bastani et al., 2025; Fan et al., 2025). Engagement is a useful outcome for examining this pattern because it responds to '
             'features of the learning environment (Fredricks et al., 2004). Teacher support is one such feature, and it predicts '
             'engagement across school levels (Skinner & Belmont, 1993; Wang & Eccles, 2012). Most studies have examined access to the '
             "tools; how students' use of the tools relates to engagement alongside teacher support remains unexamined. Prior studies "
             'also relied on net-effect models, which cannot show whether a condition is necessary for high engagement (Dul, 2016). '
             'In this study, we examine these relationships among undergraduates using structural equation modeling and necessary '
             'condition analysis. In doing so, we contribute to research on technology-supported learning by separating the routes '
             'through which teacher support and tool use relate to engagement.\n\n')
    GOOD_REVIEW = ('## 2. Teacher support and engagement\n\nStudies consistently show that students who perceive more support from '
                   'their teachers report higher engagement (Skinner & Belmont, 1993; Wang & Eccles, 2012; Li & Chiu, 2025). Self-efficacy '
                   'is one belief through which this support may operate, because feedback is a situational source of efficacy beliefs '
                   '(Schunk, 1991). Efficacious students engage more fully in classroom learning (Linnenbrink & Pintrich, 2003). '
                   'Li and Chiu (2025) found that needs satisfaction carried part of the association between teacher support and '
                   'engagement with a chatbot among university students. However, the mediating belief has been specified differently '
                   'across studies, and capability judgments for course tasks have rarely been modeled (Schunk, 1991). Based on '
                   'self-efficacy theory, we hypothesize that academic self-efficacy mediates the relationship between teacher support '
                   'and engagement in university courses.\n\n')
    GOOD_DISCUSSION = ('## 5. Discussion\n\nTeacher support was related to engagement both directly and through academic self-efficacy. '
                       'This is consistent with earlier studies of school and university students (Skinner & Belmont, 1993; Li & Chiu, 2025). '
                       "Self-efficacy theory offers one explanation: feedback from teachers is a situational source of students' efficacy "
                       'beliefs, and efficacious students engage more fully (Linnenbrink & Pintrich, 2003; Schunk, 1991). The direct path '
                       'that remained may reflect needs for relatedness and autonomy that self-efficacy does not capture (Ryan & Deci, 2000). '
                       'Contrary to H5, use of the tools was not directly related to engagement. This differs from studies that reported '
                       'adverse associations of heavy use (Abbas et al., 2024). One possible reason is that the items recorded how much '
                       'students used the tools rather than how they used them. A study that logs modes of use could test this explanation '
                       'and show whether checking and bypassing uses relate to engagement in opposite directions.\n')
    WEAK_REVIEW = ('## 2. Literature review\n\nSmith (2020) found that teacher support predicted engagement in a large sample of '
                   'secondary schools in three regions. Jones (2021) showed that self-efficacy mediated this association in a sample of '
                   'undergraduates at one university. Lee (2022) reported similar results for online courses with adult learners in '
                   'professional programmes. Kim (2023) argued that feedback is the active ingredient of teacher support in these '
                   'settings and in others. Wang (2024) found that the effect was weaker in large classes than in small seminars for '
                   'first-year students. Chen (2024) demonstrated that the association held across disciplines and across years of '
                   'study. These studies together suggest that teacher support matters for engagement in many settings. A growing body '
                   'of work also points to generative AI as a new resource for learners. Zhao (2025) found that students who used '
                   'chatbots during seminars asked their teachers fewer questions. Liu (2025) reported that this pattern was stronger '
                   'among first-year students than among older students.\n\n')
    WEAK_DISCUSSION = ("## 5. Discussion\n\nIn today's rapidly changing educational landscape, it is worth noting that our results delve "
                       'into the role of AI. Teacher support predicted engagement (β = .245, 95% CI [.093, .387]) and self-efficacy '
                       'predicted engagement (β = .447, 95% CI [.303, .585]). The indirect effect was .187 [.114, .279], which was 43% '
                       'of the total effect of .432 in the full sample of 300 students. AI use predicted self-efficacy (β = .181) but '
                       'not engagement (β = .104, p = .11) once teacher support was included. These findings pave the way for future '
                       'research on generative AI in higher education and in schools. The two explanations make different predictions, and '
                       'this could be tested with the present data by splitting the sample into heavy and light users. Overall, the '
                       'results provide valuable insights and highlight the importance of teacher support for students and teachers. '
                       'Engagement averaged 3.62 (SD = 0.71) and self-efficacy averaged 3.48 (SD = 0.69) on the five-point scales, '
                       'and both scales had acceptable reliability in this sample of undergraduates from one university.\n')

    def report(self, text):
        name = self.write('d.md', '# Title\n\n' + text)
        return json.loads(run(SKILLS / 'edu-writing/scripts/check_prose.py', ['--draft', name, '--json'], self.dir).stdout)

    def flags(self, text):
        r = self.report(text)
        return ' | '.join([f for s in r['sections'].values() for f in s['flags']] + r['document'])

    def test_sections_found_from_topic_headings(self):
        # "2. Teacher support and engagement" is a review section; the Method placeholder is skipped
        r = self.report(self.INTRO + self.GOOD_REVIEW + '## 3. Method\n\n[Written by the authors.]\n\n' + self.GOOD_DISCUSSION)
        self.assertEqual(r['language'], 'en')
        self.assertEqual(list(r['sections']), ['intro', 'review', 'discussion'])

    def test_published_style_passes(self):
        self.assertEqual(self.flags(self.INTRO + self.GOOD_REVIEW + self.GOOD_DISCUSSION), '')

    def test_weak_prose_is_flagged(self):
        out = self.flags(self.INTRO + self.WEAK_REVIEW + self.WEAK_DISCUSSION)
        for expected in ['consecutive sentences open with "Author (year)"', 'narrative_share', 'cites no earlier study',
                         'number_share', 'propose analyses', 'rare stock phrases']:
            self.assertIn(expected, out)


class Review(Folder):
    MAJOR = ('### {id}\n- 是否阻断：是\n- 维度：数据分析\n- 论断位置：第三部分第 2 段\n- 证据位置：{ev}\n- 问题：中介只用逐步法。\n'
             '- 为什么重要：检验力低。\n- 怎样算解决：{res}\n')

    def test_structure_checks(self):
        script = SKILLS / 'edu-review/scripts/check_review.py'
        self.write('R1.md', '# R1\n\n## 主要问题\n\n' + self.MAJOR.format(id='R1-M1', ev='表 3', res='报告 bootstrap 置信区间。')
                   + '\n## 次要问题\n\n### R1-m1\n- 位置：摘要\n- 问题与建议：p 值格式。\n\n## 建议\n大修。\n')
        self.write('R2.md', '# R2\n\n' + self.MAJOR.format(id='R2-M1', ev='第一部分第 2 段', res='写出与已有研究的区别。'))
        self.write('s.md', '# 综合（AI 模拟审稿）\n\n- R1-M1、R2-M1\n')
        p = run(script, ['--reports', 'R1.md', 'R2.md', '--synthesis', 's.md'], self.dir)
        self.assertEqual(p.returncode, 0, p.stdout)
        self.write('R2.md', '# R2\n\n' + self.MAJOR.format(id='R2-M1', ev='文中多处', res='……') + '\n预计 60% 的可能录用。\n')
        self.write('s.md', '# 综合\n\n- R2-M3\n')
        p = run(script, ['--reports', 'R1.md', 'R2.md', '--synthesis', 's.md', '--json'], self.dir)
        issues = ' | '.join(json.loads(p.stdout)['issues'])
        for expected in ['R2-M1: missing field resolution', 'evidence field names no location', 'acceptance-probability',
                         'cites R2-M3', 'does not say the review is simulated']:
            self.assertIn(expected, issues)

    def test_numbered_sections_back_matter_and_evidence_pointing_to_concern(self):
        script = SKILLS / 'edu-review/scripts/check_review.py'
        block = ('### R1-M{n}\n- 是否阻断：是\n- 维度：研究伦理\n- 论断位置：摘要\n- 证据位置：{ev}\n'
                 '- 问题：{concern}\n- 为什么重要：无法刊出。\n- 怎样算解决：补齐审批信息。\n')
        report = '# R1（AI 模拟审稿）\n\n' + ''.join([
            block.format(n=1, ev='3.7节第1句；声明第1项', concern='占位未填。'),
            block.format(n=2, ev='3.7首句的方括号', concern='占位未填。'),
            block.format(n=3, ev='见下文"问题"中列出的位置', concern='4.7节与表8中的写法不一。')])
        self.write('R1.md', report)
        self.write('s.md', '# 综合（AI 模拟审稿）\n- R1-M1、R1-M2、R1-M3\n')
        r = json.loads(run(script, ['--reports', 'R1.md', '--synthesis', 's.md', '--json'], self.dir).stdout)
        self.assertEqual(r['issues'], [])

    def test_empty_field_negated_probability_and_multiline_values(self):
        script = SKILLS / 'edu-review/scripts/check_review.py'
        block = ('### R1-M1\n- 是否阻断：是\n- 维度：数据分析\n- 论断位置：\n  - 摘要第 2 句；\n  - 第五部分第 1 段\n'
                 '- 证据位置：表 1\n- 问题：\n  1. 中介只用逐步法。\n- 为什么重要：\n- 怎样算解决：报告区间。\n')
        self.write('R1.md', '# R1\n\n' + block + '\n## 建议\n大修。这是模拟审稿的建议，不预测录用概率。\n')
        self.write('s.md', '# 综合（AI 模拟审稿，不预测录用概率）\n\n- R1-M1\n')
        r = json.loads(run(script, ['--reports', 'R1.md', '--synthesis', 's.md', '--json'], self.dir).stdout)
        self.assertEqual(r['issues'], ['R1-M1: missing field why'])


if __name__ == '__main__':
    unittest.main()
