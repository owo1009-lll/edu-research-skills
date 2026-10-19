"""Annotation anchors and source preservation, including adversarial mutations."""
from copy import deepcopy
import importlib.util
from pathlib import Path
import tempfile
import unittest
import zipfile

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('fixtures', ROOT / 'evals/review-annotations/create_fixture.py')
fixtures = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixtures)
a = fixtures.annotator


class Annotations(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def source(self, extension):
        path = self.folder / ('original.' + extension)
        (fixtures.create_docx if extension == 'docx' else fixtures.create_pdf)(path)
        return path, fixtures.plan_for(path)

    def altered_zip(self, path, part, mutate):
        dest = self.folder / 'altered.docx'
        with zipfile.ZipFile(path) as before, zipfile.ZipFile(dest, 'w') as after:
            for item in before.infolist():
                content = before.read(item.filename)
                after.writestr(item, mutate(content) if item.filename == part else content)
        return dest

    def test_word_preserves_text_tables_images_fields_comments_and_revisions(self):
        source, plan = self.source('docx')
        source_bytes = source.read_bytes()
        result = a.apply(source, plan)
        output = Path(result['output'])
        self.assertEqual(source_bytes, source.read_bytes())
        before, after = a.package(source), a.package(output)
        b, c = a.xml(before[a.DOC]), a.xml(after[a.DOC])
        self.assertEqual([a.ptext(p) for p in a.paragraphs(b)], [a.ptext(p) for p in a.paragraphs(c)])
        for tag in ('ins', 'del', 'fldSimple', 'drawing', 'tblPr', 'tcPr'):
            self.assertEqual([a.canonical(n) for n in b.iter(a.w(tag))],
                             [a.canonical(n) for n in c.iter(a.w(tag))], tag)
        for name in before:
            if name.startswith('word/media/'):
                self.assertEqual(before[name], after[name])
        old_comment = a.xml(before['word/comments.xml'])[0]
        self.assertEqual(a.canonical(old_comment), a.canonical(a.xml(after['word/comments.xml'])[0]))
        self.assertEqual(result['annotations'], 4)
        colors = [e.get(a.w('val')) for e in c.iter(a.w('highlight'))]
        self.assertTrue({'red', 'yellow', 'blue', 'green'} <= set(colors))
        self.assertTrue(output.stem.endswith('_审稿标注版'))
        from docx import Document
        reopened = Document(output)
        self.assertEqual(len(reopened.comments), 5)
        self.assertEqual(reopened.tables[0].rows[-1].cells[-1].text, '48')
        self.assertEqual(len(reopened.inline_shapes), 1)

    def test_word_without_comment_part_and_partial_multirun_anchor(self):
        from docx import Document
        path = self.folder / 'simple.docx'
        d = Document()
        p = d.add_paragraph('Before ')
        p.add_run('careful').italic = True
        p.add_run(' interpretation after.')
        d.save(path)
        inv = a.inventory(path)
        plan = {'schema_version': 1, 'source_sha256': inv['source_sha256'],
                'created_utc': '2026-10-19T00:00:00Z', 'annotations': [
                    {'id': 'C1', 'kind': 'improvement', 'comment': 'Clarify this interpretation.',
                     'anchors': [fixtures.anchor(inv, 'ful interpretation')]}]}
        result = a.apply(path, plan)
        self.assertTrue(a.verify(path, result['output'], plan)['preserved_source'])
        d = Document(result['output'])
        self.assertEqual(d.paragraphs[0].text, 'Before careful interpretation after.')
        self.assertEqual(len(d.comments), 1)

    def test_wrong_anchor_and_changed_source_rejected_without_output(self):
        source, plan = self.source('docx')
        wrong = deepcopy(plan)
        wrong['annotations'][0]['anchors'][0]['quote'] = '99'
        with self.assertRaisesRegex(ValueError, 'expected.*99.*actual.*52'):
            a.apply(source, wrong)
        wrong = deepcopy(plan)
        wrong['source_sha256'] = '0' * 64
        with self.assertRaisesRegex(ValueError, 'Source SHA-256 changed'):
            a.apply(source, wrong)
        self.assertEqual(list(self.folder.glob('*审稿标注版*')), [])

    def test_overlap_and_unsupported_field_anchor_rejected(self):
        source, plan = self.source('docx')
        plan['annotations'].append(deepcopy(plan['annotations'][0]))
        plan['annotations'][-1]['id'] = 'overlap'
        with self.assertRaisesRegex(ValueError, 'overlapping'):
            a.apply(source, plan)
        plan['annotations'] = [plan['annotations'][0]]
        plan['annotations'][0]['anchors'] = [fixtures.anchor(a.inventory(source), '[1]', contains='Citation field')]
        with self.assertRaisesRegex(ValueError, 'field|non-plain'):
            a.apply(source, plan)

    def test_word_tampering_with_text_comment_revision_or_image_fails(self):
        source, plan = self.source('docx')
        output = Path(a.apply(source, plan)['output'])
        for part, old, new in [
                (a.DOC, b'>48<', b'>49<'),
                (a.DOC, b'old wording', b'changed wording'),
                ('word/comments.xml', b'Original author comment', b'Removed original comment'),
                ('word/media/image1.png', b'PNG', b'BAD')]:
            with self.subTest(part=part, old=old):
                changed = self.altered_zip(output, part, lambda data: data.replace(old, new))
                with self.assertRaisesRegex(ValueError, 'differ'):
                    a.verify(source, changed, plan)

    def test_empty_plan_passes_and_never_overwrites(self):
        for extension in ('docx', 'pdf'):
            with self.subTest(extension=extension):
                source, plan = self.source(extension)
                plan['annotations'] = []
                output = Path(a.apply(source, plan)['output'])
                self.assertEqual(a.verify(source, output, plan)['annotations'], 0)
                with self.assertRaisesRegex(ValueError, 'already exists'):
                    a.apply(source, plan)
                with self.assertRaisesRegex(ValueError, 'equals source'):
                    a.apply(source, plan, source)

    def test_pdf_native_comments_existing_note_and_exact_repeated_number(self):
        import fitz
        source, plan = self.source('pdf')
        raw = source.read_bytes()
        result = a.apply(source, plan)
        output = Path(result['output'])
        self.assertEqual(raw, source.read_bytes())
        self.assertTrue(output.read_bytes().startswith(raw))
        self.assertEqual(result['annotations'], 3)
        with fitz.open(output) as doc:
            page = doc[0]
            annots = list(page.annots())
            self.assertEqual(annots[0].info['title'], 'Original Author')
            self.assertEqual([an.info['title'] for an in annots[1:]], [a.AUTHOR] * 3)
            self.assertTrue(all(an.type[0] == fitz.PDF_ANNOT_HIGHLIGHT for an in annots[1:]))
            self.assertTrue(any('明确错误' in an.info['content'] for an in annots))

    def test_pdf_rotated_and_multiline_anchors(self):
        source, plan = self.source('pdf')
        inv = a.inventory(source)
        plan['annotations'] = [
            {'id': 'rotated', 'kind': 'verify', 'comment': 'Check the rotated text.',
             'anchors': [fixtures.anchor(inv, '48', contains='Rotated')]},
            {'id': 'rows', 'kind': 'improvement', 'comment': 'Explain these groups.',
             'anchors': [fixtures.anchor(inv, 'Group A: 24'), fixtures.anchor(inv, 'Group B: 24')]},
        ]
        self.assertEqual(a.apply(source, plan)['annotations'], 2)

    def test_pdf_modified_comment_or_page_content_fails(self):
        import fitz
        source, plan = self.source('pdf')
        output = Path(a.apply(source, plan)['output'])
        for change in ('comment', 'page', 'invisible'):
            with self.subTest(change=change):
                dest = self.folder / (change + '.pdf')
                dest.write_bytes(output.read_bytes())
                with fitz.open(dest) as doc:
                    page = doc[0]
                    if change == 'comment':
                        an = list(page.annots())[-1]
                        an.set_info(content='Unplanned comment')
                        an.update()
                    elif change == 'invisible':
                        an = list(page.annots())[-1]
                        an.set_opacity(0)
                        an.update()
                    else:
                        page.insert_text((40, 700), 'Unplanned replacement data')
                    doc.saveIncr()
                with self.assertRaises(ValueError):
                    a.verify(source, dest, plan)


if __name__ == '__main__':
    unittest.main()
